"""Presentation decks: CRUD and .pptx export. See app/models/decks.py for
the slides_json shape.

A slide is a plain dict:
    {"kind": "title" | "statement" | "content" | "questions",
     "heading": str, "bullets": [str], "evidence": [{"record_type","record_id"}],
     "notes": str, "visual": None | {"type": "figures"|"timeline"|"bar"|"chart", ...}}
Only heading/bullets/evidence are required; a slide with no `kind` (decks
saved before kinds existed) renders as "content". `visual` mirrors the
presentation_draft module's per-section visual (app/ai/reasoning.py
VISUAL_TYPES) -- every data point carries its own evidence refs. A "chart"
visual is a resolved order slip (app/ai/chart_check.py): its spec, the data
the Workbench cooked from the records (every row tagged with record refs),
the check report and the chart's words, and its style (preset and options,
app/visuals/style.py) when it has one. It is drawn as a native PowerPoint
chart of its kind (a treemap as the render service's picture), and checked
again whenever the deck is saved and when it is exported.

Export draws every slide itself on a 16:9 canvas rather than relying on the
stock python-pptx template: an editorial layout in the app's Dark Academia
palette (dark title/takeaway slides, warm-paper content slides), real
hanging-indent bullets, the cited evidence quoted beside the bullets, native
(editable) PowerPoint bar charts, a drawn timeline, the section's full prose
in the speaker notes, and a numbered, human-readable sources list -- see
render_deck_pptx.
"""
from __future__ import annotations

import math
from dataclasses import dataclass
from datetime import date
from typing import Any

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.time import utcnow_naive
from app.models.domain import Deck
from app.services.presentation_style import (
    VISUAL_PRESETS, Theme, normalize_deck_style, normalize_visual_style, resolved_theme, resolved_visual_style,
)
from app.visuals.checker import CheckReport, check_visual
from app.visuals.data_service import DataRow, VisualData
from app.visuals.echarts import (
    DEFAULT_HEIGHT, DEFAULT_WIDTH, KIND_LABELS, MIN_FONT_PX, SINGLE_RECORD_MARK, BuiltChart, ChartDoesNotFit, UnsupportedChartKind, build_chart,
    format_value, mark_colour, part_colours, part_value_text, prepare_rows, singular,
)
from app.visuals.layout import TEXT_SLACK, block_width, place_on_line
from app.visuals.spec import ChartKind, VisualSpec, spec_errors
from app.visuals.style import normalize_chart_style, styled_chart

MAX_BULLETS_PER_SLIDE = 8
SLIDE_KINDS = {"title", "statement", "content", "questions"}
VISUAL_LIMITS = {"figures": 4, "timeline": 8, "bar": 8}
CHART_VISUAL = "chart"
# The kinds a deck draws as native PowerPoint charts; the others (the
# treemap) go in as the render service's picture.
NATIVE_KINDS = frozenset({ChartKind.bar, ChartKind.column, ChartKind.line, ChartKind.area, ChartKind.stacked_bar_100,
                          ChartKind.donut, ChartKind.histogram, ChartKind.scatter})


def _clean_refs(refs: Any) -> list[dict]:
    """The {record_type, record_id} pairs whose two parts are non-empty text.
    Every record id is text, so anything else (a deck the editor or a client
    saved with a number there) points at nothing and is dropped, which keeps
    the deck's declared shape (app/schemas/responses.py RecordRefOut) true."""
    if not isinstance(refs, list):
        return []
    return [
        {"record_type": e["record_type"], "record_id": e["record_id"]}
        for e in refs
        if isinstance(e, dict) and isinstance(e.get("record_type"), str) and isinstance(e.get("record_id"), str)
        and e["record_type"] and e["record_id"]
    ]


def normalize_chart_visual(visual: Any) -> dict | None:
    """A resolved chart (app/ai/chart_check.py): spec, data, report, text.
    The data must be there -- a deck never runs a query at export time, so
    a bare order slip with nothing cooked is dropped."""
    if not isinstance(visual, dict) or visual.get("type") != CHART_VISUAL:
        return None
    if spec_errors(visual.get("spec")) or not isinstance(visual.get("data"), dict):
        return None
    data = VisualData.from_dict(visual["data"])
    if not data.rows:
        return None
    report = visual.get("report") if isinstance(visual.get("report"), dict) else {"ok": True, "errors": [], "warnings": []}
    text = visual.get("text") if isinstance(visual.get("text"), dict) else {}
    spec = VisualSpec.model_validate(visual["spec"])
    kept = {"type": CHART_VISUAL, "spec": spec.model_dump(mode="json"), "data": data.as_dict(), "report": report, "text": text}
    if isinstance(visual.get("option"), dict):
        # The checked drawing travels with the deck so the preview of a
        # promoted deck is the one the reporter approved.
        kept["option"] = visual["option"]
    # The chart's preset and options (app/visuals/style.py), kept as the
    # catalogue allows them; the order slip and the data stay as they are.
    style = normalize_chart_style(spec, data, visual.get("style"))
    if style:
        kept["style"] = style
    return kept


def recheck_charts(slides: list[dict], theme: Theme) -> list[dict]:
    """Every chart in a deck drawn again in the deck's theme and its own
    style, and checked again, whenever the deck's slides or look change
    ("the checker re-runs on every change", Visual Editor design, "Keeping
    it honest"). The stored drawing, words and report become the ones the
    deck will export, so a chart a style or a theme has broken is marked as
    failing where it is kept, and the export draws a note naming the rules
    instead of the chart. A chart that cannot be drawn at all is left as it
    was stored; the export says why. New dicts are returned, the given ones
    left as they were: a deck's stored slides are compared with the new ones
    to know what to save."""
    out: list[dict] = []
    for slide in slides:
        visual = slide.get("visual")
        if isinstance(visual, dict) and visual.get("type") == CHART_VISUAL:
            checked = checked_chart(visual, theme)
            if checked.built is not None and checked.report is not None:
                slide = {**slide, "visual": {**visual, "option": checked.built.option, "text": checked.built.text(),
                                             "report": checked.report.as_dict()}}
        out.append(slide)
    return out


def chart_record_refs(visual: dict) -> list[dict]:
    """Every record a chart's rows rest on, in row order."""
    refs: list[dict] = []
    for row in ((visual.get("data") or {}).get("rows") or []):
        refs.extend(_clean_refs(row.get("record_refs")))
    return refs


def chart_rows(visual: dict) -> tuple[VisualSpec, VisualData, list[DataRow]]:
    """A stored chart's spec, data and rows in drawing order."""
    spec = VisualSpec.model_validate(visual["spec"])
    data = VisualData.from_dict(visual["data"])
    return spec, data, prepare_rows(spec, data)


def chart_row_labels(data: VisualData, rows: list[DataRow], *, separator: str = " · ") -> dict[str, str]:
    """Each row's label with how many records sit behind it (design rule
    16) and the mark for a value resting on one record (rule 20), for the
    exports, which write both on the category label. A chart whose labels
    sit side by side under it puts the count on a line of its own
    (`separator` "\\n"), so a period's label stays level instead of being
    turned on its side to fit."""
    labels: dict[str, str] = {}
    for row in rows:
        label = row.label
        if data.unit_kind == "amount":
            label += f"{separator}{row.record_count} {data.record_noun if row.record_count != 1 else singular(data.record_noun)}"
            if row.record_count == 1:
                label += f" {SINGLE_RECORD_MARK}"
        elif row.record_count == 1:
            label += f" {SINGLE_RECORD_MARK}"
        labels[row.label] = label
    return labels


def chart_bar_items(rows: list[DataRow], labels: dict[str, str] | None = None) -> list[dict]:
    """Rows as the bar renderer's items: label, value and the records."""
    return [{"label": (labels or {}).get(row.label, row.label), "value": row.value, "evidence": [r.as_dict() for r in row.record_refs]}
            for row in rows]


@dataclass
class CheckedChart:
    """A stored chart drawn again from its own data and checked once more,
    ready for an export. `problem` says why it must not be drawn."""
    spec: VisualSpec
    data: VisualData
    built: BuiltChart | None
    report: CheckReport | None
    problem: str | None


def checked_chart(visual: dict, theme: Theme, *, width: int = DEFAULT_WIDTH, height: int = DEFAULT_HEIGHT) -> CheckedChart:
    """Draw a stored chart again from the rows the data service cooked (no
    query runs at export time) in the export's theme, and check it before it
    goes in. A chart that fails a hard rule is not drawn: its place says which
    rules failed, in plain words ("Export refuses broken files", Visual
    Editor design, "Keeping it honest"). The deck or brief itself still
    exports. `width` and `height` are the picture's design size: a picture
    drawn for a slide is drawn, and checked, at the slide's own scale."""
    spec = VisualSpec.model_validate(visual["spec"])
    data = VisualData.from_dict(visual["data"])
    try:
        # The chart's own style (preset and options) on top of its order slip;
        # the rows are drawn as the data service cooked them.
        spec, draw = styled_chart(spec, data, visual.get("style"))
        built = build_chart(spec, data, theme, width=width, height=height, draw=draw)
    except (UnsupportedChartKind, ChartDoesNotFit) as exc:
        return CheckedChart(spec, data, None, None, f"Chart not drawn: {exc}.")
    report = check_visual(spec, data, theme, built)
    if not report.ok:
        failed = "; ".join(f"rule {e.rule}: {e.message}" for e in report.errors)
        return CheckedChart(spec, data, built, report, f"Chart not drawn: it fails the design checks ({failed}).")
    return CheckedChart(spec, data, built, report, None)


def _set_alt_text(shape: Any, text: str) -> None:
    """A chart frame's or a picture's alt text (rule 13). python-pptx has no
    setting for it; PowerPoint keeps it in the shape's cNvPr `descr`."""
    for props in shape._element.xpath("./p:nvGraphicFramePr/p:cNvPr | ./p:nvPicPr/p:cNvPr"):
        props.set("descr", text)


def _set_fill_alpha(fill: Any, percent: int) -> None:
    """Make a solid fill partly transparent (python-pptx has no setting for
    it, its issue #62): DrawingML's <a:alpha> on the fill colour."""
    from lxml import etree
    from pptx.oxml.ns import qn
    for colour in fill._xPr.xpath("./a:solidFill/a:srgbClr"):
        for old in colour.findall(qn("a:alpha")):
            colour.remove(old)
        etree.SubElement(colour, qn("a:alpha")).set("val", str(percent * 1000))


def _set_hole_size(plot: Any, percent: int) -> None:
    """A doughnut's hole as a percentage of its diameter (c:holeSize, which
    python-pptx writes at 50 and does not expose)."""
    from pptx.oxml.ns import qn
    hole = plot._element.find(qn("c:holeSize"))
    if hole is not None:
        hole.set("val", str(percent))


def _pin_plot_area(chart: Any, x: float, y: float, w: float, h: float) -> None:
    """Fix the plot area's inner rectangle (the bars or the ring, without
    axis labels) at these fractions of the chart frame
    (c:plotArea/c:layout/c:manualLayout, layoutTarget inner, edge mode), so
    a label can be moved to a place worked out beforehand. python-pptx has
    no setting for it."""
    from lxml import etree
    from pptx.oxml.ns import qn
    plot_area = chart._chartSpace.find(qn("c:chart")).find(qn("c:plotArea"))
    layout = plot_area.find(qn("c:layout"))
    if layout is None:
        layout = etree.Element(qn("c:layout"))
        plot_area.insert(0, layout)   # CT_PlotArea: layout comes first
    for child in list(layout):
        layout.remove(child)
    manual = etree.SubElement(layout, qn("c:manualLayout"))
    for tag, value in (("c:layoutTarget", "inner"), ("c:xMode", "edge"), ("c:yMode", "edge"),
                       ("c:x", x), ("c:y", y), ("c:w", w), ("c:h", h)):
        etree.SubElement(manual, qn(tag)).set("val", value if isinstance(value, str) else f"{value:.5f}")


def _move_data_label(label: Any, dx: float, dy: float) -> None:
    """Move a data label from its default place by (dx, dy), fractions of
    the chart frame's width and height (c:dLbl/c:layout/c:manualLayout in
    its default factor mode: the offset PowerPoint saves when a label is
    dragged). python-pptx has no setting for it."""
    from lxml import etree
    from pptx.oxml.ns import qn
    dlbl = label._get_or_add_dLbl()
    layout = dlbl.find(qn("c:layout"))
    if layout is None:
        layout = etree.Element(qn("c:layout"))
        dlbl.find(qn("c:idx")).addnext(layout)   # CT_DLbl: idx, then layout, then tx
    for child in list(layout):
        layout.remove(child)
    manual = etree.SubElement(layout, qn("c:manualLayout"))
    etree.SubElement(manual, qn("c:x")).set("val", f"{dx:.5f}")
    etree.SubElement(manual, qn("c:y")).set("val", f"{dy:.5f}")


_C15 = "http://schemas.microsoft.com/office/drawing/2012/chart"
_LEADER_LINES_EXT = "{CE6537A1-D6FC-4f65-9D91-7224C49458BB}"   # PowerPoint 2013's data-label extension


def _show_leader_lines(series: Any, colour: str) -> None:
    """Draw a leader line from each mark to its moved label, as PowerPoint
    2013 and later write it for any chart type (c15:showLeaderLines in the
    series' data labels' extension list; LibreOffice reads it too). The
    pie family's own c:showLeaderLines stays as python-pptx writes it.
    python-pptx has no setting for it."""
    from lxml import etree
    from pptx.oxml.ns import qn
    dlbls = series._element.get_or_add_dLbls()
    ext_list = dlbls.find(qn("c:extLst"))
    if ext_list is None:
        ext_list = etree.SubElement(dlbls, qn("c:extLst"))   # CT_DLbls: extLst comes last
    ext = etree.SubElement(ext_list, qn("c:ext"), nsmap={"c15": _C15})
    ext.set("uri", _LEADER_LINES_EXT)
    etree.SubElement(ext, f"{{{_C15}}}showLeaderLines").set("val", "1")
    leader_lines = etree.SubElement(ext, f"{{{_C15}}}leaderLines")
    line = etree.SubElement(etree.SubElement(leader_lines, qn("c:spPr")), qn("a:ln"), w="6350")
    etree.SubElement(etree.SubElement(line, qn("a:solidFill")), qn("a:srgbClr"), val=colour)


_NUMBER_PREFIXES = {"USD": "$", "EUR": "€", "GBP": "£"}


def _number_format(data: VisualData) -> str:
    """The Excel number format for a chart's values: '$2.4M', '$285K' for an
    amount (as the bar chart has always written them), '1,200' for a count."""
    if data.unit_kind != "amount":
        return "#,##0"
    prefix = _NUMBER_PREFIXES.get((data.currency or data.unit or "").upper(), "")
    return f'[>=1000000]{prefix}#,##0.0,,"M";[>=1000]{prefix}#,##0,"K";{prefix}#,##0'


# --- where a native chart's labels go ---------------------------------------------------
# PowerPoint writes a data label in the middle of its mark unless the label
# is moved (c:dLbl/c:layout/c:manualLayout, the offset PowerPoint itself
# saves when a label is dragged), and LibreOffice wraps a long label into a
# narrow column. A 100% bar's small parts and a donut's thin slices then
# carry labels that run over each other. So the parts' labels are laid out
# here the way the app's drawing lays them out -- a 100% bar's above and
# below the bar in turn, a donut's outside the ring on its own side -- and
# each label is moved there from its default place, as written, unwrapped.
# Sizes are in points, from the chart frame's top-left corner.
NATIVE_LABEL_SIZES = (11, 10, 9)   # pt, the largest that lets every label fit
NATIVE_LINE = 1.25                 # a label line's height, as a multiple of its size
NATIVE_LABEL_GAP = 6.0             # pt between two labels, and between a label and its mark
WHOLE_BAR_THICKNESS = {"thin": 22.0, "regular": 34.0, "thick": 50.0}   # pt
DONUT_HOLE = 45                    # percent of the diameter: a thicker ring than the 50% default
DONUT_LEADER = (10.0, 8.0)         # pt: out from the ring, then across to the label (ECharts' length, length2)
DONUT_MIN_DIAMETER = 96.0          # pt; a smaller ring is not drawn as a donut


@dataclass(frozen=True)
class PlacedLabel:
    """A data label's box (its centre, width and height) and the point it is
    moved from: the middle of its mark, where PowerPoint puts it."""
    x: float
    y: float
    width: float
    height: float
    anchor_x: float
    anchor_y: float


def label_box(text: str, size: float) -> tuple[float, float]:
    """The width and height a label of these lines takes at `size` pt, with
    the slack the app's drawing allows (layout.TEXT_SLACK)."""
    return block_width(text, size) * TEXT_SLACK, (text.count("\n") + 1) * size * NATIVE_LINE


def whole_bar_layout(labels: list[str], values: list[float], width: float, height: float,
                     thickness: float) -> tuple[float, float, list[PlacedLabel]] | None:
    """A native 100% bar in a frame `width` x `height` pt: the bar across the
    frame's middle, `thickness` tall, and each part's label above or below
    it in turn, as the app draws it (echarts._whole_bar_offsets): centred on
    its part, then moved apart and into the frame as little as the line
    needs (layout.place_on_line). A label must still stand over or under
    some of its own part, so it is never read as its neighbour's; when one
    line a side cannot manage that, two lines a side. Returns (font size,
    the bar's top, the labels in the parts' order), or None when no size
    fits."""
    total = math.fsum(values)
    spans, before = [], 0.0
    for value in values:
        start = before / total * width if total > 0 else 0.0
        before += value
        spans.append((start, before / total * width if total > 0 else width))
    centres = [(start + end) / 2 for start, end in spans]
    bar_top = (height - thickness) / 2
    middle = bar_top + thickness / 2
    for per_side in (1, 2):
        lines = 2 * per_side
        for size in NATIVE_LABEL_SIZES:
            boxes = [label_box(text, size) for text in labels]
            line_height = max((h for _, h in boxes), default=0.0)
            if per_side * (line_height + NATIVE_LABEL_GAP) > bar_top:
                continue
            placed: dict[int, PlacedLabel] = {}
            for line in range(lines):
                idx = [i for i in range(len(labels)) if i % lines == line]
                widths = [boxes[i][0] for i in idx]
                lefts = place_on_line([centres[i] for i in idx], widths, 0.0, width, NATIVE_LABEL_GAP)
                if lefts is None or any(left > spans[i][1] or left + w < spans[i][0] for i, left, w in zip(idx, lefts, widths)):
                    break
                # Even lines above the bar, odd ones below; the second line a side further out.
                reach = NATIVE_LABEL_GAP + line_height / 2 + (line // 2) * (line_height + NATIVE_LABEL_GAP)
                y = bar_top - reach if line % 2 == 0 else bar_top + thickness + reach
                for i, left, w in zip(idx, lefts, widths):
                    placed[i] = PlacedLabel(left + w / 2, y, w, boxes[i][1], centres[i], middle)
            else:
                return size, bar_top, [placed[i] for i in range(len(labels))]
    return None


def donut_layout(labels: list[str], values: list[float], width: float, height: float
                 ) -> tuple[float, float, float, list[PlacedLabel]] | None:
    """A native donut in a frame `width` x `height` pt, labelled as the app
    draws it: outside the ring on a leader line, on the ring's right for
    the slices on its right half and on its left for the others (slices
    clockwise from twelve o'clock). The labels of a side stand in a column
    clear of the ring, each at its slice's height, moved apart up and down
    as little as they need; the ring is as large as the frame's height and
    the two columns allow, and sits off centre when one side has more to
    say. Returns (font size, the ring's diameter, its centre's x, the
    labels in the slices' order), or None when no size leaves a ring of
    DONUT_MIN_DIAMETER."""
    total = math.fsum(values)
    angles, before = [], 0.0
    for value in values:
        angles.append(2 * math.pi * ((before + value / 2) / total if total > 0 else 0.0))
        before += value
    right = [math.sin(angle) >= 0 for angle in angles]
    out, across = DONUT_LEADER
    cy = height / 2
    for size in NATIVE_LABEL_SIZES:
        boxes = [label_box(text, size) for text in labels]
        left_w = max((w for (w, _), r in zip(boxes, right) if not r), default=0.0)
        right_w = max((w for (w, _), r in zip(boxes, right) if r), default=0.0)
        left_room = left_w + across if left_w else 0.0
        right_room = right_w + across if right_w else 0.0
        diameter = min(height - 2 * NATIVE_LABEL_GAP, width - left_room - right_room - 2 * NATIVE_LABEL_GAP)
        if diameter < DONUT_MIN_DIAMETER:
            continue
        radius = diameter / 2
        # The spare width split evenly either side of ring and columns.
        spare = width - left_room - right_room - diameter
        cx = spare / 2 + left_room + radius
        middle = radius * (1 + DONUT_HOLE / 100) / 2   # where PowerPoint writes a doughnut's label: mid-ring
        ys: dict[int, float] = {}
        for side in (True, False):
            members = sorted((i for i in range(len(labels)) if right[i] == side), key=lambda i: -math.cos(angles[i]))
            heights = [boxes[i][1] for i in members]
            tops = place_on_line([cy - (radius + out) * math.cos(angles[i]) for i in members], heights, 0.0, height, 2.0)
            if tops is None:
                break
            for i, top, h in zip(members, tops, heights):
                ys[i] = top + h / 2
        else:
            placed = [PlacedLabel(cx + radius + across + w / 2 if right[i] else cx - radius - across - w / 2, ys[i], w, h,
                                  cx + middle * math.sin(angle), cy - middle * math.cos(angle))
                      for i, (angle, (w, h)) in enumerate(zip(angles, boxes))]
            return size, diameter, cx, placed
    return None


PICTURE_TEXT_PT = 10.5   # the slide size of a picture's smallest text; the native charts' labels are 11 pt
PICTURE_SCALE = 3.0      # PNG pixels per design pixel: about 250 dpi on the slide


def picture_size(width_in: float, height_in: float) -> tuple[int, int]:
    """The design size a chart picture is drawn at for a slide region: the
    region's own shape, at a scale where the drawing's smallest text (12 px,
    rule 14) reads as PICTURE_TEXT_PT on the slide, and never wider than
    the app's design size. Drawn at 960 x 540 and shrunk into half a slide,
    a treemap's labels came out at 5.5 pt."""
    px_per_inch = 72 * MIN_FONT_PX / PICTURE_TEXT_PT
    design_w = min(DEFAULT_WIDTH, round(width_in * px_per_inch))
    return design_w, round(design_w * height_in / width_in)


def normalize_visual(visual: Any) -> dict | None:
    """Keep a well-formed visual, drop anything else. A deck must always
    export, so a malformed visual is discarded rather than raised on."""
    if isinstance(visual, dict) and visual.get("type") == CHART_VISUAL:
        return normalize_chart_visual(visual)
    if not isinstance(visual, dict) or visual.get("type") not in VISUAL_LIMITS:
        return None
    kind = visual["type"]
    items: list[dict[str, Any]] = []
    for item in (visual.get("items") or [])[: VISUAL_LIMITS[kind]]:
        if not isinstance(item, dict) or not str(item.get("label") or "").strip():
            continue
        clean: dict[str, Any] = {"label": str(item["label"])[:200], "evidence": _clean_refs(item.get("evidence"))}
        if kind == "figures":
            clean["value"] = str(item.get("value") or "")[:40]
        elif kind == "timeline":
            clean["date"] = str(item.get("date") or "")[:40]
        else:
            value = item.get("value")
            if isinstance(value, bool) or not isinstance(value, (int, float)):
                continue
            clean["value"] = value
        items.append(clean)
    if not items or (kind in {"timeline", "bar"} and len(items) < 2):
        return None
    out: dict[str, Any] = {"type": kind, "items": items}
    if kind == "bar":
        out["title"] = str(visual.get("title") or "")[:200]
        out["unit"] = str(visual.get("unit") or "")[:40]
    style = normalize_visual_style(kind, visual.get("style"), [i["label"] for i in items])
    if style:
        out["style"] = style
    return out


def _normalize_slides(slides: list[Any] | None) -> list[dict]:
    normalized = []
    # Stored and AI-written slides reach here unchecked, so a value of the
    # wrong type is dropped rather than raised on: a deck must always open.
    for slide in slides or []:
        if not isinstance(slide, dict):
            continue
        kind = slide.get("kind")
        bullets = slide.get("bullets")
        normalized.append({
            "kind": kind if isinstance(kind, str) and kind in SLIDE_KINDS else "content",
            "heading": str(slide.get("heading") or "Untitled slide")[:300],
            "bullets": [str(b)[:500] for b in (bullets if isinstance(bullets, list) else [])][:MAX_BULLETS_PER_SLIDE],
            "evidence": _clean_refs(slide.get("evidence")),
            "notes": str(slide.get("notes") or "")[:20000],
            "visual": normalize_visual(slide.get("visual")),
        })
    return normalized


def deck_citation_refs(slides: list[dict]) -> list[dict]:
    """Every ref a deck cites, in first-appearance order (slide evidence,
    then each visual item's evidence) -- the order sources are numbered in."""
    seen, refs = set(), []
    for slide in slides or []:
        candidates = list(slide.get("evidence") or [])
        visual = slide.get("visual") or {}
        for item in (visual.get("items") or []):
            candidates.extend(item.get("evidence") or [])
        if visual.get("type") == CHART_VISUAL:
            candidates.extend(chart_record_refs(visual))
        for ref in candidates:
            key = (ref.get("record_type"), ref.get("record_id"))
            if key[0] and key[1] and key not in seen:
                seen.add(key)
                refs.append({"record_type": key[0], "record_id": key[1]})
    return refs


def serialize_deck(row: Deck) -> dict:
    slides = _normalize_slides(row.slides_json)
    return {
        "id": row.id,
        "investigation_id": row.investigation_id,
        "title": row.title,
        "slides": slides,
        "slide_count": len(slides),
        "style": normalize_deck_style(row.style_json),
        "source_ai_candidate_id": row.source_ai_candidate_id,
        "created_by": row.created_by,
        "created_at": row.created_at,
        "updated_at": row.updated_at,
    }


def create_deck(db: Session, *, investigation_id: str, title: str, slides: list[dict] | None = None,
                 source_ai_candidate_id: str | None = None, created_by: str | None = None, style: dict | None = None) -> Deck:
    deck_style = normalize_deck_style(style)
    row = Deck(
        investigation_id=investigation_id, title=title.strip() or "Untitled deck",
        slides_json=recheck_charts(_normalize_slides(slides), resolved_theme(deck_style)), style_json=deck_style,
        source_ai_candidate_id=source_ai_candidate_id,
        created_by=created_by, created_at=utcnow_naive(),
    )
    db.add(row)
    db.commit()
    db.refresh(row)
    return row


def update_deck(db: Session, row: Deck, data: dict) -> Deck:
    if "title" in data and data["title"] is not None:
        row.title = data["title"].strip() or row.title
    slides_changed = "slides_json" in data and data["slides_json"] is not None
    style_changed = "style_json" in data and data["style_json"] is not None
    if style_changed:
        row.style_json = normalize_deck_style(data["style_json"])
    if slides_changed or style_changed:
        # A new look or new slides: every chart is drawn and checked again.
        slides = _normalize_slides(data["slides_json"]) if slides_changed else list(row.slides_json or [])
        row.slides_json = recheck_charts(slides, resolved_theme(row.style_json))
    db.commit()
    db.refresh(row)
    return row


def list_decks(db: Session, investigation_id: str) -> list[Deck]:
    query = select(Deck).where(Deck.investigation_id == investigation_id)
    return list(db.scalars(query.order_by(Deck.updated_at.desc())).all())


def delete_deck(db: Session, row: Deck) -> None:
    db.delete(row)
    db.commit()


# --- .pptx export -------------------------------------------------------------

# Dark Academia palette (frontend/app/globals.css), plus two paper tones for
# the content slides -- a projected/printed deck reads better on warm paper
# than on the app's near-black background, which is kept for the title and
# takeaway slides.
NIGHT = "141B1F"       # title / statement background (deeper than --card #1a2328)
NIGHT_LINE = "2B3439"
PARCHMENT = "ECE3D2"   # text on NIGHT
PARCHMENT_MUTED = "B5A88F"
BRASS = "B98D34"       # --primary
BRASS_DEEP = "8A6420"  # brass with enough contrast on paper
OXBLOOD = "672C1A"     # --secondary
FOREST = "3A5C52"      # --accent
PAPER = "F7F2E8"
PANEL = "EDE4D3"
RULE = "D9CDB6"
INK = "1F1B16"
TAUPE = "7A6950"

SERIF = "Georgia"
SANS = "Calibri"

SLIDE_W_IN, SLIDE_H_IN = 13.333, 7.5
MARGIN_IN = 0.75


_MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def display_date(value: str) -> str:
    """'2023-11-14' -> '14 Nov 2023', '2023-11' -> 'Nov 2023'; anything
    else (already human-written, or a year) is returned unchanged. Never
    adds precision the value didn't have."""
    parts = (value or "").strip().split("-")
    try:
        if len(parts) == 3 and len(parts[0]) == 4:
            return f"{int(parts[2])} {_MONTHS[int(parts[1]) - 1]} {parts[0]}"
        if len(parts) == 2 and len(parts[0]) == 4:
            return f"{_MONTHS[int(parts[1]) - 1]} {parts[0]}"
    except (ValueError, IndexError):
        pass
    return value


def _fit_size(texts: list[str], width_in: float, height_in: float, sizes: tuple[int, ...],
              *, char_w: float = 0.5, line_spacing: float = 1.2, gap_pt: float = 8.0) -> int:
    """Largest font size (pt) at which `texts`, one paragraph each, fit the
    box. Character-count estimate -- python-pptx can't measure text -- tuned
    on the side of fitting (char_w is a generous average glyph width)."""
    texts = [t for t in texts if t] or [""]
    for pt in sizes:
        per_line = max(8, int(width_in * 72 / (pt * char_w)))
        lines = sum(max(1, math.ceil(len(t) / per_line)) for t in texts)
        needed = lines * pt * line_spacing + gap_pt * (len(texts) - 1)
        if needed <= height_in * 72:
            return pt
    return sizes[-1]


def sort_bar_items(items: list[dict], order: str) -> list[dict]:
    """Reorder bars for display. Only the order changes: every item keeps
    its own label, value and citations."""
    if order == "desc":
        return sorted(items, key=lambda i: i["value"], reverse=True)
    if order == "asc":
        return sorted(items, key=lambda i: i["value"])
    return list(items)


class _DeckRenderer:
    def __init__(self, title: str, citation_details: dict[tuple[str, str], dict], citation_numbers: dict[tuple[str, str], int],
                 theme: Theme | None = None):
        from pptx import Presentation
        from pptx.util import Inches

        self.t = theme or resolved_theme(None)
        self.Inches = Inches
        self.title = title
        self.details = citation_details
        self.numbers = citation_numbers
        self.prs = Presentation()
        self.prs.slide_width = Inches(SLIDE_W_IN)
        self.prs.slide_height = Inches(SLIDE_H_IN)
        self.layout = self.prs.slide_layouts[5]  # "Title Only": keeps a real title placeholder for outline/accessibility
        self.page = 0

    # --- primitives ---------------------------------------------------------
    def _rgb(self, hex_color: str) -> Any:
        from pptx.dml.color import RGBColor
        return RGBColor.from_string(hex_color)

    def new_slide(self, background: str) -> Any:
        slide = self.prs.slides.add_slide(self.layout)
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = self._rgb(background)
        self.page += 1
        return slide

    def rect(self, slide: Any, x: float, y: float, w: float, h: float, color: str, *, rounded: bool = False) -> Any:
        from pptx.enum.shapes import MSO_SHAPE
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                       self.Inches(x), self.Inches(y), self.Inches(w), self.Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = self._rgb(color)
        shape.line.fill.background()
        shape.shadow.inherit = False
        if rounded:
            shape.adjustments[0] = 0.06
        return shape

    def oval(self, slide: Any, cx: float, cy: float, d: float, color: str, *, ring: str | None = None) -> Any:
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Pt
        shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, self.Inches(cx - d / 2), self.Inches(cy - d / 2), self.Inches(d), self.Inches(d))
        shape.fill.solid()
        shape.fill.fore_color.rgb = self._rgb(color)
        if ring:
            shape.line.color.rgb = self._rgb(ring)
            shape.line.width = Pt(2.5)
        else:
            shape.line.fill.background()
        shape.shadow.inherit = False
        return shape

    def _frame(self, box: Any, *, anchor: str = "top") -> Any:
        from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
        tf = box.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[anchor]
        return tf

    def _run(self, paragraph: Any, text: str, *, size: float, color: str, font: str | None = None, bold: bool = False, italic: bool = False, spacing: int | None = None) -> Any:
        from pptx.util import Pt
        run = paragraph.add_run()
        run.text = text
        run.font.name = font or self.t.body_font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = self._rgb(color)
        if spacing is not None:
            run.font._element.set("spc", str(spacing))  # letter-spacing, hundredths of a point
        return run

    def text(self, slide: Any, x: float, y: float, w: float, h: float, text: str, *, size: float, color: str, font: str | None = None, bold: bool = False, italic: bool = False,
             align: str = "left", anchor: str = "top", spacing: int | None = None, line_spacing: float | None = None) -> Any:
        from pptx.enum.text import PP_ALIGN
        box = slide.shapes.add_textbox(self.Inches(x), self.Inches(y), self.Inches(w), self.Inches(h))
        tf = self._frame(box, anchor=anchor)
        p = tf.paragraphs[0]
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
        if line_spacing:
            p.line_spacing = line_spacing
        self._run(p, text, size=size, color=color, font=font, bold=bold, italic=italic, spacing=spacing)
        return box

    def place_title(self, slide: Any, x: float, y: float, w: float, h: float, text: str, *, size: float, color: str, anchor: str = "top", line_spacing: float = 1.05) -> Any:
        """Use the layout's real title placeholder (so the slide has a title
        in PowerPoint's outline and for screen readers), restyled in place."""
        from pptx.enum.text import PP_ALIGN
        shape = slide.shapes.title
        shape.left, shape.top, shape.width, shape.height = self.Inches(x), self.Inches(y), self.Inches(w), self.Inches(h)
        tf = self._frame(shape, anchor=anchor)
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = line_spacing
        self._run(p, text, size=size, color=color, font=self.t.heading_font)
        return shape

    def bullets(self, slide: Any, x: float, y: float, w: float, h: float, items: list[str], *, color: str | None = None, bullet_color: str | None = None, sizes: tuple[int, ...] = (28, 26, 24, 22, 20, 18, 16, 15, 14)) -> Any:
        color = color or self.t.text
        bullet_color = bullet_color or self.t.accent
        from lxml import etree
        from pptx.oxml.ns import qn
        from pptx.util import Emu, Pt

        size = _fit_size(items, w - 0.35, h, sizes, gap_pt=16)
        box = slide.shapes.add_textbox(self.Inches(x), self.Inches(y), self.Inches(w), self.Inches(h))
        tf = self._frame(box)
        indent = Emu(int(self.Inches(0.32)))
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.line_spacing = 1.12
            p.space_after = Pt(16 if size >= 22 else 10)
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(indent))
            pPr.set("indent", str(-indent))
            for tag in ("a:buClr", "a:buSzPct", "a:buFont", "a:buChar"):
                existing = pPr.find(qn(tag))
                if existing is not None:
                    pPr.remove(existing)
            bu_clr = etree.SubElement(pPr, qn("a:buClr"))
            etree.SubElement(bu_clr, qn("a:srgbClr")).set("val", bullet_color)
            etree.SubElement(pPr, qn("a:buSzPct")).set("val", "80000")
            etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
            etree.SubElement(pPr, qn("a:buChar")).set("char", "■")
            self._run(p, item, size=size, color=color)
        return box

    def kicker(self, slide: Any, x: float, y: float, text: str, color: str) -> Any:
        return self.text(slide, x, y, 8.0, 0.3, text.upper(), size=11, color=color, bold=True, spacing=220)

    def notes(self, slide: Any, text: str) -> None:
        if text.strip():
            slide.notes_slide.notes_text_frame.text = text.strip()

    def footer(self, slide: Any, refs: list[dict], *, dark: bool) -> None:
        nums = sorted({self.numbers[(r["record_type"], r["record_id"])] for r in refs if (r["record_type"], r["record_id"]) in self.numbers})
        muted = self.t.cover_muted if dark else self.t.muted
        line = self.t.cover_line if dark else self.t.rule
        self.rect(slide, MARGIN_IN, 6.88, SLIDE_W_IN - 2 * MARGIN_IN, 0.012, line)
        if nums:
            self.text(slide, MARGIN_IN, 6.98, 8.0, 0.3, "Sources  " + "  ".join(f"[{n}]" for n in nums), size=10, color=muted)
        short_title = self.title if len(self.title) <= 70 else self.title[:69].rstrip() + "…"
        self.text(slide, SLIDE_W_IN - MARGIN_IN - 6.0, 6.98, 6.0, 0.3, f"{short_title}   ·   {self.page}", size=10, color=muted, align="right")

    # --- slide kinds --------------------------------------------------------
    def title_slide(self, heading: str, subtitle: str, source_count: int) -> None:
        slide = self.new_slide(self.t.cover_bg)
        self.rect(slide, 0, 0, 0.18, SLIDE_H_IN, self.t.accent)
        self.kicker(slide, 0.95, 1.35, "Investigation brief", self.t.cover_accent)
        size = _fit_size([heading], 11.2, 2.5, (44, 40, 36, 32, 28), char_w=0.58, line_spacing=1.12)
        self.place_title(slide, 0.95, 1.8, 11.2, 2.6, heading, size=size, color=self.t.cover_text)
        self.rect(slide, 0.95, 4.55, 1.1, 0.045, self.t.accent)
        if subtitle:
            sub_size = _fit_size([subtitle], 10.6, 1.3, (18, 16, 15, 14))
            self.text(slide, 0.95, 4.8, 10.6, 1.4, subtitle, size=sub_size, color=self.t.cover_muted, line_spacing=1.2)
        today = date.today()
        stamp = f"Prepared from {source_count} cited record{'s' if source_count != 1 else ''}  ·  {today.day} {today.strftime('%B %Y')}"
        self.text(slide, 0.95, 6.75, 10.0, 0.3, stamp, size=11, color=self.t.cover_muted)

    def statement_slide(self, slide_data: dict) -> None:
        slide = self.new_slide(self.t.cover_bg)
        text = " ".join(slide_data.get("bullets") or []) or slide_data.get("heading") or ""
        self.kicker(slide, 0.95, 1.2, slide_data.get("heading") or "The takeaway", self.t.cover_accent)
        size = _fit_size([text], 10.9, 4.6, (40, 36, 32, 30, 28, 26, 24, 22), char_w=0.58, line_spacing=1.3)
        lines = max(1, math.ceil(len(text) / max(8, int(10.9 * 72 / (size * 0.58)))))
        text_h = min(4.9, lines * size * 1.3 / 72 + 0.1)
        self.rect(slide, 0.95, 1.65, 0.06, text_h, self.t.accent)
        self.place_title(slide, 1.35, 1.6, 10.9, 5.0, text, size=size, color=self.t.cover_text, line_spacing=1.15)
        self.notes(slide, slide_data.get("notes") or "")
        self.footer(slide, slide_data.get("evidence") or [], dark=True)

    def _evidence_panel(self, slide: Any, x: float, y: float, w: float, h: float, refs: list[dict]) -> None:
        self.rect(slide, x, y, w, h, self.t.panel, rounded=True)
        self.kicker(slide, x + 0.3, y + 0.28, "From the record", self.t.alert)
        details = [d for d in (self.details.get((r["record_type"], r["record_id"])) for r in refs) if d]
        quoted = [d for d in details if d.get("quote")][:2]
        cursor = y + 0.75
        inner_w = w - 0.6
        bottom = y + h - 0.25
        if quoted:
            per = (bottom - cursor) / len(quoted)
            if len(quoted) == 1:
                per = min(per, 2.4)  # leave room to list the slide's other cited records
            for d in quoted:
                n = self.numbers.get((d["record_type"], d["record_id"]))
                quote = d["quote"]
                attribution = f"[{n}]  {d.get('locator') or d.get('label')}" + (f" — {d['source_title']}" if d.get("source_title") and d.get("locator") else "")
                text_w = inner_w - 0.32
                q_size = _fit_size([quote], text_w, per - 0.8, (18, 17, 16, 15, 14, 13, 12, 11), char_w=0.56, line_spacing=1.25)
                per_line = max(8, int(text_w * 72 / (q_size * 0.56)))
                max_lines = max(1, int((per - 0.8) * 72 / (q_size * 1.25)))
                if len(quote) > per_line * max_lines:
                    quote = quote[: max(20, per_line * max_lines - 1)].rstrip() + "…"
                quote_h = math.ceil(len(quote) / per_line) * q_size * 1.42 / 72
                self.text(slide, x + 0.3, cursor - 0.12, 0.4, 0.5, "“", size=34, color=self.t.accent, font=self.t.heading_font)
                self.text(slide, x + 0.62, cursor, text_w, quote_h + 0.05, quote, size=q_size, color=self.t.text, font=self.t.heading_font, italic=True, line_spacing=1.15)
                self.text(slide, x + 0.62, cursor + quote_h + 0.12, text_w, 0.5, attribution, size=11, color=self.t.muted, line_spacing=1.0)
                cursor += quote_h + 0.75
            rest = [d for d in details if d not in quoted]
            for d in rest[: max(0, int((bottom - cursor) / 0.62))]:
                n = self.numbers.get((d["record_type"], d["record_id"]))
                self.text(slide, x + 0.3, cursor, inner_w, 0.25, f"[{n}]  {d['type_label'].upper()}", size=9, color=self.t.muted, bold=True, spacing=120)
                label = d.get("label") or ""
                self.text(slide, x + 0.3, cursor + 0.24, inner_w, 0.35, label if len(label) <= 55 else label[:54] + "…", size=12, color=self.t.text)
                cursor += 0.62
        else:
            rows = details[:5]
            if not rows:
                self.text(slide, x + 0.3, cursor, inner_w, 0.6, "No cited records on this slide.", size=13, color=self.t.muted, italic=True)
                return
            per = min(0.95, (bottom - cursor) / len(rows))
            for d in rows:
                n = self.numbers.get((d["record_type"], d["record_id"]))
                self.text(slide, x + 0.3, cursor, inner_w, 0.25, f"[{n}]  {d['type_label'].upper()}", size=9, color=self.t.muted, bold=True, spacing=120)
                self.text(slide, x + 0.3, cursor + 0.25, inner_w, per - 0.3, d.get("label") or "", size=13, color=self.t.text, line_spacing=1.05)
                cursor += per

    def _cite_marks(self, refs: Any) -> str:
        nums = sorted({self.numbers[(r["record_type"], r["record_id"])] for r in refs or [] if (r["record_type"], r["record_id"]) in self.numbers})
        return " ".join(f"[{n}]" for n in nums)

    def _figures(self, slide: Any, x: float, y: float, w: float, h: float, items: list[dict], style: dict) -> None:
        """Cards / big numbers stacked in a column, or a row of stats."""
        opts = style["options"]
        value_color = {"alert": self.t.alert, "accent": self.t.accent_text, "text": self.t.text}[opts.get("value_color", "alert")]
        row = style["layout"].get("arrangement") == "row"
        gap = 0.25 if row else 0.2
        count = len(items)
        for i, item in enumerate(items):
            if row:
                cell_w, cell_h = (w - gap * (count - 1)) / count, h
                left, top = x + i * (cell_w + gap), y
            else:
                cell_w, cell_h = w, (h - gap * (count - 1)) / count
                left, top = x, y + i * (cell_h + gap)
            if opts.get("card_fill", True):
                self.rect(slide, left, top, cell_w, cell_h, self.t.panel, rounded=True)
            elif i > 0:
                if row:
                    self.rect(slide, left - gap / 2, top + 0.1, 0.012, cell_h - 0.2, self.t.rule)
                else:
                    self.rect(slide, left, top - gap / 2, cell_w, 0.012, self.t.rule)
            if opts.get("accent_edge", True):
                self.rect(slide, left, top + 0.18, 0.06, cell_h - 0.36, self.t.alert)
            pad = 0.35 if (opts.get("card_fill", True) or opts.get("accent_edge", True)) else 0.05
            value_size = _fit_size([item["value"]], cell_w - pad - 0.25, min(cell_h * 0.5, 1.0), (54, 48, 44, 40, 36, 32, 28, 24), char_w=0.55, line_spacing=1.0)
            self.text(slide, left + pad, top + 0.15, cell_w - pad - 0.25, cell_h * 0.5, item["value"], size=value_size, color=value_color, font=self.t.heading_font, anchor="bottom")
            marks = self._cite_marks(item["evidence"])
            label = item["label"] + ("  " + marks if marks else "")
            label_size = 15 if cell_h > 1.6 and cell_w > 3 else 13 if cell_h > 1.1 else 11
            self.text(slide, left + pad, top + 0.2 + cell_h * 0.5, cell_w - pad - 0.25, cell_h * 0.5 - 0.3, label, size=label_size, color=self.t.text, line_spacing=1.05)

    def _bar_chart(self, slide: Any, x: float, y: float, w: float, h: float, visual: dict, style: dict) -> Any:
        """A native (editable) PowerPoint bar chart. The value axis always
        includes zero -- no style can start a bar anywhere else. Returns the
        chart's frame on the slide."""
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
        from pptx.util import Pt

        opts, layout = style["options"], style["layout"]
        horizontal = layout.get("orientation", "horizontal") == "horizontal"
        unit = (visual.get("unit") or "").strip()
        currency = {"USD": "$", "US$": "$", "$": "$", "EUR": "€", "€": "€", "GBP": "£", "£": "£"}.get(unit.upper() if unit else "", None)
        if visual.get("title"):
            self.text(slide, x, y, w, 0.35, visual["title"] + (f"  ({unit})" if unit and not currency else ""), size=13, color=self.t.muted, bold=True)
            y, h = y + 0.4, h - 0.4
        items = sort_bar_items(visual["items"], opts.get("sort", "as_given"))
        if horizontal:
            items = list(reversed(items))  # bar charts draw bottom-up; keep the reading order top-down
        data = CategoryChartData()
        data.categories = [i["label"] for i in items]
        data.add_series(visual.get("title") or "Value", [i["value"] for i in items])
        chart_type = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
        frame = slide.shapes.add_chart(chart_type, self.Inches(x), self.Inches(y), self.Inches(w), self.Inches(h), data)
        chart = frame.chart
        chart.has_legend = False
        chart.has_title = False
        chart.font.name = self.t.body_font
        chart.font.size = Pt(12)
        chart.font.color.rgb = self._rgb(self.t.text)
        plot = chart.plots[0]
        plot.gap_width = {"thin": 150, "regular": 60, "thick": 25}[opts.get("thickness", "regular")]
        plot.vary_by_categories = False
        series = plot.series[0]
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = self._rgb(self.t.accent)
        highlight = set(opts.get("highlight") or [])
        if highlight:
            for idx, item in enumerate(items):
                point = series.points[idx]
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = self._rgb(self.t.accent if item["label"] in highlight else self.t.bar_muted)
        prefix = currency or ""
        fmt = f'[>=1000000]{prefix}#,##0.0,,"M";[>=1000]{prefix}#,##0,"K";{prefix}#,##0'
        value_labels = bool(opts.get("value_labels", True))
        plot.has_data_labels = value_labels
        if value_labels:
            labels = plot.data_labels
            labels.number_format = fmt
            labels.number_format_is_linked = False
            labels.position = XL_LABEL_POSITION.OUTSIDE_END
            labels.font.size = Pt(12)
            labels.font.bold = True
            labels.font.color.rgb = self._rgb(self.t.text)
        value_axis = chart.value_axis
        value_axis.minimum_scale = min(0, min(i["value"] for i in items))
        # With no value labels the axis is the only way to read a value, so it stays.
        show_axis = bool(opts.get("gridlines", False)) or not value_labels
        value_axis.visible = show_axis
        value_axis.has_major_gridlines = show_axis
        if show_axis:
            value_axis.major_gridlines.format.line.color.rgb = self._rgb(self.t.rule)
            value_axis.tick_labels.number_format = fmt
            value_axis.tick_labels.number_format_is_linked = False
            value_axis.tick_labels.font.size = Pt(11)
            value_axis.tick_labels.font.color.rgb = self._rgb(self.t.muted)
            value_axis.format.line.fill.background()
        category_axis = chart.category_axis
        if layout.get("axis_line", True):
            category_axis.format.line.color.rgb = self._rgb(self.t.rule)
        else:
            category_axis.format.line.fill.background()
        category_axis.tick_labels.font.size = Pt(12)
        category_axis.tick_labels.font.color.rgb = self._rgb(self.t.text)
        return frame

    # --- checked charts (app/visuals), one native drawing per kind -------------
    def _chart(self, slide: Any, x: float, y: float, w: float, h: float, visual: dict) -> str:
        """A resolved chart (app/ai/chart_check.py), drawn again from the rows
        the data service cooked and checked once more in this deck's theme
        before it goes in (a chart that now fails a hard rule is not drawn;
        its place says which rules failed). Every kind python-pptx can write
        is a native, editable PowerPoint chart (XL_CHART_TYPE: bar and
        column, line, area, 100% bar, doughnut, XY scatter, and a histogram as
        touching columns); a treemap, which PowerPoint charts cannot be made
        as from python-pptx, goes in as the picture from the render service,
        drawn and checked at the slide's own scale with its title and source
        line in it, or as native bars with a line saying why when there is no
        picture (ADR-0005, line 2). Under every native chart: the source line
        (rule 15); on every chart frame: the alt text (rule 13). Returns the
        chart's description for the speaker notes."""
        checked = checked_chart(visual, self.t)
        if checked.problem or checked.built is None:
            self.text(slide, x, y, w, 1.6, checked.problem or "Chart not available.", size=13, color=self.t.alert, italic=True, line_spacing=1.1)
            return ""
        built = checked.built
        kind = built.kind
        missing: str | None = None
        if kind not in NATIVE_KINDS:
            picture = self._chart_picture(slide, x, y, w, h, visual)
            if not isinstance(picture, str):
                _set_alt_text(picture, built.alt_text)
                return built.description
            missing = picture
        # The chart's own title and subtitle (the finding, then the unit), as the app draws them.
        self.text(slide, x, y, w, 0.32, checked.spec.title, size=13, color=self.t.muted, bold=True)
        titles = built.option.get("title") or []
        subtitle = str(titles[0].get("subtext") or "") if titles and isinstance(titles[0], dict) else ""
        if subtitle:
            self.text(slide, x, y + 0.3, w, 0.28, subtitle, size=11, color=self.t.muted)
        top = y + (0.62 if subtitle else 0.4)
        plot_h = h - (top - y) - 0.36
        if kind in (ChartKind.bar, ChartKind.column):
            frame = self._native_bars(slide, x, top, w, plot_h, checked)
        elif kind in (ChartKind.line, ChartKind.area):
            frame = self._native_line(slide, x, top, w, plot_h, checked)
        elif kind == ChartKind.stacked_bar_100:
            frame = self._native_whole_bar(slide, x, top, w, plot_h, checked)
        elif kind == ChartKind.donut:
            frame = self._native_donut(slide, x, top, w, plot_h, checked)
        elif kind == ChartKind.histogram:
            frame = self._native_histogram(slide, x, top, w, plot_h, checked)
        elif kind == ChartKind.scatter:
            frame = self._native_scatter(slide, x, top, w, plot_h, checked)
        else:
            frame = self._parts_as_bars(slide, x, top, w, plot_h, checked, missing or "")
        _set_alt_text(frame, built.alt_text)
        self.text(slide, x, y + h - 0.32, w, 0.3, built.source_line, size=10, color=self.t.muted)
        return built.description

    def _new_chart(self, slide: Any, chart_type: Any, x: float, y: float, w: float, h: float, data: Any) -> Any:
        """A native chart frame in the deck's type: no legend (every part is
        labelled directly, rule 12), no chart title (the slide gives it)."""
        from pptx.util import Pt
        frame = slide.shapes.add_chart(chart_type, self.Inches(x), self.Inches(y), self.Inches(w), self.Inches(h), data)
        chart = frame.chart
        chart.has_legend = False
        chart.has_title = False
        chart.font.name = self.t.body_font
        chart.font.size = Pt(12)
        chart.font.color.rgb = self._rgb(self.t.text)
        return frame

    def _point_label(self, point: Any, text: str, *, position: Any = None, colour: str | None = None, bold: bool = False,
                     size: float = 11, moved: tuple[float, float] | None = None) -> None:
        """Write `text` beside one mark as its own data label (the value, its
        record count and the single-record mark, as the app's chart does),
        its lines as written: never wrapped into a narrower column, as
        LibreOffice otherwise does (a:bodyPr wrap="none"). `moved` shifts it
        from its default place, in fractions of the chart frame
        (_move_data_label)."""
        from pptx.util import Pt
        label = point.data_label
        frame = label.text_frame
        lines = text.split("\n")
        frame.text = lines[0]
        for line in lines[1:]:
            frame.add_paragraph().text = line
        for paragraph in frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(size)
                run.font.bold = bold
                run.font.color.rgb = self._rgb(colour or self.t.text)
        frame.word_wrap = False
        if position is not None:
            label.position = position
        if moved is not None:
            _move_data_label(label, *moved)

    def _style_value_axis(self, axis: Any, *, minimum: float | None, number_format: str, gridlines: bool, visible: bool = True) -> None:
        from pptx.util import Pt
        if minimum is not None:
            axis.minimum_scale = minimum
        axis.visible = visible
        axis.has_major_gridlines = gridlines
        if gridlines:
            axis.major_gridlines.format.line.color.rgb = self._rgb(self.t.rule)
        axis.tick_labels.number_format = number_format
        axis.tick_labels.number_format_is_linked = False
        axis.tick_labels.font.size = Pt(11)
        axis.tick_labels.font.color.rgb = self._rgb(self.t.muted)
        axis.format.line.fill.background()

    def _style_category_axis(self, axis: Any, *, dark_zero_line: bool = True) -> None:
        from pptx.util import Pt
        if dark_zero_line:
            # The dark zero line (BBC style), as in the app's drawing.
            axis.format.line.color.rgb = self._rgb(self.t.text)
            axis.format.line.width = Pt(1.5)
        else:
            axis.format.line.fill.background()
        axis.tick_labels.font.size = Pt(12)
        axis.tick_labels.font.color.rgb = self._rgb(self.t.text)

    def _native_bars(self, slide: Any, x: float, y: float, w: float, h: float, checked: CheckedChart, *, horizontal: bool | None = None) -> Any:
        """bar and column: the native bar chart pass 2 drew, labels carrying
        the record count and the single-record mark (rules 16 and 20).
        `horizontal` overrides the kind (the treemap's fallback ranks its parts)."""
        built, data = checked.built, checked.data
        assert built is not None
        if horizontal is None:
            horizontal = built.kind == ChartKind.bar
        # Columns stand side by side, so a column's count goes on a line of its own.
        decorated = chart_row_labels(data, built.rows, separator=" · " if horizontal else "\n")
        style = {
            "preset": "ranked",
            "layout": {"orientation": "horizontal" if horizontal else "vertical", "axis_line": built.draw.axis_line},
            "options": {"sort": "as_given", "value_labels": True, "gridlines": built.draw.gridlines and not horizontal, "thickness": built.draw.thickness,
                        "highlight": [decorated[label] for label in built.highlighted_labels if label in decorated]},
        }
        bar_visual = {"type": "bar", "unit": data.unit, "items": chart_bar_items(built.rows, decorated)}
        return self._bar_chart(slide, x, y, w, h, bar_visual, style)

    def _native_line(self, slide: Any, x: float, y: float, w: float, h: float, checked: CheckedChart) -> Any:
        """line and area: periods in order on a zero baseline; a point that
        rests on one record drawn hollow (rule 20 by shape) and its period
        marked; the last value written at the end of the line. Each period's
        label carries its record count on a line of its own (the app's chart
        has it in the point's tooltip), so a value written on a point is the
        value alone."""
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_MARKER_STYLE
        from pptx.util import Pt

        built, data = checked.built, checked.data
        assert built is not None
        rows, draw = built.rows, built.draw
        area = built.kind == ChartKind.area
        colour = mark_colour(self.t)
        decorated = chart_row_labels(data, rows, separator="\n")
        chart_data = CategoryChartData()
        chart_data.categories = [decorated[row.label] for row in rows]
        chart_data.add_series(checked.spec.title, [row.value for row in rows])
        frame = self._new_chart(slide, XL_CHART_TYPE.AREA if area else XL_CHART_TYPE.LINE_MARKERS, x, y, w, h, chart_data)
        chart = frame.chart
        series = chart.plots[0].series[0]
        if area:
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = self._rgb(colour)
            _set_fill_alpha(series.format.fill, 30)   # python-pptx has no transparency setting (its issue #62)
            series.format.line.color.rgb = self._rgb(colour)
            series.format.line.width = Pt(2)
        else:
            series.smooth = draw.smooth
            series.format.line.color.rgb = self._rgb(colour)
            series.format.line.width = Pt({"thin": 1.5, "regular": 2.5, "thick": 4.0}[draw.line_width])
            series.marker.style = XL_MARKER_STYLE.CIRCLE if draw.markers else XL_MARKER_STYLE.NONE
            series.marker.size = 7
            series.marker.format.fill.solid()
            series.marker.format.fill.fore_color.rgb = self._rgb(colour)
            series.marker.format.line.color.rgb = self._rgb(colour)
            for index, row in enumerate(rows):
                if row.record_count == 1:
                    point = series.points[index]
                    point.marker.style = XL_MARKER_STYLE.CIRCLE
                    point.marker.size = 8
                    point.marker.format.fill.background()   # hollow: rests on one record
                    point.marker.format.line.color.rgb = self._rgb(colour)
        highlighted = set(built.highlighted_labels)
        for index, row in enumerate(rows):
            if draw.value_labels or row.label in highlighted or index == len(rows) - 1:
                self._point_label(series.points[index], format_value(row.value, data), position=None if area else XL_LABEL_POSITION.ABOVE,
                                  bold=row.label in highlighted)
        self._style_value_axis(chart.value_axis, minimum=min(0.0, min((row.value for row in rows), default=0.0)),
                               number_format=_number_format(data), gridlines=draw.gridlines)
        self._style_category_axis(chart.category_axis)
        return frame

    def _parts_as_bars(self, slide: Any, x: float, y: float, w: float, h: float, checked: CheckedChart, note: str) -> Any:
        """Parts of a whole that cannot be drawn as their own kind here, as
        native bars ranked largest first, under a line that says why."""
        lines = max(1, math.ceil(block_width(note, 10) * TEXT_SLACK / (w * 72)))
        note_h = lines * 10 * 1.2 / 72 + 0.12
        self.text(slide, x, y, w, note_h, note, size=10, color=self.t.alert, italic=True, line_spacing=1.05)
        return self._native_bars(slide, x, y + note_h, w, h - note_h, checked, horizontal=True)

    def _part_labels(self, checked: CheckedChart) -> list[str]:
        built = checked.built
        assert built is not None
        return [f"{row.label}\n{part_value_text(row, share, checked.data)}" for row, share in zip(built.rows, built.shares)]

    def _native_whole_bar(self, slide: Any, x: float, y: float, w: float, h: float, checked: CheckedChart) -> Any:
        """stacked_bar_100: PowerPoint's 100% stacked bar, one bar of one
        series per part, across the middle of its space. Each part's name and
        share go above and below the bar in turn, as the app draws them, moved
        apart so small neighbours do not run over each other
        (whole_bar_layout); when no size lets them, the parts are shown as
        ranked bars with a line saying so."""
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE

        built = checked.built
        assert built is not None
        labels = self._part_labels(checked)
        width, height = w * 72, h * 72
        thickness = WHOLE_BAR_THICKNESS[built.draw.thickness]
        layout = whole_bar_layout(labels, [row.value for row in built.rows], width, height, thickness)
        if layout is None:
            return self._parts_as_bars(slide, x, y, w, h, checked, "The parts' names do not fit beside a 100% bar this size; "
                                                                    "the same parts are shown as bars.")
        size, bar_top, placed = layout
        colours = part_colours(built.rows, self.t, set(built.highlighted_labels))
        chart_data = CategoryChartData()
        chart_data.categories = [""]
        for row in built.rows:
            chart_data.add_series(row.label, [row.value])
        frame = self._new_chart(slide, XL_CHART_TYPE.BAR_STACKED_100, x, y, w, h, chart_data)
        chart = frame.chart
        plot = chart.plots[0]
        plot.gap_width = 0   # the bar fills the band pinned for it
        _pin_plot_area(chart, 0.0, bar_top / height, 1.0, thickness / height)
        for series, row, text, label, colour in zip(plot.series, built.rows, labels, placed, colours):
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = self._rgb(colour)
            series.format.line.color.rgb = self._rgb(self.t.page)
            self._point_label(series.points[0], text, bold=row.label in built.highlighted_labels, size=size,
                              moved=((label.x - label.anchor_x) / width, (label.y - label.anchor_y) / height))
            _show_leader_lines(series, self.t.muted)
        chart.value_axis.visible = False
        chart.value_axis.has_major_gridlines = False
        chart.category_axis.visible = False
        return frame

    def _native_donut(self, slide: Any, x: float, y: float, w: float, h: float, checked: CheckedChart) -> Any:
        """donut: PowerPoint's doughnut (never the exploded one), slices from
        twelve o'clock, each labelled outside the ring on its own side as the
        app draws it (donut_layout; PowerPoint and LibreOffice draw the
        leader line), the whole written in the hole. When the ring would be
        too small to leave room for the labels, the parts are shown as
        ranked bars with a line saying so."""
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE

        built, data = checked.built, checked.data
        assert built is not None
        labels = self._part_labels(checked)
        width, height = w * 72, h * 72
        layout = donut_layout(labels, [row.value for row in built.rows], width, height)
        if layout is None:
            return self._parts_as_bars(slide, x, y, w, h, checked, "The slices' names do not fit around a donut this size; "
                                                                    "the same parts are shown as bars.")
        size, diameter, centre, placed = layout
        colours = part_colours(built.rows, self.t, set(built.highlighted_labels))
        chart_data = CategoryChartData()
        chart_data.categories = [row.label for row in built.rows]
        chart_data.add_series(checked.spec.title, [row.value for row in built.rows])
        frame = self._new_chart(slide, XL_CHART_TYPE.DOUGHNUT, x, y, w, h, chart_data)
        chart = frame.chart
        _pin_plot_area(chart, (centre - diameter / 2) / width, (height - diameter) / 2 / height, diameter / width, diameter / height)
        plot = chart.plots[0]
        plot.vary_by_categories = True
        _set_hole_size(plot, DONUT_HOLE)
        series = plot.series[0]
        for index, (row, text, label, colour) in enumerate(zip(built.rows, labels, placed, colours)):
            point = series.points[index]
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = self._rgb(colour)
            point.format.line.color.rgb = self._rgb(self.t.page)
            self._point_label(point, text, bold=row.label in built.highlighted_labels, size=size,
                              moved=((label.x - label.anchor_x) / width, (label.y - label.anchor_y) / height))
        _show_leader_lines(series, self.t.muted)
        if built.draw.centre_total and built.total is not None:
            hole = diameter * DONUT_HOLE / 100 / 72   # inches
            self.text(slide, x + centre / 72 - hole / 2, y + h / 2 - 0.2, hole, 0.4, format_value(built.total, data),
                      size=max(11, min(18, round(hole * 72 / 4.5))), color=self.t.text, font=self.t.heading_font, bold=True,
                      align="center", anchor="middle")
        return frame

    def _native_histogram(self, slide: Any, x: float, y: float, w: float, h: float, checked: CheckedChart) -> Any:
        """histogram: touching columns (gap width 0), one per bin, the count
        written on each, a hairline in the page colour between them."""
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
        from pptx.util import Pt

        built, data = checked.built, checked.data
        assert built is not None
        chart_data = CategoryChartData()
        chart_data.categories = [row.label for row in built.rows]
        chart_data.add_series(checked.spec.title, [row.value for row in built.rows])
        frame = self._new_chart(slide, XL_CHART_TYPE.COLUMN_CLUSTERED, x, y, w, h, chart_data)
        chart = frame.chart
        plot = chart.plots[0]
        plot.gap_width = 0
        plot.vary_by_categories = False
        series = plot.series[0]
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = self._rgb(self.t.accent)
        series.format.line.color.rgb = self._rgb(self.t.page)
        series.format.line.width = Pt(0.75)
        for index, row in enumerate(built.rows):
            text = f"{int(row.value)}" + (f" {SINGLE_RECORD_MARK}" if row.record_count == 1 else "")
            self._point_label(series.points[index], text, position=XL_LABEL_POSITION.OUTSIDE_END, bold=True)
        self._style_value_axis(chart.value_axis, minimum=0, number_format="0", gridlines=built.draw.gridlines)
        self._style_category_axis(chart.category_axis)
        axis = chart.category_axis
        axis.has_title = True
        axis.axis_title.text_frame.text = (built.option.get("xAxis") or {}).get("name") or "Amount"
        return frame

    def _native_scatter(self, slide: Any, x: float, y: float, w: float, h: float, checked: CheckedChart) -> Any:
        """scatter: PowerPoint's XY scatter with markers only, the two
        numbers' names on the axes, the named points labelled, a point that
        rests on one record hollow."""
        from pptx.chart.data import XyChartData
        from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_MARKER_STYLE

        built, data = checked.built, checked.data
        assert built is not None
        colour = mark_colour(self.t)
        chart_data = XyChartData()
        series_data = chart_data.add_series(checked.spec.title)
        for row in built.rows:
            series_data.add_data_point(row.x or 0.0, row.value)
        frame = self._new_chart(slide, XL_CHART_TYPE.XY_SCATTER, x, y, w, h, chart_data)
        chart = frame.chart
        series = chart.plots[0].series[0]
        series.marker.style = XL_MARKER_STYLE.CIRCLE
        series.marker.size = 9
        series.marker.format.fill.solid()
        series.marker.format.fill.fore_color.rgb = self._rgb(colour)
        series.marker.format.line.color.rgb = self._rgb(colour)
        named = set(built.labelled_points) | set(built.highlighted_labels)
        for index, row in enumerate(built.rows):
            point = series.points[index]
            if row.record_count == 1:
                point.marker.style = XL_MARKER_STYLE.CIRCLE
                point.marker.size = 9
                point.marker.format.fill.background()   # hollow: rests on one record
                point.marker.format.line.color.rgb = self._rgb(colour)
            if row.label in named:
                text = row.label + (f" {SINGLE_RECORD_MARK}" if row.record_count == 1 else "")
                self._point_label(point, text, position=XL_LABEL_POSITION.RIGHT, bold=row.label in built.highlighted_labels)
        names = data.measure_labels
        x_axis, y_axis = chart.category_axis, chart.value_axis   # for an XY chart: the axis across, then the axis up
        self._style_value_axis(x_axis, minimum=min(0.0, min((row.x or 0.0 for row in built.rows), default=0.0)),
                               number_format="0" if data.x_unit_kind == "count" else "General", gridlines=False)
        self._style_value_axis(y_axis, minimum=min(0.0, min((row.value for row in built.rows), default=0.0)),
                               number_format=_number_format(data), gridlines=built.draw.gridlines)
        for axis, title in ((x_axis, names.get("x", "x")), (y_axis, names.get("y", "y"))):
            axis.has_title = True
            axis.axis_title.text_frame.text = title
        x_axis.format.line.color.rgb = self._rgb(self.t.text)
        return frame

    def _chart_picture(self, slide: Any, x: float, y: float, w: float, h: float, visual: dict) -> Any:
        """A kind PowerPoint cannot hold natively (the treemap): the picture
        the render service draws, filling the chart's whole space with the
        chart's title, unit and source line in it, as the app draws them.
        It is drawn for the slide -- at a design size where its smallest
        text reads at PICTURE_TEXT_PT (picture_size) -- and checked again at
        that size, since a part that can name itself at 960 x 540 may not
        at half a slide. Returns the picture, or why there is none: the
        chart fails at this size, or the render service cannot be reached."""
        from io import BytesIO

        from app.visuals.render_client import RenderUnavailable, render_png

        design_w, design_h = picture_size(w, h)
        sized = checked_chart(visual, self.t, width=design_w, height=design_h)
        label = KIND_LABELS.get(sized.spec.kind, sized.spec.kind.value)
        if sized.built is None or sized.report is None or not sized.report.ok:
            first = sized.report.errors[0] if sized.report and sized.report.errors else None
            reason = f"rule {first.rule}: {first.message.split(';')[0]}" if first else (sized.problem or "it cannot be drawn")
            return f"{label} not drawn at this size ({reason}). The same parts are shown as bars."
        try:
            png = render_png(sized.built.option, design_w, design_h, scale=PICTURE_SCALE)
        except RenderUnavailable as exc:
            return f"{label} image not available ({exc}); the same parts are shown as bars."
        return slide.shapes.add_picture(BytesIO(png), self.Inches(x), self.Inches(y), self.Inches(w), self.Inches(h))

    def _timeline(self, slide: Any, x: float, y: float, w: float, h: float, items: list[dict], style: dict) -> None:
        opts, layout = style["options"], style["layout"]
        n = len(items)
        ring = opts.get("marker", "ring") == "ring"
        if layout.get("direction") == "vertical":
            row_h = min(0.7, h / n)
            if opts.get("show_line", True):
                self.rect(slide, x + 2.25, y + row_h / 2, 0.025, row_h * (n - 1), self.t.rule)
            size = 16 if row_h >= 0.6 else 14 if row_h >= 0.5 else 12
            for i, item in enumerate(items):
                cy = y + i * row_h + row_h / 2
                self.text(slide, x, cy - 0.2, 2.0, 0.4, display_date(item.get("date") or ""), size=size, color=self.t.accent_text, bold=True, align="right", anchor="middle")
                if ring:
                    self.oval(slide, x + 2.26, cy, 0.22, self.t.page, ring=self.t.accent)
                self.oval(slide, x + 2.26, cy, 0.1 if ring else 0.16, self.t.accent)
                marks = self._cite_marks(item["evidence"])
                label = item["label"] + (" " + marks if marks else "")
                self.text(slide, x + 2.6, cy - row_h / 2 + 0.05, w - 2.6, row_h - 0.1, label, size=size, color=self.t.text, anchor="middle", line_spacing=1.05)
            return
        compact = bool(layout.get("compact"))
        line_y = y + (0.55 if compact else 0.75)
        if opts.get("show_line", True):
            self.rect(slide, x, line_y - 0.015, w, 0.03, self.t.rule)
        step = w / n
        label_size = (14 if n <= 3 else 12 if n <= 5 else 11) if compact else (16 if n <= 3 else 14 if n <= 5 else 12)
        for i, item in enumerate(items):
            cx = x + step * (i + 0.5)
            if ring:
                self.oval(slide, cx, line_y, 0.16 if compact else 0.22, self.t.page, ring=self.t.accent)
            self.oval(slide, cx, line_y, (0.07 if compact else 0.1) if ring else (0.12 if compact else 0.16), self.t.accent)
            self.text(slide, cx - step / 2 + 0.05, y, step - 0.1, 0.4, display_date(item.get("date") or ""), size=label_size + 2, color=self.t.accent_text, bold=True, align="center")
            marks = self._cite_marks(item["evidence"])
            label = item["label"] + (" " + marks if marks else "")
            self.text(slide, cx - step / 2 + 0.08, line_y + 0.25, step - 0.16, h - (0.85 if compact else 1.05), label, size=label_size, color=self.t.text, align="center", line_spacing=1.05)

    def content_slide(self, slide_data: dict, index: int, total: int) -> None:
        slide = self.new_slide(self.t.page)
        visual = slide_data.get("visual")
        self.kicker(slide, MARGIN_IN, 0.42, f"{index:02d} / {total:02d}", self.t.accent_text)
        heading = slide_data.get("heading") or "Untitled slide"
        h_size = _fit_size([heading], 11.8, 0.9, (34, 30, 28, 26, 24), char_w=0.58, line_spacing=1.1)
        self.place_title(slide, MARGIN_IN, 0.72, 11.8, 0.95, heading, size=h_size, color=self.t.text, anchor="middle")
        self.rect(slide, MARGIN_IN, 1.78, 0.9, 0.045, self.t.accent)
        bullets = slide_data.get("bullets") or []
        refs = list(slide_data.get("evidence") or [])
        vis: dict = visual if isinstance(visual, dict) else {}
        kind = vis.get("type")
        # A chart visual has no preset layer: its look comes from the theme and its own spec.
        style: dict = resolved_visual_style(vis) if kind in VISUAL_PRESETS else {}
        chart_description = ""
        if kind == "timeline":
            layout = style["layout"]
            if layout.get("direction") == "vertical":
                if bullets:
                    self.bullets(slide, MARGIN_IN, 2.1, 5.0, 4.55, bullets, sizes=(22, 20, 18, 16, 15, 14))
                self._timeline(slide, 6.0 if bullets else MARGIN_IN, 2.1, 6.55 if bullets else 11.8, 4.55, vis["items"], style)
            else:
                compact = bool(layout.get("compact"))
                band = 1.75 if compact else 2.45
                if bullets:
                    self.bullets(slide, MARGIN_IN, 2.05, 11.8, 4.6 - band, bullets, sizes=(22, 20, 18, 16, 15, 14))
                self._timeline(slide, MARGIN_IN + 0.1, (6.7 - band) if bullets else 2.6, 11.6, band if bullets else 3.6, vis["items"], style)
        elif kind == "bar":
            if bullets:
                self.bullets(slide, MARGIN_IN, 2.1, 5.3, 4.55, bullets, sizes=(22, 20, 18, 16, 15, 14))
            self._bar_chart(slide, 6.45 if bullets else MARGIN_IN, 2.0, 6.1 if bullets else 11.8, 4.7, vis, style)
        elif kind == CHART_VISUAL:
            if bullets:
                self.bullets(slide, MARGIN_IN, 2.1, 5.3, 4.55, bullets, sizes=(22, 20, 18, 16, 15, 14))
            chart_description = self._chart(slide, 6.45 if bullets else MARGIN_IN, 2.0, 6.1 if bullets else 11.8, 4.7, vis)
        elif kind == "figures":
            if style["layout"].get("arrangement") == "row":
                if bullets:
                    self.bullets(slide, MARGIN_IN, 2.05, 11.8, 2.3, bullets, sizes=(22, 20, 18, 16, 15, 14))
                self._figures(slide, MARGIN_IN, 4.6 if bullets else 2.3, 11.8, 2.0 if bullets else 3.0, vis["items"], style)
            else:
                if bullets:
                    self.bullets(slide, MARGIN_IN, 2.1, 6.9, 4.55, bullets)
                self._figures(slide, 8.15 if bullets else MARGIN_IN, 2.05, 4.45 if bullets else 11.8, 4.6, vis["items"], style)
        else:
            if bullets:
                self.bullets(slide, MARGIN_IN, 2.1, 6.9, 4.55, bullets)
            self._evidence_panel(slide, 8.15, 2.05, 4.45, 4.6, refs)
        for item in (visual or {}).get("items") or []:
            refs.extend(item.get("evidence") or [])
        notes = slide_data.get("notes") or ""
        if kind == CHART_VISUAL:
            refs.extend(chart_record_refs(vis))
            if chart_description:
                # Rule 13's longer description travels with the deck: every
                # value with the records behind it, readable without the chart.
                notes = f"{notes}\n\nChart: {chart_description}".strip()
        self.notes(slide, notes)
        self.footer(slide, refs, dark=False)

    def questions_slide(self, slide_data: dict) -> None:
        slide = self.new_slide(self.t.page)
        self.kicker(slide, MARGIN_IN, 0.42, "Open questions", self.t.alert)
        self.place_title(slide, MARGIN_IN, 0.72, 11.8, 0.95, slide_data.get("heading") or "What the record doesn't yet show", size=32, color=self.t.text, anchor="middle")
        self.rect(slide, MARGIN_IN, 1.78, 0.9, 0.045, self.t.alert)
        questions = (slide_data.get("bullets") or [])[:6]
        if questions:
            gap = 0.12
            per = min(1.15, (4.8 - gap * (len(questions) - 1)) / len(questions))
            longest = max(questions, key=len)
            size = _fit_size([longest], 10.6, per - 0.1, (24, 22, 20, 18, 16, 15, 14), gap_pt=0)
            for i, question in enumerate(questions):
                top = 2.1 + i * (per + gap)
                self.text(slide, MARGIN_IN, top, 0.9, per, f"{i + 1:02d}", size=26, color=self.t.accent_text, font=self.t.heading_font)
                self.text(slide, MARGIN_IN + 1.0, top + 0.05, 10.8, per, question, size=size, color=self.t.text, line_spacing=1.1)
        self.notes(slide, slide_data.get("notes") or "")
        self.footer(slide, slide_data.get("evidence") or [], dark=False)

    def sources_slides(self, refs_in_order: list[tuple[str, str]]) -> None:
        per_slide = 10
        chunks = [refs_in_order[i:i + per_slide] for i in range(0, len(refs_in_order), per_slide)]
        for c, chunk in enumerate(chunks):
            slide = self.new_slide(self.t.page)
            self.kicker(slide, MARGIN_IN, 0.42, "Sources" + (f"  ·  {c + 1} of {len(chunks)}" if len(chunks) > 1 else ""), self.t.accent_text)
            self.place_title(slide, MARGIN_IN, 0.72, 11.8, 0.95, "Where every statement comes from", size=30, color=self.t.text, anchor="middle")
            self.rect(slide, MARGIN_IN, 1.78, 0.9, 0.045, self.t.accent)
            col_w, row_h = 5.75, 0.9
            for i, key in enumerate(chunk):
                col, row = divmod(i, per_slide // 2)
                x = MARGIN_IN + col * (col_w + 0.3)
                y = 2.05 + row * row_h
                d = self.details.get(key) or {"type_label": key[0], "label": key[1], "record_type": key[0], "record_id": key[1]}
                n = self.numbers[key]
                self.text(slide, x, y, 0.6, 0.4, f"[{n}]", size=13, color=self.t.accent_text, bold=True)
                self.text(slide, x + 0.6, y, col_w - 0.6, 0.25, d["type_label"].upper() + ("  ·  NO LONGER AVAILABLE" if d.get("missing") else ""), size=9, color=self.t.muted, bold=True, spacing=120)
                second = d.get("source_title") if d.get("source_title") and d.get("source_title") != d.get("label") else (d.get("detail") or "")
                main = d.get("label") or key[1]
                self.text(slide, x + 0.6, y + 0.22, col_w - 0.6, 0.32, main if len(main) <= 80 else main[:79] + "…", size=13, color=self.t.text, bold=True)
                if second:
                    self.text(slide, x + 0.6, y + 0.5, col_w - 0.6, 0.3, second if len(second) <= 90 else second[:89] + "…", size=11, color=self.t.muted, italic=True)
            self.footer(slide, [], dark=False)


def render_deck_pptx(row: Deck, citation_details: list[dict] | None = None) -> Any:
    """Render a deck to a python-pptx Presentation (16:9).

    `citation_details` is describe_record_refs()'s output for this deck's
    refs (see deck_citation_refs); without it the sources slide falls back
    to record type + id, so an export never fails for want of labels.
    """
    slides = _normalize_slides(row.slides_json)
    refs = deck_citation_refs(slides)
    numbers = {(r["record_type"], r["record_id"]): i + 1 for i, r in enumerate(refs)}
    details = {(d["record_type"], d["record_id"]): d for d in (citation_details or [])}
    renderer = _DeckRenderer(row.title, details, numbers, resolved_theme(getattr(row, "style_json", None)))

    title_slide = next((s for s in slides if s["kind"] == "title"), None)
    renderer.title_slide(
        title_slide["heading"] if title_slide else row.title,
        " ".join(title_slide["bullets"]) if title_slide else "",
        len(refs),
    )
    if title_slide and title_slide.get("notes"):
        renderer.notes(renderer.prs.slides[0], title_slide["notes"])

    content = [s for s in slides if s["kind"] == "content"]
    content_index = 0
    for slide_data in slides:
        if slide_data["kind"] == "statement":
            renderer.statement_slide(slide_data)
        elif slide_data["kind"] == "questions":
            renderer.questions_slide(slide_data)
        elif slide_data["kind"] == "content":
            content_index += 1
            renderer.content_slide(slide_data, content_index, len(content))
    if refs:
        renderer.sources_slides(list(numbers.keys()))
    return renderer.prs
