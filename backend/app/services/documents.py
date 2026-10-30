from __future__ import annotations
import hashlib, re, uuid
from pathlib import Path
from typing import Any, Iterable, NamedTuple
from sqlalchemy import select
from sqlalchemy.orm import Session
from app.core.time import utcnow_naive
from app.models.domain import Document, DocumentChunk, DocumentCorpusSync, ExtractionCandidate, Source, Evidence, Claim, ClaimEvidenceLink, Entity, Statement, Investigation, RelationshipEdge, TableMapping
from app.services.ftm import make_ftm_entity
from app.services.identifier_proposals import accept_identifier
from app.services.security import contained_path, resolve_storage_root
from app.services.resolution import BASIS_SAME, candidate_entities, entity_names
from app.services.pdf_ocr import extract_pdf_chunks
from app.services.image_reader import IMAGE_EXTENSIONS, extract_image_chunks
from app.services.container_readers import CONTAINER_EXTENSIONS, read_container
from app.services.document_reader import read_chunks
from app.services.anchors import anchor_document_quotes
from app.services.entity_signals import clean_names, norm_name, read_name, read_passage, split_title
from app.core.config import settings
from app.db.locking import lock_investigation_transaction
from app.services.photos import PHOTO_EXTENSIONS
from app.services.transcription import MEDIA_EXTENSIONS
from app.jobs.corpus import request_corpus_sync

# Pictures are documents too (app/services/image_reader.py): kept as pictures, their text read by OCR.
# Emails, mailboxes, .zip archives and spreadsheets are read by app/services/container_readers.py.
# Audio and video are transcribed on this computer (app/services/transcription.py, ADR-0020).
SUPPORTED_EXTENSIONS = {'.txt', '.md', '.csv', '.html', '.htm', '.pdf', '.docx'} | set(IMAGE_EXTENSIONS) | set(CONTAINER_EXTENSIONS) | set(MEDIA_EXTENSIONS) | set(PHOTO_EXTENSIONS)
# Capitals and letters of the Latin alphabet, accented ones too (Société Vérandine, Brühlmark & Söhne).
_NAME_UPPER = "A-ZÀ-ÖØ-Þ"
_NAME_LETTER = "A-Za-zÀ-ÖØ-öø-ÿĀ-ž"
NAME_RE = re.compile(rf"\b([{_NAME_UPPER}][{_NAME_LETTER}'’.-]+(?:\s+(?:[{_NAME_UPPER}][{_NAME_LETTER}'’.-]+|of|the|and|&)){{1,5}})\b")
ORG_SUFFIX = re.compile(r"\b(Inc\.?|LLC|Ltd\.?|Corporation|Corp\.?|Company|Co\.?|Department|Commission|Authority|University|Foundation|Association|Committee|Council|Bank)\b", re.I)
SENTENCE_SPLIT_RE = re.compile(r'(?<=[.!?])\s+')
CLAIM_TRIGGER_RE = re.compile(r'\b(is|are|was|were|has|have|had|will|did|does|announced|reported|paid|owns|served|worked|received|filed|approved|voted)\b', re.I)
CLAIM_MIN_CHARS, CLAIM_MAX_CHARS = 35, 1200


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _clean(text: str) -> str:
    return re.sub(r"[ \t]+", " ", (text or '').replace('\x00','')).strip()


def extract_chunks(path: Path) -> tuple[list[dict], list[str]]:
    """Extract locator-preserving text chunks from a supported document.

    Returns ``(chunks, issues)``. ``issues`` are non-fatal problems worth
    surfacing to the reporter (a blank/unreadable page, a DOCX whose content
    lives only in tables/images we don't read, an empty file) -- the document
    still reaches ``extraction_status='complete'`` with these attached, unlike
    a raised exception, which is reserved for a total extraction failure.
    """
    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported document type: {ext or 'unknown'}")
    rows: list[dict] = []
    issues: list[str] = []
    if ext in MEDIA_EXTENSIONS:
        # The reading job transcribes a recording itself: its passages come with segments a chunk list cannot carry.
        raise ValueError('A recording is read by transcription (app/services/transcription.py), not as text')
    if ext in PHOTO_EXTENSIONS:
        # A HEIC or RAW photo is read through its viewable copy, which the reading job makes and keeps.
        raise ValueError('A HEIC or RAW photo is read through its viewable copy (app/services/photos.py), not as text')
    if ext in CONTAINER_EXTENSIONS:
        return read_container(path)
    if ext in IMAGE_EXTENSIONS:
        return extract_image_chunks(path, enable_ocr=settings.enable_pdf_ocr, language=settings.pdf_ocr_language, dpi=settings.pdf_ocr_dpi)
    if ext == '.pdf':
        result = extract_pdf_chunks(
            path,
            enable_ocr=settings.enable_pdf_ocr,
            language=settings.pdf_ocr_language,
            dpi=settings.pdf_ocr_dpi,
            min_native_chars=settings.pdf_ocr_min_native_chars,
        )
        rows.extend(result.chunks)
        if not rows and result.ocr_error:
            raise RuntimeError(f'PDF contained no extractable text and OCR failed: {result.ocr_error}')
        if result.blank_pages:
            total_pages = result.native_pages + result.ocr_pages + result.blank_pages
            issues.append(
                f'{result.blank_pages} of {total_pages} page(s) produced no extractable text '
                '(blank, image quality too low for OCR, or an unsupported page type).'
            )
        if result.ocr_error and rows:
            # Some pages extracted fine; OCR only degraded, not a total failure.
            issues.append(f'OCR was degraded on at least one page: {result.ocr_error}')
    elif ext == '.docx':
        from docx import Document as DocxDocument
        doc = DocxDocument(str(path))
        for i, para in enumerate(doc.paragraphs, 1):
            text = _clean(para.text)
            if text:
                rows.append({'locator': f'paragraph {i}', 'page_number': None, 'text': text})
        if not rows:
            has_tables = len(doc.tables) > 0
            issues.append(
                'No paragraph text found in this DOCX.'
                + (' It contains tables, which this app does not extract yet -- '
                   'review the original file directly for that content.' if has_tables
                   else ' If it contains only images or text boxes, that content is not extracted.')
            )
    elif ext in ('.html', '.htm'):
        from app.services.html_reader import html_blocks
        for i, block in enumerate(html_blocks(path.read_text(encoding='utf-8', errors='replace')), 1):
            rows.append({'locator': f'paragraph {i}', 'page_number': None, 'text': block})
        if not rows:
            issues.append('This page has no readable text outside its markup.')
    else:
        text = path.read_text(encoding='utf-8', errors='replace')
        # Paragraphs retain useful deterministic line ranges when possible.
        lines = text.splitlines()
        start = None; buf=[]
        for idx, line in enumerate(lines + [''], 1):
            clean = _clean(line)
            if clean:
                if start is None: start = idx
                buf.append(clean)
            elif buf:
                end = idx - 1
                loc = f'line {start}' if start == end else f'lines {start}-{end}'
                rows.append({'locator': loc, 'page_number': None, 'text': ' '.join(buf)})
                start=None; buf=[]
        if not rows:
            issues.append('This file has no readable text content.')
    return rows, issues


def _marker(text: str) -> re.Pattern[str]:
    """A marker's words with any run of space between them, as a reader may break lines anywhere."""
    return re.compile(r'\s+'.join(re.escape(word) for word in text.split()))


def own_part(chunks: list[dict], part: dict | None) -> tuple[list[tuple[int, int] | None], str | None]:
    """Where in each chunk a document's own text is, when its file holds more than it.

    A gazette page carries the notices printed beside the one a source is
    (a Federal Register PDF of one notice starts with the end of the notice
    before it). ``part`` names the first words of the document's own text and
    the last (``{"from": ..., "to": ...}``, the last included); each chunk gets
    the (start, end) of its share of that text, or None for none of it. Without
    ``part`` every chunk is wholly the document's. When the first words are
    not found, nothing is cut and the second value says so."""
    if not part:
        return [(0, len(c.get('text') or '')) for c in chunks], None
    start, stop = _marker(part['from']), _marker(part['to'])
    out: list[tuple[int, int] | None] = []
    inside, ended = False, False
    for c in chunks:
        text = c.get('text') or ''
        a = 0
        if not inside and not ended:
            found = start.search(text)
            if found is None:
                out.append(None)
                continue
            inside, a = True, found.start()
        if ended:
            out.append(None)
            continue
        end = stop.search(text, a)
        if end is not None:
            inside, ended = False, True
            out.append((a, end.end()))
        else:
            out.append((a, len(text)))
    if not any(out):
        return [(0, len(c.get('text') or '')) for c in chunks], (
            f"The start of this document's own text (\"{part['from']}\") was not found, so the whole file was read for proposals.")
    return out, None


class _Part(NamedTuple):
    text: str
    locator: str | None


def _part_payloads(chunk: DocumentChunk, share: tuple[int, int], **kwargs: Any) -> Iterable[tuple[str, dict, float]]:
    """``_candidate_payloads`` over one chunk's share of the document's own text, with spans counted in the whole chunk."""
    a, b = share
    if (a, b) == (0, len(chunk.text)):
        yield from _candidate_payloads(chunk, **kwargs)
        return
    part = chunk.text[a:b]
    lead = len(part) - len(part.lstrip())
    for typ, payload, conf in _candidate_payloads(_Part(part, chunk.locator), **kwargs):  # type: ignore[arg-type]
        span = payload.get('span')
        if span:
            payload = {**payload, 'span': {'start_char': span['start_char'] + a + lead, 'end_char': span['end_char'] + a + lead}}
        yield typ, payload, conf


class FoundSoFar:
    """What the reader will propose from the pages read so far, counted as a read goes (Case home's Bring in: "4 names
    and 2 claims found so far"): the same proposals, one per name and per claim sentence, from the document's own
    part. The proposals themselves are made once the whole file is read."""

    def __init__(self, part: dict | None = None) -> None:
        self.part = part
        self.pages: list[dict] = []
        self.seen: set[tuple[str, str]] = set()
        self.counts = {'entity': 0, 'claim': 0}

    def add(self, page_no: int, text: str) -> None:
        self.pages.append({'locator': f'page {page_no}', 'text': text})
        share = own_part(self.pages, self.part)[0][-1]
        if share is None:
            return
        part = text[share[0]:share[1]]
        for typ, payload, _conf in _candidate_payloads(_Part(part, None), include_entities=settings.enable_local_entity_suggestions):  # type: ignore[arg-type]
            key = _once_key(typ, payload)
            if typ in self.counts and key is not None and key not in self.seen:
                self.seen.add(key)
                self.counts[typ] += 1

    def words(self) -> str:
        """"4 names and 2 claims found so far", or "" before anything is found."""
        parts = [f"{n} {one if n == 1 else many}" for n, one, many in
                 ((self.counts['entity'], 'name', 'names'), (self.counts['claim'], 'claim', 'claims')) if n]
        return f"{' and '.join(parts)} found so far" if parts else ""


def _candidate_payloads(chunk: DocumentChunk, *, include_entities: bool = True,
                        existing: dict[str, tuple[str, str]] | None = None) -> Iterable[tuple[str, dict, float]]:
    """Proposals from one chunk. Each carries why its rule fired and how sure
    the text lets it be (app/services/entity_signals.py); ``existing`` maps the
    investigation's record names to (id, schema)."""
    text = chunk.text.strip()
    if not text: return
    passage_confidence, passage_rationale = read_passage(text)
    yield ('evidence', {'quote': text, 'locator': chunk.locator, 'span': {'start_char': 0, 'end_char': len(text)}, 'rationale': passage_rationale}, passage_confidence)
    # Claims are suggestions only: sentence-like assertions long enough to be useful.
    # Each payload records *why* the rule fired (#190) so review can show it later
    # even if these rules change; see app/services/extraction_explanations.py.
    for sentence in SENTENCE_SPLIT_RE.split(text):
        sentence = sentence.strip()
        trigger = CLAIM_TRIGGER_RE.search(sentence) if CLAIM_MIN_CHARS <= len(sentence) <= CLAIM_MAX_CHARS else None
        if trigger:
            start = text.find(sentence)
            yield ('claim', {'text': sentence, 'span': {'start_char': start, 'end_char': start + len(sentence)},
                             'rationale': {'rule': 'sentence_assertion', 'trigger': trigger.group(1), 'length': len(sentence)}}, 0.55)
    if not include_entities:
        return
    seen=set()
    for match in NAME_RE.finditer(text):
        for found, offset, stop in clean_names(match.group(1)):
            caption, title_chars = split_title(found)
            start = match.start(1) + offset + title_chars
            if len(caption) < 5 or caption in seen:
                continue
            seen.add(caption)
            reading = read_name(caption, text[:start], existing)
            payload = {'caption': caption, 'suggested_schema': reading.schema,
                       'span': {'start_char': start, 'end_char': match.start(1) + stop}, 'rationale': reading.rationale()}
            if reading.choices:
                payload['schema_choices'] = reading.choices
            yield ('entity', payload, reading.confidence)


def store_document(db: Session, *, investigation_id: str, title: str, filename: str, mime_type: str | None, data: bytes, storage_dir: Path,
                   parent_document_id: str | None = None, source_metadata: dict | None = None) -> Document:
    """Persist one uploaded document without racing investigation deletion.

    Saves the file, its Source and a Document row with
    ``extraction_status='queued'``; reading the text is ``extract_document``.

    PostgreSQL writers take the same transaction advisory lock used by destructive
    restore/delete operations, then re-check that the investigation still exists.
    File bytes are first written under ``.upload-staging`` and are only promoted to
    the canonical document path after the database rows have been prepared. If the
    database commit fails, a newly promoted file is removed so a failed upload does
    not leave reporter data orphaned on disk.

    ``parent_document_id`` is the archive or email a file was found in, and
    ``source_metadata`` is added to its Source's metadata (openaleph_children.py).
    """
    storage_dir = resolve_storage_root(storage_dir)
    storage_dir.mkdir(parents=True, exist_ok=True)

    # Coordinate with delete/restore before checking mutable investigation state.
    lock_investigation_transaction(db, investigation_id)
    if db.get(Investigation, investigation_id) is None:
        raise ValueError('Investigation not found')

    digest = sha256_bytes(data)
    safe_name = re.sub(r'[^A-Za-z0-9._-]+', '_', Path(filename).name) or 'document'
    stored = contained_path(storage_dir, f'{digest[:16]}_{safe_name}')
    staging_root = contained_path(storage_dir, '.upload-staging')
    staging_root.mkdir(parents=True, exist_ok=True)
    staged = contained_path(staging_root, f'{uuid.uuid4().hex}_{safe_name}')
    staged.write_bytes(data)
    created_final = False

    try:
        source = Source(investigation_id=investigation_id, title=title or filename, source_type='document', metadata_json={**(source_metadata or {}), 'filename': filename, 'mime_type': mime_type, 'sha256': digest})
        db.add(source); db.flush()
        doc = Document(investigation_id=investigation_id, source_id=source.id, filename=filename, mime_type=mime_type, sha256=digest, storage_path=str(stored), extraction_status='queued', parent_document_id=parent_document_id)
        db.add(doc); db.flush()

        # Promote the staged bytes only after relational state has been prepared.
        if not stored.exists():
            staged.replace(stored)
            created_final = True
        else:
            staged.unlink(missing_ok=True)

        db.commit(); db.refresh(doc)
        return doc
    except Exception:
        db.rollback()
        staged.unlink(missing_ok=True)
        if created_final:
            stored.unlink(missing_ok=True)
        raise
    finally:
        try:
            staging_root.rmdir()
        except OSError:
            pass


MAX_MENTIONS_KEPT = 500  # places kept per proposal; the count goes on past it


def _once_key(typ: str, payload: dict) -> tuple[str, str] | None:
    """What makes two proposals in one document the same thing to decide: the
    same name (as the name-matcher compares names), or the same claim sentence.
    Passages are each their own page and are not folded."""
    if typ == 'entity' and payload.get('caption'):
        return ('entity', norm_name(str(payload['caption'])))
    if typ == 'claim' and payload.get('text'):
        return ('claim', re.sub(r"\s+", " ", str(payload['text'])).strip().lower())
    return None


def extract_document(db: Session, document_id: str, *, text: tuple[list[dict], list[str]] | None = None,
                     find_names: bool = True) -> Document | None:
    """Turn a stored document into locator-preserving chunks and proposals.

    ``text`` is text already read elsewhere (OpenAleph; see
    app/services/openaleph_reader.py) as ``(chunks, issues)``; a chunk may
    carry a ``provenance`` dict, which every proposal made from it records.
    Without it the Workbench reads the file itself.

    ``find_names`` runs the Workbench's own pattern name-finder (when
    ENABLE_LOCAL_ENTITY_SUGGESTIONS is on). The caller turns it off when
    OpenAleph's ftm-analyze has analysed the document: the pattern finder is
    the fallback only (ADR-0004, decision 5; ADR-0006). It never runs on a
    table that already has a column mapping.

    Returns None when the document (or its investigation) was deleted before
    this ran. Extraction failures remain a valid committed document state: the
    reporter keeps the source file and sees ``extraction_status='failed'``
    with the reason, and nothing half-read is kept. A document that already
    has chunks is never read twice.
    """
    doc = db.get(Document, document_id)
    if doc is None:
        return None
    lock_investigation_transaction(db, doc.investigation_id)
    db.refresh(doc)
    if db.get(Investigation, doc.investigation_id) is None:
        return None
    if db.scalar(select(DocumentChunk.id).where(DocumentChunk.document_id == doc.id).limit(1)) is not None:
        raise ValueError('Document has already been read')
    doc.extraction_status = 'extracting'
    doc.extraction_error = None
    # The investigation's records by name (caption, other names and aliases): a name it already
    # holds is proposed as that record's kind. A caption wins over another record's alias.
    held = db.scalars(select(Entity).where(Entity.investigation_id == doc.investigation_id, Entity.merged_into_entity_id.is_(None))).all()
    existing: dict[str, tuple[str, str]] = {norm_name(record.caption): (record.id, record.schema) for record in held}
    for record in held:
        for name in entity_names(record)[1:]:
            existing.setdefault(norm_name(name), (record.id, record.schema))
    # A table (CSV, TSV, XLSX) gets its names, links and payments from its columns, through a mapping
    # (app/services/table_mappings.py), whether it is mapped yet or not: its cells read as running text
    # make junk names, claims and one whole-block passage. Its text is still kept, so it can be searched.
    from app.services.table_mappings import is_table

    table = is_table(doc)
    mapped = db.scalar(select(TableMapping.id).where(TableMapping.document_id == doc.id).limit(1)) is not None
    find_names = find_names and not mapped and not table
    try:
        chunks, issues = text if text is not None else read_chunks(Path(doc.storage_path))
        source = db.get(Source, doc.source_id)
        shares, not_found = own_part(chunks, ((source.metadata_json or {}) if source else {}).get('part'))
        if not_found:
            issues = [*(issues or []), not_found]
        # One proposal per name (and per claim sentence) in a document, with every place it is
        # mentioned: a name on 146 pages is one thing to decide, mentioned 146 times.
        once: dict[tuple[str, str], ExtractionCandidate] = {}
        for ordinal, raw in enumerate(chunks, 1):
            item = dict(raw)
            provenance = item.pop('provenance', None)
            chunk = DocumentChunk(document_id=doc.id, ordinal=ordinal, **item)
            db.add(chunk); db.flush()
            share = shares[ordinal - 1]
            if table or share is None:
                continue
            for typ, payload, conf in _part_payloads(chunk, share, include_entities=find_names and settings.enable_local_entity_suggestions,
                                                     existing=existing):
                if provenance:
                    payload = {**payload, 'provenance': provenance}
                key = _once_key(typ, payload)
                mention = {'chunk_id': chunk.id, 'locator': chunk.locator, 'span': payload.get('span')}
                if key is not None and key in once:
                    first = once[key]
                    mentions = [*first.payload.get('mentions', []), mention][:MAX_MENTIONS_KEPT]
                    first.payload = {**first.payload, 'mentions': mentions, 'mention_count': int(first.payload.get('mention_count', 1)) + 1}
                    first.confidence = max(first.confidence, conf)
                    continue
                if key is not None:
                    payload = {**payload, 'mentions': [mention], 'mention_count': 1}
                row = ExtractionCandidate(investigation_id=doc.investigation_id, document_id=doc.id, chunk_id=chunk.id, candidate_type=typ, payload=payload, confidence=conf)
                db.add(row)
                if key is not None:
                    once[key] = row
        doc.extraction_status = 'complete'
        doc.extraction_issues = issues or None
        db.flush()
        # Quotes saved from this document before it was read, or kept from an earlier reading, find their place.
        anchor_document_quotes(db, doc)
        db.commit()
    except Exception as exc:
        # Drop any chunks/proposals added before the failure, then record why.
        db.rollback()
        failed = db.get(Document, document_id)
        if failed is None:
            return None
        failed.extraction_status = 'failed'
        failed.extraction_error = str(exc) or type(exc).__name__
        db.commit()
        doc = failed
    db.refresh(doc)
    return doc


class OpenAlephRef(NamedTuple):
    """How OpenAleph holds a document: its id there, and the schema it gave
    the document once read (an Email, a Package...)."""
    entity_id: str
    schema: str | None


def openaleph_document_ids(db: Session, document_ids: Iterable[str]) -> dict[str, OpenAlephRef]:
    """OpenAleph's id (and schema) for each of these documents that OpenAleph holds, in one query."""
    ids = list(document_ids)
    if not ids:
        return {}
    rows = db.scalars(select(DocumentCorpusSync).where(
        DocumentCorpusSync.document_id.in_(ids), DocumentCorpusSync.provider == 'openaleph',
        DocumentCorpusSync.status == 'synced', DocumentCorpusSync.provider_record_id.is_not(None))).all()
    refs: dict[str, OpenAlephRef] = {}
    for row in rows:
        schema = (row.response_json or {}).get('schema') if isinstance(row.response_json, dict) else None
        refs[row.document_id] = OpenAlephRef(str(row.provider_record_id), schema if isinstance(schema, str) else None)
    return refs


def serialize_document(db: Session, doc: Document, *, openaleph_ids: dict[str, OpenAlephRef] | None = None) -> dict:
    """``openaleph_document_id`` is OpenAleph's id for the document when
    OpenAleph holds it (ADR-0006: names, similar documents and threads come
    from there), and ``openaleph_schema`` what OpenAleph read it as.
    ``parent_document_id`` is the archive or email it was found in. A list
    passes ``openaleph_ids``, looked up for all its rows at once
    (openaleph_document_ids)."""
    chunks = db.scalars(select(DocumentChunk).where(DocumentChunk.document_id == doc.id).order_by(DocumentChunk.ordinal)).all()
    candidates = db.scalars(select(ExtractionCandidate).where(ExtractionCandidate.document_id == doc.id).order_by(ExtractionCandidate.created_at)).all()
    counts: dict[str, int] = {}
    for c in candidates: counts[c.candidate_type] = counts.get(c.candidate_type,0)+1
    known = openaleph_ids if openaleph_ids is not None else openaleph_document_ids(db, [doc.id])
    ref = known.get(doc.id)
    source = db.get(Source, doc.source_id)
    return {'id':doc.id,'investigation_id':doc.investigation_id,'source_id':doc.source_id,'filename':doc.filename,
            'title':source.title if source is not None else None,'mime_type':doc.mime_type,'sha256':doc.sha256,'extraction_status':doc.extraction_status,'extraction_error':doc.extraction_error,'extraction_issues':doc.extraction_issues or [],'created_at':doc.created_at,'chunk_count':len(chunks),'candidate_counts':counts,
            'openaleph_document_id':ref.entity_id if ref else None, 'openaleph_schema':ref.schema if ref else None,
            'parent_document_id':doc.parent_document_id}


def serialize_chunk(chunk: DocumentChunk) -> dict:
    """One extracted chunk as the API sends it. Rows used to go out as ORM
    objects, relying on FastAPI's encoder to read their columns; with a
    declared response model that is fragile, so the shape is explicit."""
    return {'id': chunk.id, 'document_id': chunk.document_id, 'locator': chunk.locator, 'page_number': chunk.page_number,
            'ordinal': chunk.ordinal, 'text': chunk.text, 'created_at': chunk.created_at, 'ocr_legibility': chunk.ocr_legibility}


def serialize_candidate(candidate: ExtractionCandidate) -> dict:
    """One extraction proposal as the API sends it, the same shape in every
    reply that carries one (list, review, staged relationship)."""
    return {
        'id': candidate.id, 'investigation_id': candidate.investigation_id, 'document_id': candidate.document_id,
        'chunk_id': candidate.chunk_id, 'candidate_type': candidate.candidate_type, 'payload': candidate.payload or {},
        'confidence': candidate.confidence, 'review_status': candidate.review_status,
        'accepted_record_type': candidate.accepted_record_type, 'accepted_record_id': candidate.accepted_record_id,
        'reviewer_note': candidate.reviewer_note, 'created_at': candidate.created_at, 'reviewed_at': candidate.reviewed_at,
    }


def serialize_extraction_lineage(db: Session, candidate: ExtractionCandidate) -> dict:
    """Return the complete read-only document -> chunk -> proposal -> review -> accepted-record trail."""
    doc = db.get(Document, candidate.document_id)
    chunk = db.get(DocumentChunk, candidate.chunk_id) if candidate.chunk_id else None
    source = db.get(Source, doc.source_id) if doc else None
    locator = chunk.locator if chunk else None
    provenance = (candidate.payload or {}).get('provenance') if isinstance(candidate.payload, dict) else None
    recorded_method: object = None
    if isinstance(provenance, dict):
        # A name OpenAleph found was located in a chunk read earlier; the
        # chunk's own reading method is recorded beside it (openaleph_analysis.py).
        recorded_method = (provenance.get('chunk_extraction_method') if provenance.get('extraction_method') == 'openaleph_ftm_analyze'
                           else provenance.get('extraction_method'))
    extraction_method = (recorded_method if isinstance(recorded_method, str) and recorded_method else None) or ('ocr' if locator and '(OCR)' in locator else ('native_pdf' if doc and (doc.mime_type == 'application/pdf' or doc.filename.lower().endswith('.pdf')) else 'native_text'))
    return {
        'candidate': {
            'id': candidate.id, 'investigation_id': candidate.investigation_id,
            'document_id': candidate.document_id, 'chunk_id': candidate.chunk_id,
            'candidate_type': candidate.candidate_type, 'payload': candidate.payload or {},
            'confidence': candidate.confidence, 'review_status': candidate.review_status,
            'accepted_record_type': candidate.accepted_record_type,
            'accepted_record_id': candidate.accepted_record_id,
            'reviewer_note': candidate.reviewer_note, 'created_at': candidate.created_at,
            'reviewed_at': candidate.reviewed_at,
            'decided_by_id': candidate.decided_by_id, 'decided_by_name': candidate.decided_by_name,
        },
        'document': ({
            'id': doc.id, 'investigation_id': doc.investigation_id, 'source_id': doc.source_id,
            'filename': doc.filename, 'mime_type': doc.mime_type, 'sha256': doc.sha256,
            'extraction_status': doc.extraction_status, 'created_at': doc.created_at,
        } if doc else None),
        'source': ({
            'id': source.id, 'investigation_id': source.investigation_id, 'title': source.title,
            'url': source.url, 'source_type': source.source_type, 'metadata_json': source.metadata_json or {},
        } if source else None),
        'chunk': ({
            'id': chunk.id, 'document_id': chunk.document_id, 'ordinal': chunk.ordinal,
            'locator': chunk.locator, 'page_number': chunk.page_number, 'text': chunk.text,
            'extraction_method': extraction_method,
            'candidate_span': ((candidate.payload or {}).get('span') if isinstance(candidate.payload, dict) else None),
        } if chunk else None),
    }


def propose_relationship_from_extracted_claim(db: Session, *, claim_id: str, schema: str, source_entity_id: str, target_entity_id: str, properties: dict | None = None, evidence_id: str | None = None, note: str | None = None) -> ExtractionCandidate:
    """Stage a reporter-selected relationship without promoting it to canonical FtM data.

    The claim must itself have come from an accepted extraction candidate. Endpoints and
    relationship semantics are explicitly chosen by the reporter; no NLP/entity matcher is
    allowed to create a canonical edge. Promotion happens only in review_candidate().
    """
    claim = db.get(Claim, claim_id)
    if claim is None:
        raise ValueError('Claim not found')
    lineage = db.scalar(select(ExtractionCandidate).where(
        ExtractionCandidate.accepted_record_type == 'claim',
        ExtractionCandidate.accepted_record_id == claim_id,
        ExtractionCandidate.review_status == 'accepted',
    ).order_by(ExtractionCandidate.reviewed_at.desc()))
    if lineage is None:
        raise ValueError('Relationship proposals require a claim accepted from document extraction')
    source_entity = db.get(Entity, source_entity_id); target_entity = db.get(Entity, target_entity_id)
    if source_entity is None or target_entity is None:
        raise ValueError('Relationship endpoints must reference existing entities')
    if source_entity.investigation_id != claim.investigation_id or target_entity.investigation_id != claim.investigation_id:
        raise ValueError('Relationship endpoints must belong to the claim investigation')
    if source_entity_id == target_entity_id:
        raise ValueError('Relationship endpoints must be different entities')
    if evidence_id is None:
        support = db.scalar(select(ClaimEvidenceLink).where(
            ClaimEvidenceLink.claim_id == claim_id, ClaimEvidenceLink.stance == 'supports', ClaimEvidenceLink.superseded_at.is_(None)
        ).order_by(ClaimEvidenceLink.created_at))
        evidence_id = support.evidence_id if support else None
    payload = {
        'schema': schema, 'source_entity_id': source_entity_id, 'target_entity_id': target_entity_id,
        'properties': properties or {}, 'claim_id': claim_id, 'evidence_id': evidence_id,
        'note': note, 'span': (lineage.payload or {}).get('span'),
    }
    row = ExtractionCandidate(
        investigation_id=claim.investigation_id, document_id=lineage.document_id, chunk_id=lineage.chunk_id,
        candidate_type='relationship', payload=payload, confidence=0.0,
    )
    db.add(row); db.commit(); db.refresh(row); return row



def preview_entity_candidate_matches(db: Session, candidate: ExtractionCandidate) -> dict:
    """Return canonical entity matches before a proposed extraction entity can be accepted.

    Provider/OpenAleph resolved IDs are surfaced only as lead metadata. They never select a
    Workbench entity or create an identity decision automatically.
    """
    if candidate.candidate_type != 'entity':
        raise ValueError('Candidate is not an entity proposal')
    doc = db.get(Document, candidate.document_id)
    if doc is None:
        raise ValueError('Candidate document is missing')
    payload = candidate.payload or {}
    caption = str(payload.get('caption') or '').strip()
    schema = str(payload.get('suggested_schema') or 'Thing')
    # A kind the text did not show (the pattern finder's low band, or a name
    # OpenAleph gave no kind) is only a guess: compare as any legal entity.
    guessed = payload.get('schema_question') is True or (candidate.confidence or 0) < 0.5
    matches = candidate_entities(db, doc.investigation_id, caption, schema, schema_is_guess=guessed)
    raw_provenance = payload.get('provenance')
    provenance: dict = raw_provenance if isinstance(raw_provenance, dict) else {}
    return {
        'candidate_id': candidate.id, 'caption': caption, 'schema': schema,
        'matches': matches,
        # The same name (resolution.py SAME_NAME): accepting waits for the reporter's same-or-new answer.
        'strong_match': next((row for row in matches if row['basis'] == BASIS_SAME), None),
        'provider_lead': {
            'provider': provenance.get('provider'),
            'provider_entity_id': provenance.get('provider_entity_id'),
            'provider_resolved_entity_id': payload.get('provider_resolved_entity_id'),
        },
        'identity_rule': 'Provider resolution is a lead only. Reporter must choose same, different, or create_new.',
    }

def review_candidate(db: Session, candidate: ExtractionCandidate, *, decision: str, note: str|None=None, entity_schema: str|None=None, caption: str|None=None, claim_status: str='lead', confidence: float|None=None, matched_entity_id: str|None=None, entity_identity_decision: str|None=None, target_entity_id: str|None=None) -> dict:
    # Two people deciding at once: the second waits here for the first and then
    # sees its decision, instead of a status read before either wrote.
    lock_investigation_transaction(db, candidate.investigation_id)
    db.refresh(candidate)
    # A proposal held back by the OCR quality gate (app/jobs/handlers.py hold_illegible) can still be decided by a person.
    if candidate.review_status not in ('proposed', 'held'):
        raise ValueError('Candidate has already been reviewed')
    if decision not in {'accept','reject'}: raise ValueError('Decision must be accept or reject')
    candidate.reviewer_note = note
    candidate.reviewed_at = utcnow_naive()
    if decision == 'reject':
        candidate.review_status='rejected'; db.commit(); db.refresh(candidate); return {'candidate': serialize_candidate(candidate), 'record': None}
    doc=db.get(Document,candidate.document_id); source=db.get(Source,doc.source_id) if doc else None; chunk=db.get(DocumentChunk,candidate.chunk_id) if candidate.chunk_id else None
    if not doc or not source: raise ValueError('Document/source missing')
    span = candidate.payload.get('span') if isinstance(candidate.payload, dict) else None
    span_suffix = ''
    if isinstance(span, dict) and isinstance(span.get('start_char'), int) and isinstance(span.get('end_char'), int):
        span_suffix = f"@chars:{span['start_char']}-{span['end_char']}"
    # A proposal without a chunk (a spreadsheet row, table_mappings.py) points at its own locator.
    payload_locator = candidate.payload.get('locator') if isinstance(candidate.payload, dict) else None
    origin=f'document:{doc.id}#{chunk.locator if chunk else (payload_locator or "document")}{span_suffix}'
    # One typed slot per accepted-record kind: the branches below produce different
    # ORM types, and the serialized `record_out` is built after the commit so it
    # reflects persisted state (ids, defaults) exactly as before.
    evidence_record: Evidence | None = None
    claim_record: Claim | None = None
    entity_record: Entity | None = None
    edge_record: RelationshipEdge | None = None
    statement_record: Statement | None = None
    accepted_id: str
    if candidate.candidate_type=='evidence':
        evidence_quote = candidate.payload.get('quote')
        evidence_locator = candidate.payload.get('locator') or (chunk.locator if chunk else None)
        # A reviewed extraction span is part of the evidence locator itself, not only
        # side-channel candidate metadata. This keeps Evidence independently traceable
        # after it flows into claims, relationships, dossiers, exports, or search.
        if span_suffix:
            evidence_locator = f'{evidence_locator or "document"}{span_suffix}'
        if chunk is not None and isinstance(span, dict) and evidence_quote is not None:
            start, end = span.get('start_char'), span.get('end_char')
            if not isinstance(start, int) or not isinstance(end, int) or start < 0 or end < start or end > len(chunk.text):
                raise ValueError('Extraction candidate character span is outside the source chunk')
            if chunk.text[start:end] != evidence_quote:
                raise ValueError('Extraction candidate evidence quote does not match its exact source span')
        evidence_record=Evidence(source_id=source.id, quote=evidence_quote, locator=evidence_locator, notes=note)
        db.add(evidence_record); db.flush(); rtype='evidence'; accepted_id=evidence_record.id
    elif candidate.candidate_type=='claim':
        claim_text = candidate.payload.get('text','')
        claim_record=Claim(investigation_id=doc.investigation_id, text=claim_text, status=claim_status, confidence=confidence if confidence is not None else 0.0)
        db.add(claim_record); db.flush(); accepted_id=claim_record.id
        # Accepting an extracted claim must not sever it from the exact document text
        # the reporter reviewed. Materialize that precise span as Evidence and record
        # an explicit SUPPORTS link. This is provenance, not automatic verification:
        # the claim keeps the reporter-selected status/confidence unchanged.
        evidence_locator = chunk.locator if chunk else None
        if span_suffix:
            evidence_locator = f'{evidence_locator or "document"}{span_suffix}'
        extracted_evidence = Evidence(
            source_id=source.id,
            quote=claim_text,
            locator=evidence_locator,
            notes=f'Created from accepted extraction candidate {candidate.id}',
        )
        db.add(extracted_evidence); db.flush()
        db.add(ClaimEvidenceLink(
            claim_id=claim_record.id,
            evidence_id=extracted_evidence.id,
            stance='supports',
            note=f'Exact reviewed document span from extraction candidate {candidate.id}',
        ))
        rtype='claim'
    elif candidate.candidate_type=='entity':
        cap=(caption or candidate.payload.get('caption') or '').strip()
        # A kind the text gave no sign of (the name-finder's low band, or OpenAleph's name with no kind) is never
        # taken from the guess: the person says which kind, or "not sure" (LegalEntity).
        if not entity_schema and (candidate.confidence < 0.5 or candidate.payload.get('schema_question')):
            raise ValueError(f'The text does not say what kind of record “{cap}” is: say which (or “not sure”) before it is added')
        schema=entity_schema or candidate.payload.get('suggested_schema') or 'Thing'
        preview = preview_entity_candidate_matches(db, candidate)
        strong = preview.get('strong_match')
        if strong and entity_identity_decision is None:
            raise ValueError('A likely canonical entity match exists; review identity before accepting this proposal')
        if entity_identity_decision not in {None, 'same', 'different', 'create_new'}:
            raise ValueError('entity_identity_decision must be same, different, or create_new')
        if entity_identity_decision in {'same', 'different'} and not matched_entity_id:
            raise ValueError('matched_entity_id is required for same/different identity review')
        matched = db.get(Entity, matched_entity_id) if matched_entity_id else None
        if matched_entity_id and (matched is None or matched.investigation_id != doc.investigation_id):
            raise ValueError('Matched entity is missing or outside this investigation')
        candidate_properties = candidate.payload.get('properties') if isinstance(candidate.payload, dict) else None
        candidate_properties = candidate_properties if isinstance(candidate_properties, dict) else {}
        normalized_properties = {
            str(prop): [str(value) for value in values if str(value).strip()]
            for prop, values in candidate_properties.items()
            if isinstance(values, list)
        }
        if entity_identity_decision == 'same':
            assert matched is not None  # guaranteed by the matched_entity_id checks above
            entity_record = matched
        else:
            ftm=make_ftm_entity(schema, cap, normalized_properties)
            entity_record=Entity(investigation_id=doc.investigation_id, ftm_id=ftm['id'], schema=ftm['schema'], caption=cap, properties=ftm['properties'])
            db.add(entity_record); db.flush()
        accepted_id=entity_record.id
        provenance = candidate.payload.get('provenance') if isinstance(candidate.payload, dict) else None
        if isinstance(provenance, dict) and provenance.get('provider'):
            provider = str(provenance.get('provider'))
            collection_id = str(provenance.get('collection_id') or '').strip()
            provider_entity_id = str(provenance.get('provider_entity_id') or '').strip()
            dataset = f'{provider}:{collection_id}' if collection_id else provider
            # A name the Workbench found again in the document text points at
            # that exact place; one it could not find points at the provider's
            # record of the document (openaleph_analysis.import_analysis_proposals).
            if chunk is not None and span_suffix:
                statement_origin = origin
            else:
                statement_origin = f'{provider}:entity:{provider_entity_id}' if provider_entity_id else origin
        else:
            dataset = f'document:{doc.id}'
            statement_origin = origin
        # Preserve extraction/provider provenance even when the entity has no
        # extra properties. Reporter acceptance is the point at which this
        # proposal becomes a Workbench entity; provider resolution metadata is
        # not treated as a canonical identity decision.
        db.add(Statement(entity_id=entity_record.id, prop='name', value=cap, dataset=dataset, origin=statement_origin, original_value=cap))
        # Every other place the document names it, one statement per mention: the one record carries
        # all its mentions in this document, so each can be traced (extract_document folds them).
        seen_places = {statement_origin}
        for m in (candidate.payload.get('mentions') or [])[1:]:
            m_span = m.get('span') if isinstance(m, dict) else None
            suffix = f"@chars:{m_span['start_char']}-{m_span['end_char']}" if isinstance(m_span, dict) and isinstance(m_span.get('start_char'), int) else ''
            place = f"document:{doc.id}#{m.get('locator') or 'document'}{suffix}"
            if place not in seen_places:
                seen_places.add(place)
                db.add(Statement(entity_id=entity_record.id, prop='name', value=cap, dataset=dataset, origin=place, original_value=cap))
        if entity_identity_decision == 'different' and matched is not None:
            from app.services.resolution import record_canonical_resolution
            # Record the reporter's explicit negative identity decision without auto-merging.
            record_canonical_resolution(db, entity_record, matched, decision='different', confidence=confidence if confidence is not None else 1.0, rationale=note, commit=False)
        rtype='entity'
    elif candidate.candidate_type=='relationship' and (candidate.payload or {}).get('source_candidate_id'):
        # A link a spreadsheet mapping made (table_mappings.py): its two ends are
        # entity proposals of the same mapping, which the reporter decides first;
        # the rows it was read from become its evidence.
        from app.services.relationships import create_relationship
        payload = candidate.payload or {}
        ends = []
        for key in ('source_candidate_id', 'target_candidate_id'):
            end = db.get(ExtractionCandidate, payload.get(key))
            caption = ((end.payload or {}).get('caption') if end else None) or 'the other entity'
            if end is None or end.review_status == 'rejected':
                raise ValueError(f'“{caption}” was rejected, so this link cannot be accepted')
            if end.review_status != 'accepted' or end.accepted_record_type != 'entity' or db.get(Entity, end.accepted_record_id) is None:
                raise ValueError(f'Accept “{caption}” first')
            ends.append(str(end.accepted_record_id))
        evidence_record = Evidence(source_id=source.id, quote=payload.get('row_text'), locator=payload.get('locator'),
                                   notes=f'The spreadsheet row(s) a column mapping read this link from (proposal {candidate.id})')
        db.add(evidence_record); db.flush()
        edge_record = create_relationship(
            db, investigation_id=doc.investigation_id, schema=str(payload.get('schema') or ''),
            source_entity_id=ends[0], target_entity_id=ends[1], properties=payload.get('properties') or {},
            dataset=f'document:{doc.id}', origin=f'extraction-candidate:{candidate.id}', evidence_id=evidence_record.id, commit=False,
        )
        rtype='relationship'; accepted_id=edge_record.id
    elif candidate.candidate_type=='relationship':
        from app.services.relationships import create_relationship
        payload = candidate.payload or {}
        claim_id = payload.get('claim_id'); evidence_id = payload.get('evidence_id')
        claim = db.get(Claim, claim_id) if claim_id else None
        if claim is None or claim.investigation_id != doc.investigation_id:
            raise ValueError('Relationship proposal claim is missing or outside this investigation')
        # create_relationship performs endpoint/evidence scope validation. The relationship
        # statement origin points back to this exact reviewed extraction candidate.
        edge_record = create_relationship(
            db, investigation_id=doc.investigation_id, schema=str(payload.get('schema') or ''),
            source_entity_id=str(payload.get('source_entity_id') or ''), target_entity_id=str(payload.get('target_entity_id') or ''),
            properties=payload.get('properties') or {}, dataset=f'document:{doc.id}',
            origin=f'extraction-candidate:{candidate.id}', evidence_id=evidence_id, commit=False,
        )
        rtype='relationship'; accepted_id=edge_record.id
    elif candidate.candidate_type=='identifier':
        # An email or phone number found by OpenAleph: the reporter chose
        # whose it is; it becomes a statement on that entity (identifier_proposals.py).
        statement_record = accept_identifier(db, candidate, doc, origin=origin, target_entity_id=target_entity_id)
        rtype='statement'; accepted_id=statement_record.id
    else: raise ValueError('Unsupported candidate type')
    candidate.review_status='accepted'; candidate.accepted_record_type=rtype; candidate.accepted_record_id=accepted_id
    db.commit(); db.refresh(candidate)
    if rtype in {'entity', 'relationship'}:
        # An accepted entity or relationship goes to the OpenAleph collection
        # (ADR-0006, Wave B), queued once the acceptance is committed.
        request_corpus_sync(db, doc.investigation_id)
    record_out: dict
    if entity_record is not None:
        record_out = {'id': entity_record.id, 'investigation_id': entity_record.investigation_id, 'ftm_id': entity_record.ftm_id, 'schema': entity_record.schema, 'caption': entity_record.caption, 'properties': entity_record.properties}
    elif claim_record is not None:
        record_out = {'id': claim_record.id, 'investigation_id': claim_record.investigation_id, 'text': claim_record.text, 'status': claim_record.status, 'confidence': claim_record.confidence}
    elif edge_record is not None:
        from app.services.relationships import serialize_relationship
        record_out = serialize_relationship(db, edge_record)
    elif statement_record is not None:
        record_out = {'id': statement_record.id, 'entity_id': statement_record.entity_id, 'prop': statement_record.prop,
                      'value': statement_record.value, 'dataset': statement_record.dataset, 'origin': statement_record.origin,
                      'original_value': statement_record.original_value}
    else:
        assert evidence_record is not None
        record_out = {'id': evidence_record.id, 'source_id': evidence_record.source_id, 'quote': evidence_record.quote, 'locator': evidence_record.locator, 'notes': evidence_record.notes}
    return {'candidate': serialize_candidate(candidate), 'record': record_out}


def document_languages(db: Session, document_id: str) -> list[dict]:
    """The languages OpenAleph detected in the document, as its proposals record them (`detected_languages`, ISO 639
    codes), first seen first, each with its English name. The app's own reader detects no language, so a document
    OpenAleph has not read has none."""
    from followthemoney.types import registry

    codes: list[str] = []
    for payload in db.scalars(select(ExtractionCandidate.payload).where(ExtractionCandidate.document_id == document_id)):
        provenance = payload.get('provenance') if isinstance(payload, dict) else None
        found = provenance.get('detected_languages') if isinstance(provenance, dict) else None
        for code in found if isinstance(found, list) else []:
            clean = registry.language.clean(str(code))
            if clean and clean not in codes:
                codes.append(clean)
    return [{'code': code, 'name': registry.language.caption(code) or code} for code in codes]


def get_document_detail(db: Session, row: Document) -> dict:
    """Serialize a Document with its chunks attached, and the languages detected in it."""
    result = serialize_document(db, row)
    chunks = db.scalars(select(DocumentChunk).where(DocumentChunk.document_id == row.id).order_by(DocumentChunk.ordinal)).all()
    result["chunks"] = [serialize_chunk(chunk) for chunk in chunks]
    result["languages"] = document_languages(db, row.id)
    return result


def list_document_candidates(db: Session, document_id: str, status: str | None = None, candidate_type: str | None = None, *, explain: bool = False) -> list[dict]:
    """A document's proposals, oldest first. With `explain`, each carries the
    per-part reasoning behind it (#190) so the review list can show why every
    proposal exists without one request per row."""
    from app.services.extraction_explanations import explain_extraction_candidate

    stmt = select(ExtractionCandidate).where(ExtractionCandidate.document_id == document_id)
    if status:
        stmt = stmt.where(ExtractionCandidate.review_status == status)
    if candidate_type:
        stmt = stmt.where(ExtractionCandidate.candidate_type == candidate_type)
    rows = db.scalars(stmt.order_by(ExtractionCandidate.created_at)).all()
    if not explain:
        return [serialize_candidate(row) for row in rows]
    return [{**serialize_candidate(row), "explanation": explain_extraction_candidate(db, row)} for row in rows]
