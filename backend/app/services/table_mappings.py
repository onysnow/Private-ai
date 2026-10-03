"""Spreadsheets to entities (ADR-0006, Wave C): a reporter maps a ledger's
columns to FollowTheMoney entities, and what the mapping makes arrives as
proposals to review.

The mapping is a FollowTheMoney mapping (the ``entities`` part of
https://followthemoney.tech/docs/mappings/), the format OpenAleph's
``/collections/{id}/mappings`` stores, and it is run by the same engine
OpenAleph runs (``model.make_mapping`` -> ``QueryMapping.map`` per row,
aleph/logic/mapping.py ``map_to_aggregator``). The Workbench runs it itself
rather than through OpenAleph's endpoint because OpenAleph writes the mapped
entities straight into the collection's aggregator: they would be in the
corpus without a reporter having looked at them, and the Workbench's own
records (which it pushes to that collection, Wave B) are the ones a reporter
confirmed. Here, every entity becomes an ``entity`` proposal (reviewed with
the same same-or-new identity step as any name) and every link (Payment,
Ownership, ...) a ``relationship`` proposal whose two ends are those entity
proposals; accepting a link records the spreadsheet row as its evidence.

Choices, each in the mapping's own terms:

* Entity kinds are the legal entities the Workbench's records use
  (Person, Company, Organization, PublicBody, LegalEntity); links are the
  relationship kinds it draws (app/services/relationships.py
  RELATIONSHIP_SCHEMAS). A link names its two ends with ``entity``; other
  properties pointing at entities are refused, since the Workbench keeps
  links as relationships, not as references inside a record.
* FollowTheMoney derives an entity's id from its key columns only, so two
  entities keyed on the same column would collide; each entity's name is
  used as its ``key_literal`` unless the mapping gives one, and the
  document's id is the key prefix (OpenAleph uses the collection's).
* Row numbers are the spreadsheet's own: the header is row 1.
* A value FollowTheMoney cannot read for its property (a date it cannot
  parse, an amount that is not a number) is dropped by the engine; it is
  reported against its row instead of disappearing silently.
"""

from __future__ import annotations

import csv
import io
import logging
import threading
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

from followthemoney import model
from followthemoney.exc import FollowTheMoneyException
from followthemoney.mapping.query import QueryMapping
from followthemoney.types import registry
from normality import guess_encoding
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.domain import Document, ExtractionCandidate, TableMapping
from app.services.relationships import RELATIONSHIP_SCHEMAS, validate_payment_properties

TABLE_EXTENSIONS = {"csv": ",", "tsv": "\t", "xlsx": None}
ENTITY_SCHEMAS = ["Person", "Company", "Organization", "PublicBody", "LegalEntity"]
MAX_ROWS = 5000          # rows a mapping reads; every result is reviewed by hand
SAMPLE_ROWS = 20         # rows shown with the columns
PREVIEW_ROWS = 10        # rows a preview maps
METHOD = "column_mapping"

_engine_log = logging.getLogger("followthemoney.mapping.entity")


class TableError(ValueError):
    """The file cannot be read as a table, or the mapping cannot run; the
    message says why in plain words."""


@dataclass
class Table:
    columns: list[str]
    rows: list[tuple[int, dict[str, str]]]      # (spreadsheet row number, column -> cell)
    sheets: list[str]
    sheet: str | None
    truncated: bool = False


def _extension(doc: Document) -> str:
    return Path(doc.filename or "").suffix.lower().lstrip(".")


def is_table(doc: Document) -> bool:
    return _extension(doc) in TABLE_EXTENSIONS


def _headers(raw: list[Any]) -> list[str]:
    """Blank headers get "Column N"; a repeated one gets " (2)", " (3)"."""
    seen: dict[str, int] = {}
    columns = []
    for index, value in enumerate(raw, 1):
        name = _cell(value) or f"Column {index}"
        seen[name] = seen.get(name, 0) + 1
        columns.append(name if seen[name] == 1 else f"{name} ({seen[name]})")
    return columns


def _cell(value: Any) -> str:
    """A cell as the text a reporter sees: a date without a midnight time,
    a whole number without ".0"."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _rows(raw_rows: Any, limit: int | None) -> tuple[list[str], list[tuple[int, dict[str, str]]], bool]:
    columns: list[str] | None = None
    rows: list[tuple[int, dict[str, str]]] = []
    for number, raw in enumerate(raw_rows, 1):
        values = list(raw)
        if columns is None:
            if not any(_cell(v) for v in values):
                continue      # leading blank rows before the header
            columns = _headers(values)
            continue
        cells = {column: _cell(values[i]) if i < len(values) else "" for i, column in enumerate(columns)}
        if not any(cells.values()):
            continue
        if limit is not None and len(rows) >= limit:
            return columns, rows, True
        rows.append((number, cells))
    if columns is None:
        raise TableError("The file has no header row")
    return columns, rows, False


def read_table(doc: Document, *, sheet: str | None = None, limit: int | None = MAX_ROWS) -> Table:
    """The document's rows. CSV and TSV are decoded with the encoding
    normality guesses (the library FollowTheMoney uses); a CSV's delimiter is
    sniffed among , ; tab and |. XLSX is read with openpyxl (cell values, not
    formulas); ``sheet`` picks the sheet, the first by default."""
    extension = _extension(doc)
    if extension not in TABLE_EXTENSIONS:
        raise TableError("Only CSV, TSV and XLSX files can be mapped to entities")
    path = Path(doc.storage_path)
    if not path.is_file():
        raise TableError("The stored file is missing")
    if extension == "xlsx":
        from openpyxl import load_workbook

        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:   # openpyxl raises several kinds for a damaged file
            raise TableError(f"The workbook could not be opened: {exc}") from None
        try:
            names = list(workbook.sheetnames)
            chosen = sheet or (names[0] if names else None)
            if chosen not in names:
                raise TableError(f"The workbook has no sheet named “{sheet}”")
            columns, rows, truncated = _rows(workbook[chosen].iter_rows(values_only=True), limit)
        finally:
            workbook.close()
        return Table(columns, rows, names, chosen, truncated)
    data = path.read_bytes()
    text = data.decode(guess_encoding(data) or "utf-8", errors="replace").lstrip("﻿")
    delimiter = TABLE_EXTENSIONS[extension]
    if extension == "csv":
        try:
            delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    columns, rows, truncated = _rows(csv.reader(io.StringIO(text), delimiter=delimiter or ","), limit)
    return Table(columns, rows, [], None, truncated)


# --- the mapping ----------------------------------------------------------------------


def _columns_of(spec: dict[str, Any]) -> list[str]:
    out = [spec["column"]] if spec.get("column") else []
    return out + list(spec.get("columns") or [])


def check_query(query: dict[str, Any], table: Table) -> QueryMapping:
    """Refuse a mapping the Workbench would not keep (see the module notes),
    then let FollowTheMoney check the rest as OpenAleph does on save
    (aleph/views/mappings_api.py ``load_query``)."""
    if not query:
        raise TableError("Map at least one entity")
    known = set(table.columns)
    for name, entity in query.items():
        schema = str(entity.get("schema") or "")
        if schema not in ENTITY_SCHEMAS and schema not in RELATIONSHIP_SCHEMAS:
            raise TableError(f"“{name}”: {schema or 'no kind'} is not a kind the Workbench keeps")
        missing = [c for c in list(entity.get("keys") or []) if c not in known]
        props = entity.get("properties") or {}
        for prop, spec in props.items():
            missing += [c for c in _columns_of(spec) if c not in known]
            ftm_schema = model.get(schema)
            ftm_prop = ftm_schema.get(prop) if ftm_schema is not None else None
            if ftm_prop is not None and ftm_prop.type == registry.entity:
                if schema not in RELATIONSHIP_SCHEMAS:
                    raise TableError(f"“{name}”: {prop} points at another entity; map that as a link")
                info = RELATIONSHIP_SCHEMAS[schema]
                if prop not in {info["source_prop"], info["target_prop"]}:
                    raise TableError(f"“{name}”: a {schema} link only names its {info['source_label'].lower()} and {info['target_label'].lower()}")
        if missing:
            raise TableError(f"“{name}”: no column named {', '.join(f'“{c}”' for c in sorted(set(missing)))}")
        if schema in RELATIONSHIP_SCHEMAS:
            info = RELATIONSHIP_SCHEMAS[schema]
            for end in ("source_prop", "target_prop"):
                ref = (props.get(info[end]) or {}).get("entity")
                if not ref or ref not in query or query[ref].get("schema") in RELATIONSHIP_SCHEMAS:
                    raise TableError(f"“{name}”: choose the entity that is the {info[end.replace('prop', 'label')].lower()}")
    try:
        return _engine(query, key_prefix=None)
    except FollowTheMoneyException as exc:
        raise TableError(str(exc)) from None


def _engine(query: dict[str, Any], *, key_prefix: str | None) -> QueryMapping:
    entities = {name: {**entity, "key_literal": entity.get("key_literal") or name} for name, entity in query.items()}
    return model.make_mapping({"csv_url": "file:///dev/null", "entities": entities}, key_prefix=key_prefix)


class _Dropped(logging.Handler):
    """The engine's "Discarding unclean value" warnings for this thread's row."""

    def __init__(self) -> None:
        super().__init__(logging.WARNING)
        self.thread = threading.get_ident()
        self.messages: list[str] = []

    def emit(self, record: logging.LogRecord) -> None:
        if record.thread == self.thread:
            self.messages.append(record.getMessage())


def _dropped_text(message: str) -> str | None:
    """The engine's warning in plain words; None for an empty cell, which is
    not a problem (the value is simply absent)."""
    name = message[1:message.index("]")] if message.startswith("[") and "]" in message else "?"
    if "no ID could be computed" in message:
        return f"“{name}” is left out: its key column is empty in this row"
    if "unclean value" not in message:
        return message
    # '[pay] Discarding unclean value "3/14" for property "Interval:date".'
    value = message.split('value "', 1)[-1].split('"', 1)[0]
    if not value.strip():
        return None
    prop = message.rsplit(":", 1)[-1].rstrip('".')
    return f"“{value}” is not a valid {prop}; it was left out"


@dataclass
class MappedRow:
    number: int
    cells: dict[str, str]
    entities: dict[str, Any] = field(default_factory=dict)      # mapping name -> EntityProxy
    problems: list[str] = field(default_factory=list)


def map_rows(engine: QueryMapping, table: Table, *, limit: int | None = None) -> list[MappedRow]:
    out = []
    handler = _Dropped()
    _engine_log.addHandler(handler)
    try:
        for number, cells in table.rows[:limit] if limit else table.rows:
            handler.messages = []
            row = MappedRow(number, cells, dict(engine.map(cells)))
            row.problems = [text for text in map(_dropped_text, handler.messages) if text]
            out.append(row)
    finally:
        _engine_log.removeHandler(handler)
    return out


def _caption(proxy: Any) -> str | None:
    for prop in proxy.schema.caption:
        for value in proxy.get(prop, quiet=True):
            if str(value).strip():
                return str(value).strip()
    return None


def _plain_properties(proxy: Any) -> dict[str, list[str]]:
    return {prop.name: list(proxy.get(prop)) for prop in proxy.iterprops() if prop.type != registry.entity}


def describe_row(row: MappedRow, query: dict[str, Any]) -> dict[str, Any]:
    """One previewed row: each entity and link it makes, in plain terms."""
    entities = []
    for name, proxy in row.entities.items():
        schema = proxy.schema.name
        item: dict[str, Any] = {"name": name, "schema": schema, "caption": _caption(proxy), "properties": _plain_properties(proxy)}
        if schema in RELATIONSHIP_SCHEMAS:
            info = RELATIONSHIP_SCHEMAS[schema]
            ends = {end: (query[name]["properties"][info[f"{end}_prop"]]["entity"]) for end in ("source", "target")}
            source, target = row.entities.get(ends["source"]), row.entities.get(ends["target"])
            item["source"] = _caption(source) if source is not None else None
            item["target"] = _caption(target) if target is not None else None
            item["caption"] = f"{item['source'] or '?'} → {item['target'] or '?'}"
        entities.append(item)
    return {"number": row.number, "cells": row.cells, "entities": entities, "problems": list(row.problems)}


def preview(doc: Document, query: dict[str, Any], *, sheet: str | None) -> dict[str, Any]:
    table = read_table(doc, sheet=sheet, limit=PREVIEW_ROWS)
    engine = check_query(query, table)
    rows = map_rows(engine, table)
    return {"document_id": doc.id, "sheet": table.sheet, "rows": [describe_row(row, query) for row in rows]}


# --- proposals --------------------------------------------------------------------------


def _row_text(cells: dict[str, str]) -> str:
    return "; ".join(f"{column}: {value}" for column, value in cells.items() if value)


def _where(table: Table, numbers: list[int]) -> str:
    rows = f"row {numbers[0]}" if len(numbers) == 1 else "rows " + ", ".join(str(n) for n in numbers)
    return f"{table.sheet}, {rows}" if table.sheet else rows


def _add_row(item: dict[str, Any], table: Table, row: MappedRow, properties: dict[str, list[str]]) -> None:
    """One more row that makes this entity or link: its number, its cells
    (the evidence quote) and any values it adds."""
    item["rows"].append(row.number)
    item["locator"] = _where(table, item["rows"])
    item["row_text"] = "\n".join(filter(None, [item.get("row_text"), _row_text(row.cells)]))
    for prop, values in properties.items():
        kept = item["properties"].setdefault(prop, [])
        kept += [v for v in values if v not in kept]


def run_mapping(db: Session, doc: Document, query: dict[str, Any], *, sheet: str | None, created_by: str | None) -> dict[str, Any]:
    """Map every row and propose what it makes. An entity made by several
    rows (the same key values) is one proposal listing those rows; an entity
    or link this document's mappings already proposed is not proposed again.
    Returns the mapping and what it proposed and skipped."""
    table = read_table(doc, sheet=sheet)
    if table.truncated:
        raise TableError(f"The sheet has more than {MAX_ROWS} rows; split it, or map it in OpenAleph and review the results there")
    check_query(query, table)
    mapping = TableMapping(investigation_id=doc.investigation_id, document_id=doc.id, sheet=table.sheet, query=query,
                           created_by=created_by, rows_read=len(table.rows))
    db.add(mapping)
    db.flush()
    rows = map_rows(_engine(query, key_prefix=doc.id), table)

    existing: dict[str, ExtractionCandidate] = {}
    for candidate in db.scalars(select(ExtractionCandidate).where(ExtractionCandidate.document_id == doc.id)):
        mapped = (candidate.payload or {}).get("mapping") if isinstance(candidate.payload, dict) else None
        if isinstance(mapped, dict) and mapped.get("ftm_id"):
            existing[str(mapped["ftm_id"])] = candidate

    skipped: list[dict[str, Any]] = []
    entities: dict[str, dict[str, Any]] = {}       # ftm id -> payload being built
    links: dict[str, dict[str, Any]] = {}
    for row in rows:
        for problem in row.problems:
            skipped.append({"row": row.number, "reason": problem})
        for name, proxy in row.entities.items():
            schema = proxy.schema.name
            if schema in RELATIONSHIP_SCHEMAS:
                continue
            caption = _caption(proxy)
            if caption is None:
                skipped.append({"row": row.number, "reason": f"“{name}” has no name in this row"})
                continue
            item = entities.setdefault(proxy.id, {
                "caption": caption, "suggested_schema": schema, "properties": {}, "rows": [],
                "mapping": {"id": mapping.id, "entity": name, "ftm_id": proxy.id},
                "rationale": {"rule": METHOD, "entity": name, "schema": schema,
                              "columns": {prop: _columns_of(spec) for prop, spec in (query[name].get("properties") or {}).items()}},
            })
            _add_row(item, table, row, _plain_properties(proxy))
        for name, proxy in row.entities.items():
            schema = proxy.schema.name
            if schema not in RELATIONSHIP_SCHEMAS:
                continue
            info = RELATIONSHIP_SCHEMAS[schema]
            ends = {end: query[name]["properties"][info[f"{end}_prop"]]["entity"] for end in ("source", "target")}
            source, target = row.entities.get(ends["source"]), row.entities.get(ends["target"])
            if source is None or target is None or source.id not in entities or target.id not in entities:
                skipped.append({"row": row.number, "reason": f"“{name}” is missing its {info['source_label'].lower() if source is None or source.id not in entities else info['target_label'].lower()} in this row"})
                continue
            properties = _plain_properties(proxy)
            if schema == "Payment":
                try:
                    validate_payment_properties(properties)
                except ValueError as exc:
                    skipped.append({"row": row.number, "reason": f"“{name}”: {exc}"})
                    continue
            link = links.setdefault(proxy.id, {
                "schema": schema, "caption": f"{entities[source.id]['caption']} → {entities[target.id]['caption']}",
                "source_ftm_id": source.id, "target_ftm_id": target.id, "properties": {},
                "rows": [],
                "mapping": {"id": mapping.id, "entity": name, "ftm_id": proxy.id},
                "rationale": {"rule": METHOD, "entity": name, "schema": schema, "source": ends["source"], "target": ends["target"],
                              "columns": {prop: _columns_of(spec) for prop, spec in (query[name].get("properties") or {}).items()}},
            })
            _add_row(link, table, row, properties)

    made_entities = made_links = already = 0
    candidate_for: dict[str, str] = {ftm_id: c.id for ftm_id, c in existing.items()}
    for ftm_id, payload in entities.items():
        if ftm_id in existing:
            already += 1
            continue
        proposal = ExtractionCandidate(investigation_id=doc.investigation_id, document_id=doc.id, chunk_id=None,
                                       candidate_type="entity", payload=payload, confidence=1.0)
        db.add(proposal)
        db.flush()
        candidate_for[ftm_id] = proposal.id
        made_entities += 1
    for ftm_id, payload in links.items():
        if ftm_id in existing:
            already += 1
            continue
        payload["source_candidate_id"] = candidate_for[payload.pop("source_ftm_id")]
        payload["target_candidate_id"] = candidate_for[payload.pop("target_ftm_id")]
        db.add(ExtractionCandidate(investigation_id=doc.investigation_id, document_id=doc.id, chunk_id=None,
                                   candidate_type="relationship", payload=payload, confidence=1.0))
        made_links += 1
    mapping.entities_proposed = made_entities
    mapping.links_proposed = made_links
    db.commit()
    db.refresh(mapping)
    return {"mapping": serialize_mapping(mapping), "entities_proposed": made_entities, "links_proposed": made_links,
            "already_proposed": already, "skipped": skipped}


def serialize_mapping(mapping: TableMapping) -> dict[str, Any]:
    return {"id": mapping.id, "document_id": mapping.document_id, "sheet": mapping.sheet, "query": mapping.query,
            "created_by": mapping.created_by, "created_at": mapping.created_at, "rows_read": mapping.rows_read,
            "entities_proposed": mapping.entities_proposed, "links_proposed": mapping.links_proposed}


def list_mappings(db: Session, doc: Document) -> list[dict[str, Any]]:
    rows = db.scalars(select(TableMapping).where(TableMapping.document_id == doc.id).order_by(TableMapping.created_at.desc()))
    return [serialize_mapping(row) for row in rows]


def table_view(doc: Document, *, sheet: str | None) -> dict[str, Any]:
    table = read_table(doc, sheet=sheet)
    return {"document_id": doc.id, "filename": doc.filename, "sheets": table.sheets, "sheet": table.sheet,
            "columns": table.columns, "row_count": len(table.rows), "truncated": table.truncated,
            "rows": [{"number": number, "cells": cells} for number, cells in table.rows[:SAMPLE_ROWS]]}


def mapping_schemas() -> dict[str, Any]:
    """The kinds a mapping can make and the properties each can take from a
    column, from the FollowTheMoney model (labels as FollowTheMoney writes them)."""
    def props(schema_name: str, skip: set[str]) -> list[dict[str, Any]]:
        schema = model.get(schema_name)
        assert schema is not None
        found = [{"name": p.name, "label": p.label, "type": p.type.name}
                 for p in schema.properties.values()
                 if not p.hidden and not p.stub and p.type != registry.entity and p.name not in skip]
        return sorted(found, key=lambda p: (p["name"] != "name", p["label"]))

    meta = {"proof", "sourceUrl", "publisher", "publisherUrl", "modifiedAt", "retrievedAt", "indexText", "createdAt"}
    return {
        "entities": [{"name": n, "label": model.get(n).label, "properties": props(n, meta)} for n in ENTITY_SCHEMAS],  # type: ignore[union-attr]
        "links": [{"name": n, "label": info["label"], "source_prop": info["source_prop"], "target_prop": info["target_prop"],
                   "source_label": info["source_label"], "target_label": info["target_label"], "properties": props(n, meta)}
                  for n, info in RELATIONSHIP_SCHEMAS.items()],
    }
