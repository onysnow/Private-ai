"""The Workbench's own reader for emails, mailboxes, .zip archives and
spreadsheets (.eml, .mbox, .zip, .xlsx).

``documents.extract_chunks`` hands these four kinds of file here. Each
reader returns what every other reader returns: ``(chunks, issues)``, where
a chunk is a passage of text with a locator that says where in the file it
was found, and an issue is a plain-language note about something that was
not read (an attachment, a file inside an archive that was skipped, rows
past the limit). Every proposal made from these chunks therefore has a
passage to show and a place to point to.

What each reader produces:

- **.eml** – one chunk ``headers`` ("From: … · To: … · Date: … · Subject: …"),
  then the body as paragraphs, ``body, paragraph N``. The plain-text body is
  read when there is one, otherwise the HTML body's text. Attachments are
  listed by name in an issue; they are not read.
- **.mbox** – each message as an .eml, its locators starting
  ``message N · `` (``message 2 · body, paragraph 1``). At most
  ``MAX_MBOX_MESSAGES`` messages are read.
- **.zip** – each file inside that the Workbench can read on its own is read
  as if it had been uploaded, its locators starting with its path in the
  archive (``letters/first.eml · body, paragraph 1``). Archives inside the
  archive and pictures are not read, and a file with an unsafe path, a
  password, or too many bytes is skipped; the skipped files are listed in
  an issue. The limits (``MAX_ZIP_MEMBERS``, the per-file and total sizes)
  keep a small archive that unpacks to gigabytes (a "zip bomb") from
  filling the disk or the reader's memory.
- **.xlsx** – each row with a value is a chunk ``<sheet> row N`` reading
  "Header: value · Header: value", the sheet's first row with a value
  being its headers. At most ``MAX_XLSX_ROWS`` rows are read.
"""

from __future__ import annotations

import email
import mailbox
import re
import tempfile
import zipfile
from email import policy
from email.message import EmailMessage
from pathlib import Path, PurePosixPath
from typing import Any

from app.core.config import settings

CONTAINER_EXTENSIONS = frozenset({".eml", ".mbox", ".zip", ".xlsx"})

MAX_MBOX_MESSAGES = 2000
MAX_ZIP_MEMBERS = 500
# Bytes read out of one archive in all: four uploads' worth.
MAX_ZIP_TOTAL_FACTOR = 4
MAX_XLSX_ROWS = 20000

SEP = " · "
HEADER_FIELDS = ("From", "To", "Cc", "Date", "Subject")
_PARAGRAPH_SPLIT = re.compile(r"\n\s*\n")


def is_container_file(filename: str | None) -> bool:
    """An email, mailbox, .zip or workbook, by its file name."""
    return bool(filename) and PurePosixPath(str(filename).replace("\\", "/")).suffix.lower() in CONTAINER_EXTENSIONS


def _clean(text: str) -> str:
    return re.sub(r"\s+", " ", (text or "").replace("\x00", "")).strip()


def _listed(names: list[str], limit: int = 20) -> str:
    shown = ", ".join(names[:limit])
    return shown + (f" and {len(names) - limit} more" if len(names) > limit else "")


def read_container(path: Path) -> tuple[list[dict], list[str]]:
    """Read one of the four kinds of file this module knows, by its extension."""
    ext = path.suffix.lower()
    if ext == ".eml":
        return read_eml(path)
    if ext == ".mbox":
        return read_mbox(path)
    if ext == ".zip":
        return read_zip(path)
    if ext == ".xlsx":
        return read_xlsx(path)
    raise ValueError(f"Unsupported document type: {ext or 'unknown'}")


# --- email ------------------------------------------------------------------------------


def _header(message: EmailMessage, name: str) -> str:
    """A header as a person reads it (encoded words decoded); empty when it
    is missing or too damaged to read."""
    try:
        values = message.get_all(name) or []
    except Exception:
        return ""
    return _clean(", ".join(str(value) for value in values))


def _part_text(part: Any) -> str:
    """A body part's text, even when its declared character set is wrong
    or unknown."""
    try:
        return str(part.get_content())
    except Exception:
        payload = part.get_payload(decode=True) or b""
        return payload.decode("utf-8", errors="replace") if isinstance(payload, bytes) else str(payload)


def _body_paragraphs(message: EmailMessage) -> list[str]:
    """The body's paragraphs: the plain-text body when there is one, else
    the text of the HTML body."""
    try:
        plain = message.get_body(preferencelist=("plain",))
        html = None if plain is not None else message.get_body(preferencelist=("html",))
    except Exception:
        plain = html = None
    if plain is not None:
        text = _part_text(plain).replace("\r\n", "\n")
        return [p for p in (_clean(chunk) for chunk in _PARAGRAPH_SPLIT.split(text)) if p]
    if html is not None:
        from app.services.html_reader import html_blocks

        return [p for p in (_clean(block) for block in html_blocks(_part_text(html))) if p]
    return []


def _attachment_names(message: EmailMessage) -> list[str]:
    names = []
    try:
        attachments = list(message.iter_attachments()) if message.is_multipart() else []
    except Exception:
        attachments = []
    for part in attachments:
        try:
            name = part.get_filename()
        except Exception:
            name = None
        names.append(name or f"an unnamed {part.get_content_type()} part")
    return names


def _message_chunks(message: EmailMessage, prefix: str = "") -> tuple[list[dict], list[str]]:
    """One email's chunks (headers, then body paragraphs) and its issues;
    ``prefix`` goes before every locator (an mbox's "message N · ")."""
    rows: list[dict] = []
    issues: list[str] = []
    headers = [(name, _header(message, name)) for name in HEADER_FIELDS]
    header_text = SEP.join(f"{name}: {value}" for name, value in headers if value)
    if header_text:
        rows.append({"locator": f"{prefix}headers", "page_number": None, "text": header_text})
    paragraphs = _body_paragraphs(message)
    for number, paragraph in enumerate(paragraphs, 1):
        rows.append({"locator": f"{prefix}body, paragraph {number}", "page_number": None, "text": paragraph})
    if not paragraphs:
        issues.append("This email has no body text.")
    attachments = _attachment_names(message)
    if attachments:
        issues.append(f"This email has {len(attachments)} attachment(s), which the Workbench's own reader "
                      f"does not read: {_listed(attachments)}.")
    return rows, issues


def _parse(data: bytes) -> EmailMessage:
    return email.message_from_bytes(data, policy=policy.default)


def read_eml(path: Path) -> tuple[list[dict], list[str]]:
    """An email file: its headers, its body's paragraphs, its attachments named."""
    return _message_chunks(_parse(path.read_bytes()))


def read_mbox(path: Path) -> tuple[list[dict], list[str]]:
    """A mailbox file: each message read as an email, numbered from 1."""
    box = mailbox.mbox(str(path), create=False)
    rows: list[dict] = []
    issues: list[str] = []
    count = 0
    try:
        for key in box.iterkeys():
            count += 1
            if count > MAX_MBOX_MESSAGES:
                continue   # counted, not read
            message_rows, message_issues = _message_chunks(_parse(box.get_bytes(key)), prefix=f"message {count}{SEP}")
            rows.extend(message_rows)
            issues.extend(f"Message {count}: {issue}" for issue in message_issues)
    finally:
        box.close()
    if count == 0:
        issues.append("This mailbox has no messages.")
    if count > MAX_MBOX_MESSAGES:
        issues.append(f"This mailbox has {count} messages; only the first {MAX_MBOX_MESSAGES} were read.")
    return rows, issues


# --- .zip -------------------------------------------------------------------------------


def _size(count: int) -> str:
    return f"{count // (1024 * 1024)} MB" if count >= 1024 * 1024 else f"{count} bytes"


def _unsafe(name: str) -> bool:
    """A path that would land outside the archive's own folder when unpacked:
    absolute, on a drive, or climbing out with "..". """
    if name.startswith(("/", "\\")) or re.match(r"^[A-Za-z]:", name):
        return True
    return ".." in re.split(r"[\\/]", name)


def _copy_capped(archive: zipfile.ZipFile, info: zipfile.ZipInfo, target: Path, limit: int) -> int | None:
    """Unpack one file to ``target``, stopping past ``limit`` bytes (the size
    an archive declares can be false). Returns the bytes written, or None when
    the file was larger than the limit."""
    written = 0
    with archive.open(info) as source, target.open("wb") as out:
        while True:
            block = source.read(min(1 << 20, limit + 1 - written))
            if not block:
                return written
            written += len(block)
            if written > limit:
                return None
            out.write(block)


def read_zip(path: Path) -> tuple[list[dict], list[str]]:
    """A .zip archive: each file inside that the Workbench reads on its own,
    read as if uploaded, its locators starting with its path in the archive."""
    from app.services.documents import SUPPORTED_EXTENSIONS, extract_chunks
    from app.services.image_reader import IMAGE_EXTENSIONS

    try:
        archive = zipfile.ZipFile(path)
    except (zipfile.BadZipFile, OSError) as exc:
        raise ValueError(f"This .zip file could not be opened: {exc}") from None
    per_file = settings.max_document_bytes
    total_limit = settings.max_document_bytes * MAX_ZIP_TOTAL_FACTOR
    rows: list[dict] = []
    issues: list[str] = []
    skipped: list[str] = []
    total = 0
    read = 0
    with archive, tempfile.TemporaryDirectory(prefix="workbench-zip-") as scratch:
        files = [info for info in archive.infolist() if not info.is_dir()]
        if len(files) > MAX_ZIP_MEMBERS:
            issues.append(f"This archive holds {len(files)} files; only the first {MAX_ZIP_MEMBERS} were looked at.")
            files = files[:MAX_ZIP_MEMBERS]
        for info in files:
            name = info.filename
            if name.startswith("__MACOSX/") or PurePosixPath(name).name.startswith("._"):
                continue   # macOS's hidden copies of file details, not files of the case
            ext = PurePosixPath(name.replace("\\", "/")).suffix.lower()
            if _unsafe(name):
                skipped.append(f"{name} (its path points outside the archive)")
            elif info.flag_bits & 0x1:
                skipped.append(f"{name} (protected by a password)")
            elif ext == ".zip":
                skipped.append(f"{name} (an archive inside the archive)")
            elif ext in IMAGE_EXTENSIONS:
                skipped.append(f"{name} (a picture)")
            elif ext not in SUPPORTED_EXTENSIONS:
                skipped.append(f"{name} (a kind of file the Workbench's own reader does not read)")
            elif info.file_size > per_file:
                skipped.append(f"{name} (larger than the {_size(per_file)} limit)")
            elif total >= total_limit:
                skipped.append(f"{name} (the archive's files had reached {_size(total_limit)} in all)")
            else:
                target = Path(scratch) / f"file{read}{ext}"
                try:
                    written = _copy_capped(archive, info, target, min(per_file, total_limit - total))
                except Exception as exc:   # damaged entry, unsupported compression
                    skipped.append(f"{name} (it could not be unpacked: {type(exc).__name__})")
                    continue
                if written is None:
                    skipped.append(f"{name} (it unpacks to more than the size limit)")
                    target.unlink(missing_ok=True)
                    continue
                total += written
                read += 1
                try:
                    member_rows, member_issues = extract_chunks(target)
                except Exception as exc:
                    skipped.append(f"{name} (it could not be read: {exc or type(exc).__name__})")
                    continue
                finally:
                    target.unlink(missing_ok=True)
                for row in member_rows:
                    rows.append({**row, "locator": f"{name}{SEP}{row['locator']}"})
                issues.extend(f"{name}: {issue}" for issue in member_issues)
    if skipped:
        issues.append(f"{len(skipped)} file(s) in this archive were not read: {_listed(skipped)}.")
    if not rows and not skipped:
        issues.append("This archive has no files the Workbench's own reader could read text from.")
    return rows, issues


# --- .xlsx ------------------------------------------------------------------------------


def read_xlsx(path: Path) -> tuple[list[dict], list[str]]:
    """A workbook: each row with a value, in every sheet, as "Header: value"
    pairs, the sheet's first row with a value being its headers. Cell values
    are read, not formulas (a formula reads as the value Excel last saved)."""
    from openpyxl import load_workbook

    from app.services.table_mappings import _cell, _headers

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:   # openpyxl raises several kinds for a damaged file
        raise ValueError(f"The workbook could not be opened: {exc}") from None
    rows: list[dict] = []
    issues: list[str] = []
    capped = False
    try:
        for sheet in workbook.worksheets:
            columns: list[str] | None = None
            for number, raw in enumerate(sheet.iter_rows(values_only=True), 1):
                values: list[Any] = list(raw)
                if not any(_cell(v) for v in values):
                    continue
                if columns is None:
                    columns = _headers(values)
                    continue
                if len(rows) >= MAX_XLSX_ROWS:
                    capped = True
                    break
                pairs = []
                for index, value in enumerate(values):
                    text = _cell(value)
                    if text:
                        column = columns[index] if index < len(columns) else f"Column {index + 1}"
                        pairs.append(f"{column}: {_clean(text)}")
                rows.append({"locator": f"{sheet.title} row {number}", "page_number": None, "text": SEP.join(pairs)})
            if capped:
                break
    finally:
        workbook.close()
    if capped:
        issues.append(f"Only the first {MAX_XLSX_ROWS} rows of this workbook were read.")
    if not rows:
        issues.append("This workbook has no rows with values under a header row.")
    return rows, issues
