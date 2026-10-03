"""Deck themes, font pairings, and per-visual presets/options
(app/services/presentation_style.py): the catalogue endpoint, validation,
promotion with styles, and what the .pptx export does with them.

Style is presentation only, so several tests check that no preset or option
changes a value, drops a citation, or starts a bar anywhere but zero.
"""

import json
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

import app.ai.reasoning as reasoning_module
from app.ai.governor import governor
from app.core.config import settings
from app.main import app
from app.services import presentation_style as ps

client = TestClient(app)


class _FakeLLMClient:
    def __init__(self, response_text: str):
        self._response_text = response_text

    def generate(self, system_prompt: str, user_prompt: str, *, max_tokens: int = 4096) -> str:
        return self._response_text


@pytest.fixture(autouse=True)
def _restore_ai_settings():
    original = (settings.enable_ai_features, settings.ai_provider, settings.ai_runs_per_minute, settings.ai_max_concurrent_runs)
    governor.reset()
    yield
    settings.enable_ai_features, settings.ai_provider, settings.ai_runs_per_minute, settings.ai_max_concurrent_runs = original
    governor.reset()


def _accepted_candidate(monkeypatch) -> dict:
    settings.enable_ai_features = True
    settings.ai_provider = "anthropic"
    inv = client.post("/api/investigations", json={"name": "Style tests"}).json()
    claim = client.post("/api/claims", json={"investigation_id": inv["id"], "text": "Meridian paid Bellcastle a fee."}).json()
    source = client.post("/api/sources", json={"investigation_id": inv["id"], "title": "Wire log", "source_type": "filing"}).json()
    evidence = client.post("/api/evidence", json={"source_id": source["id"], "locator": "row 41", "quote": (
        "$2.4M wired to Bellcastle. Agreement signed on 2023-11-03; paid on 2023-11-14, 11 days after signing. "
        "Fees: Arden 285,000 USD; Bellcastle 2,400,000 USD; Corvo 900,000 USD.")}).json()
    ev = [{"record_type": "evidence", "record_id": evidence["id"], "quote": "$2.4M wired to Bellcastle."}]

    def cites(quote: str) -> list[dict]:
        # Each visual item cites the words that state its own value or date.
        return [{"record_type": "evidence", "record_id": evidence["id"], "quote": quote}]
    payload = {
        "title": "The Bellcastle fee",
        "goal_restated": "Explain the fee.",
        "executive_takeaway": "Bellcastle got the largest fee.",
        "sections": [
            {"heading": "Headline", "body": "The fee was $2.4M.", "key_points": [{"text": "$2.4M", "evidence": ev}], "evidence": ev,
             "visual": {"type": "figures", "items": [{"value": "$2.4M", "label": "to Bellcastle", "evidence": ev},
                                                     {"value": "11 days", "label": "after signing", "evidence": cites("11 days after signing")}]}},
            {"heading": "Sequence", "body": "Signed, then wired.", "key_points": [{"text": "Signed, then wired", "evidence": ev}], "evidence": ev,
             "visual": {"type": "timeline", "items": [{"date": "2023-11-03", "label": "Signed", "evidence": cites("Agreement signed on 2023-11-03")},
                                                      {"date": "2023-11-14", "label": "Wired", "evidence": cites("paid on 2023-11-14")}]}},
            {"heading": "Compared", "body": "Largest of three.", "key_points": [{"text": "Largest fee", "evidence": ev}], "evidence": ev,
             "visual": {"type": "bar", "title": "Consulting fees", "unit": "USD", "items": [
                 {"label": "Arden", "value": 285000, "evidence": cites("Arden 285,000 USD")},
                 {"label": "Bellcastle", "value": 2400000, "evidence": cites("Bellcastle 2,400,000 USD")},
                 {"label": "Corvo", "value": 900000, "evidence": cites("Corvo 900,000 USD")},
             ]}},
        ],
        "open_questions": ["Who owns Bellcastle?"],
    }
    monkeypatch.setattr(reasoning_module, "get_llm_client", lambda *a, **k: _FakeLLMClient(json.dumps(payload)))
    r = client.post(f"/api/investigations/{inv['id']}/assistant/presentation-draft", json={
        "goal": "Explain the fee.",
        "record_refs": [{"record_type": "claim", "record_id": claim["id"]}, {"record_type": "evidence", "record_id": evidence["id"]}],
    })
    assert r.status_code == 200, r.text
    candidate_id = r.json()["candidate_id"]
    client.post(f"/api/ai-analysis-candidates/{candidate_id}/review", json={"decision": "accept"})
    return {"candidate_id": candidate_id, "evidence_id": evidence["id"]}


def _export(deck_id: str):
    r = client.get(f"/api/decks/{deck_id}/export.pptx")
    assert r.status_code == 200, r.text
    return Presentation(BytesIO(r.content))


def _chart(prs):
    return next(sh.chart for s in prs.slides for sh in s.shapes if sh.has_chart)


def _background(slide) -> str:
    return str(slide.background.fill.fore_color.rgb)


def _font_names(prs) -> set[str]:
    return {run.font.name for s in prs.slides for sh in s.shapes if sh.has_text_frame
            for p in sh.text_frame.paragraphs for run in p.runs if run.font.name}


# --- catalogue -----------------------------------------------------------------

def test_catalogue_lists_themes_fonts_presets_and_options():
    r = client.get("/api/presentation-styles")
    assert r.status_code == 200
    data = r.json()
    assert data["defaults"] == {"theme": "dark_academia", "fonts": "theme"}
    assert {t["id"] for t in data["themes"]} == {"dark_academia", "newsprint", "ledger", "high_contrast"}
    assert len([f for f in data["font_pairs"] if f["id"] != "theme"]) == 8
    for kind in ("bar", "timeline", "figures"):
        assert len(data["visual_presets"][kind]) >= 3
        assert data["visual_options"][kind]
    # The refused options have no control at all.
    bar_options = set(data["visual_options"]["bar"])
    assert not bar_options & {"axis_min", "minimum", "three_d", "dual_axis", "secondary_axis"}


@pytest.mark.parametrize("theme_id", sorted(ps.THEMES))
def test_every_theme_meets_wcag_text_contrast(theme_id):
    t = ps.THEMES[theme_id]
    pairs = {
        "text on page": (t.text, t.page), "text on panel": (t.text, t.panel),
        "muted on page": (t.muted, t.page), "muted on panel": (t.muted, t.panel),
        "accent text on page": (t.accent_text, t.page), "alert on page": (t.alert, t.page), "alert on panel": (t.alert, t.panel),
        "cover text": (t.cover_text, t.cover_bg), "cover muted": (t.cover_muted, t.cover_bg), "cover accent": (t.cover_accent, t.cover_bg),
    }
    failing = {name: round(ps.contrast_ratio(fg, bg), 2) for name, (fg, bg) in pairs.items() if ps.contrast_ratio(fg, bg) < 4.5}
    assert failing == {}


# --- validation ------------------------------------------------------------------

def test_visual_style_validation_drops_unknown_values_and_foreign_labels():
    style = ps.normalize_visual_style("bar", {
        "preset": "nonexistent",
        "options": {"sort": "sideways", "value_labels": "yes", "thickness": "thick",
                    "highlight": ["Bellcastle", "Not a bar"], "axis_min": 1000},
    }, ["Arden", "Bellcastle"])
    assert style == {"preset": "horizontal", "options": {"thickness": "thick", "highlight": ["Bellcastle"]}}
    assert ps.normalize_visual_style("pie", {"preset": "donut"}, []) is None
    assert ps.normalize_deck_style({"theme": "neon", "fonts": "comic"}) == {"theme": "dark_academia", "fonts": "theme"}


# --- promotion and export ----------------------------------------------------------

def test_promote_with_styles_saves_them_and_the_export_follows(monkeypatch):
    seeded = _accepted_candidate(monkeypatch)
    deck = client.post(f"/api/ai-analysis-candidates/{seeded['candidate_id']}/promote-to-deck", json={
        "style": {"theme": "newsprint", "fonts": "modern"},
        "visual_styles": {
            "0": {"preset": "row", "options": {"value_color": "accent"}},
            "1": {"preset": "list"},
            "2": {"preset": "column", "options": {"highlight": ["Bellcastle"]}},
        },
    }).json()
    assert deck["style"] == {"theme": "newsprint", "fonts": "modern"}
    content = [s for s in deck["slides"] if s["kind"] == "content"]
    assert [s["visual"]["style"]["preset"] for s in content] == ["row", "list", "column"]

    prs = _export(deck["id"])
    slides = list(prs.slides)
    theme = ps.THEMES["newsprint"]
    assert _background(slides[0]) == theme.cover_bg
    assert _background(slides[2]) == theme.page
    fonts = _font_names(prs)
    assert {"Newsreader", "Inter"} <= fonts and "Georgia" not in fonts

    from pptx.enum.chart import XL_CHART_TYPE
    chart = _chart(prs)
    assert chart.chart_type == XL_CHART_TYPE.COLUMN_CLUSTERED
    categories = list(chart.plots[0].categories)
    series = chart.plots[0].series[0]
    assert categories == ["Arden", "Bellcastle", "Corvo"]  # column keeps the drafted order
    fills = [str(series.points[i].format.fill.fore_color.rgb) for i in range(3)]
    assert fills == [theme.bar_muted, theme.accent, theme.bar_muted]


def test_a_headline_figures_number_size_sets_how_large_its_number_is_written(monkeypatch):
    """The design doc's options for headline figures: number size, card fill,
    accent edge. A smaller size starts lower down the same ladder."""
    seeded = _accepted_candidate(monkeypatch)
    sizes = {}
    for size in ("large", "medium", "small"):
        deck = client.post(f"/api/ai-analysis-candidates/{seeded['candidate_id']}/promote-to-deck", json={
            "visual_styles": {"0": {"preset": "cards", "options": {"number_size": size}}}}).json()
        # The figure's number is the run that is the value alone (the slide's headline and bullets say more).
        runs = [run for slide in _export(deck["id"]).slides for shape in slide.shapes if shape.has_text_frame
                for p in shape.text_frame.paragraphs for run in p.runs if run.text == "$2.4M"]
        sizes[size] = runs[0].font.size.pt
    assert sizes["large"] > sizes["medium"] > sizes["small"]
    assert sizes["large"] in ps.FIGURE_NUMBER_SIZES["large"] and sizes["small"] in ps.FIGURE_NUMBER_SIZES["small"]


def test_presets_and_options_never_change_values_or_start_bars_above_zero(monkeypatch):
    seeded = _accepted_candidate(monkeypatch)
    seen = []
    for preset in ("horizontal", "ranked", "column", "minimal"):
        for options in ({}, {"value_labels": False}, {"gridlines": True, "sort": "asc"}):
            deck = client.post(f"/api/ai-analysis-candidates/{seeded['candidate_id']}/promote-to-deck", json={
                "visual_styles": {"2": {"preset": preset, "options": options}},
            }).json()
            chart = _chart(_export(deck["id"]))
            pairs = dict(zip(chart.plots[0].categories, chart.plots[0].series[0].values))
            seen.append(pairs)
            assert chart.value_axis.minimum_scale == 0
            if options.get("value_labels") is False:
                # No labels on the bars: the value axis must stay so values can be read.
                assert chart.value_axis.visible and chart.value_axis.has_major_gridlines
    assert all(p == {"Arden": 285000, "Bellcastle": 2400000, "Corvo": 900000} for p in seen)


def test_ranked_preset_sorts_largest_first(monkeypatch):
    seeded = _accepted_candidate(monkeypatch)
    deck = client.post(f"/api/ai-analysis-candidates/{seeded['candidate_id']}/promote-to-deck", json={
        "visual_styles": {"2": {"preset": "ranked"}},
    }).json()
    chart = _chart(_export(deck["id"]))
    # Horizontal bars draw bottom-up, so the top bar is the last category.
    assert list(reversed(list(chart.plots[0].categories))) == ["Bellcastle", "Corvo", "Arden"]


def test_unstyled_deck_keeps_the_v2_look(monkeypatch):
    seeded = _accepted_candidate(monkeypatch)
    deck = client.post(f"/api/ai-analysis-candidates/{seeded['candidate_id']}/promote-to-deck", json={}).json()
    assert deck["style"] == {"theme": "dark_academia", "fonts": "theme"}
    prs = _export(deck["id"])
    assert _background(list(prs.slides)[0]) == "141B1F"
    assert {"Georgia", "Calibri"} <= _font_names(prs)
    chart = _chart(prs)
    assert list(reversed(list(chart.plots[0].categories))) == ["Arden", "Bellcastle", "Corvo"]
    text = "\n".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    assert "3 Nov 2023" in text and "11 days" in text


def test_every_preset_exports_with_citations_intact(monkeypatch):
    seeded = _accepted_candidate(monkeypatch)
    for figures in ("cards", "big_numbers", "row"):
        for timeline in ("track", "list", "compact"):
            deck = client.post(f"/api/ai-analysis-candidates/{seeded['candidate_id']}/promote-to-deck", json={
                "style": {"theme": "high_contrast"},
                "visual_styles": {"0": {"preset": figures}, "1": {"preset": timeline, "options": {"marker": "dot", "show_line": False}}},
            }).json()
            prs = _export(deck["id"])
            text = "\n".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
            assert "$2.4M" in text and "11 days" in text and "14 Nov 2023" in text
            assert "[1]" in text and "row 41" in text  # numbered citation and its source entry


def test_patch_deck_style_is_validated(monkeypatch):
    seeded = _accepted_candidate(monkeypatch)
    deck = client.post(f"/api/ai-analysis-candidates/{seeded['candidate_id']}/promote-to-deck", json={}).json()
    r = client.patch(f"/api/decks/{deck['id']}", json={"style_json": {"theme": "ledger", "fonts": "nope"}})
    assert r.status_code == 200
    assert r.json()["style"] == {"theme": "ledger", "fonts": "theme"}
    prs = _export(deck["id"])
    assert _background(list(prs.slides)[0]) == ps.THEMES["ledger"].cover_bg
    assert {"IBM Plex Serif", "IBM Plex Sans"} <= _font_names(prs)
