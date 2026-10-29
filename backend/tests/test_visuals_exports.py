"""A chart visual in the exports: a native, editable PowerPoint chart in a
deck; in a Word write-up the image from the export bridge, or the table
and a line saying the image was not available (ADR-0005, line 2). Every
value keeps its record citations either way."""

import copy
from collections import Counter
from io import BytesIO

import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Inches

import app.services.writeups as writeups_module
from app.ai.chart_check import resolve_chart
from app.db.session import SessionLocal
from app.main import app
from app.models.domain import Deck
from app.services.assistant_context import describe_record_refs
from app.services.decks import chart_record_refs, deck_citation_refs, normalize_visual, picture_size, render_deck_pptx
from app.services.demo_seed import seed_demo_investigation
from app.services.writeups import create_writeup, extract_cited_refs, extract_plain_text, render_writeup_docx
from app.visuals.render_client import RenderUnavailable

client = TestClient(app)

TINY_PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4890000000d4944415478da63f8cfc0f01f0005020203"
    "f1a2b9780000000049454e44ae426082"
)


@pytest.fixture()
def demo():
    with SessionLocal() as db:
        inv, _ = seed_demo_investigation(db)
        chart, record = resolve_chart(db, inv.id, {"type": "chart", "spec": {
            "kind": "bar", "title": "Meridian received all of the money", "query": {"name": "payments_by_party", "params": {"direction": "received"}},
        }})
        assert chart is not None, record
        yield db, inv.id, chart


def test_resolved_chart_carries_its_data_report_and_words(demo):
    _db, _inv, chart = demo
    assert chart["type"] == "chart" and chart["spec"]["kind"] == "bar"
    assert chart["data"]["rows"][0]["label"] == "Meridian Holdings Ltd" and chart["data"]["rows"][0]["value"] == 2400000.0
    assert chart["data"]["rows"][0]["record_refs"][0]["record_type"] == "relationship"
    assert chart["report"]["ok"] is True and 20 in {w["rule"] for w in chart["report"]["warnings"]}
    assert chart["text"]["alt_text"].startswith("Bar chart: Meridian received all of the money")
    assert chart["text"]["data_table"]["rows"][0]["record_count"] == 1


def test_normalize_visual_keeps_a_resolved_chart_and_drops_a_bare_order(demo):
    _db, _inv, chart = demo
    # A chart keeps its style as the catalogue allows it (app/visuals/style.py);
    # an unknown preset or option is dropped, and nothing else changes.
    kept = normalize_visual({**chart, "style": {"preset": "ranked", "options": {"thickness": "thin", "start_above_zero": True}}})
    assert kept is not None and kept["type"] == "chart" and kept["style"] == {"preset": "ranked", "options": {"thickness": "thin"}}
    assert kept["data"] == chart["data"] and kept["spec"] == chart["spec"]
    assert "style" not in (normalize_visual({**chart, "style": {"preset": "three_d"}}) or {})
    assert normalize_visual({"type": "chart", "spec": chart["spec"]}) is None
    assert normalize_visual({"type": "chart", "spec": {"kind": "bar"}, "data": chart["data"]}) is None
    assert chart_record_refs(chart) == [{"record_type": "relationship", "record_id": chart["data"]["rows"][0]["record_refs"][0]["record_id"]}]


def test_deck_export_draws_a_native_chart_and_cites_its_records(demo):
    db, inv_id, chart = demo
    row = Deck(investigation_id=inv_id, title="Money", slides_json=[
        {"kind": "content", "heading": "Who got paid", "bullets": ["Meridian received the whole sum"], "evidence": [], "notes": "", "visual": chart},
    ])
    db.add(row)
    db.commit()
    db.refresh(row)
    refs = deck_citation_refs(row.slides_json)
    assert refs == chart_record_refs(chart)
    details = describe_record_refs(db, investigation_id=inv_id, refs=refs)
    prs = render_deck_pptx(row, details)
    slides = list(prs.slides)
    assert len(slides) == 3  # title, content, sources
    charts = [shape for shape in slides[1].shapes if shape.has_chart]
    assert len(charts) == 1
    categories = list(charts[0].chart.plots[0].categories)
    assert categories == ["Meridian Holdings Ltd · 1 payment †"]
    assert list(charts[0].chart.plots[0].series[0].values) == [2400000.0]
    text = "\n".join(shape.text_frame.text for slide in slides for shape in slide.shapes if shape.has_text_frame)
    assert "Meridian received all of the money" in text
    assert "Source: 1 payment in this investigation's records" in text
    assert "Sources  [1]" in text
    assert "Bellcastle Trading Group — Payment → Meridian Holdings Ltd" in text  # the relationship, by its caption

    exported = client.get(f"/api/decks/{row.id}/export.pptx")
    assert exported.status_code == 200
    assert sum(1 for s in Presentation(BytesIO(exported.content)).slides for sh in s.shapes if sh.has_chart) == 1


def test_deck_with_a_column_chart_exports_too(demo):
    db, inv_id, _chart = demo
    column, record = resolve_chart(db, inv_id, {"type": "chart", "spec": {
        "kind": "column", "title": "Events cluster in 2023 and 2024", "query": {"name": "events_over_time", "params": {"granularity": "year"}},
    }})
    assert column is not None, record
    row = Deck(investigation_id=inv_id, title="When", slides_json=[{"kind": "content", "heading": "Timing", "bullets": [], "evidence": [], "notes": "", "visual": column}])
    db.add(row)
    db.commit()
    db.refresh(row)
    prs = render_deck_pptx(row, describe_record_refs(db, investigation_id=inv_id, refs=deck_citation_refs(row.slides_json)))
    chart_shape = next(shape for shape in list(prs.slides)[1].shapes if shape.has_chart)
    assert list(chart_shape.chart.plots[0].categories) == ["2019 †", "2020", "2021 †", "2022", "2023 †", "2024"]
    assert list(chart_shape.chart.plots[0].series[0].values) == [1.0, 0.0, 1.0, 0.0, 1.0, 2.0]


def _writeup_with_chart(db, inv_id, chart):
    content = {"type": "doc", "content": [
        {"type": "heading", "attrs": {"level": 2}, "content": [{"type": "text", "text": "Who got paid"}]},
        {"type": "paragraph", "content": [{"type": "text", "text": "The whole sum went to Meridian."}]},
        {"type": "evidenceVisual", "attrs": {"visual": chart}},
    ]}
    return create_writeup(db, investigation_id=inv_id, title="Money brief", content_json=content)


def test_writeup_counts_chart_records_as_citations_and_lists_them_in_plain_text(demo):
    db, inv_id, chart = demo
    row = _writeup_with_chart(db, inv_id, chart)
    assert extract_cited_refs(row.content_json) == chart_record_refs(chart)
    assert row.cited_record_refs == chart_record_refs(chart)
    text = extract_plain_text(row.content_json)
    assert "Meridian received all of the money" in text and "2400000.0 — Meridian Holdings Ltd" in text


def test_writeup_export_falls_back_to_a_table_when_the_render_service_is_down(demo, monkeypatch):
    db, inv_id, chart = demo
    row = _writeup_with_chart(db, inv_id, chart)

    def unreachable(*args, **kwargs):
        raise RenderUnavailable("render service at http://render:8070 could not be reached: ConnectError")

    monkeypatch.setattr("app.visuals.render_client.render_png", unreachable)
    docx = render_writeup_docx(row, describe_record_refs(db, investigation_id=inv_id, refs=row.cited_record_refs))
    paragraphs = [p.text for p in docx.paragraphs]
    assert any("Chart image not available" in p and "could not be reached" in p for p in paragraphs)
    assert not docx.inline_shapes
    cells = "\n".join(c.text for t in docx.tables for r in t.rows for c in r.cells)
    assert "Meridian Holdings Ltd · 1 payment †" in cells and "$2.4M" in cells and "[1]" in cells
    assert any(p.startswith("Source: 1 payment") for p in paragraphs)
    assert "Sources" in paragraphs


def test_writeup_export_embeds_the_png_when_the_bridge_answers(demo, monkeypatch):
    db, inv_id, chart = demo
    row = _writeup_with_chart(db, inv_id, chart)
    calls = {}

    def fake_render(option, width, height, **kwargs):
        calls["option"] = option
        return TINY_PNG

    monkeypatch.setattr("app.visuals.render_client.render_png", fake_render)
    docx = render_writeup_docx(row, describe_record_refs(db, investigation_id=inv_id, refs=row.cited_record_refs))
    assert calls["option"]["title"][0]["text"] == "Meridian received all of the money"
    assert len(docx.inline_shapes) == 1
    paragraphs = [p.text for p in docx.paragraphs]
    assert not any("not available" in p for p in paragraphs)
    # The image carries the title and source line, so the document does not
    # repeat them; its description goes in Word's Alt Text field.
    assert not any("Meridian received all of the money" in p for p in paragraphs)
    assert not any(p.startswith("Source:") for p in paragraphs)
    assert docx.inline_shapes[0]._inline.docPr.get("descr").startswith("Bar chart: Meridian received all of the money")
    cells = "\n".join(c.text for t in docx.tables for r in t.rows for c in r.cells)
    assert "Meridian Holdings Ltd · 1 payment †" in cells and "$2.4M" in cells and "[1]" in cells

    exported = client.get(f"/api/writeups/{row.id}/export.docx")
    assert exported.status_code == 200
    assert len(DocxDocument(BytesIO(exported.content)).inline_shapes) == 1


# --- every statistical kind (pass 3) --------------------------------------------------

EVERY_KIND = [
    ("line", {"title": "Payments fell after the November transfer", "query": {"name": "payments_over_time"}}),
    ("area", {"title": "Payments fell after the November transfer", "query": {"name": "payments_over_time"}}),
    ("stacked_bar_100", {"title": "Meridian received most of the money", "query": {"name": "payments_by_party"}}),
    ("donut", {"title": "Meridian received most of the money", "query": {"name": "payments_by_party"}, "limit": 3}),
    # A treemap needs parts big enough to carry their names (rule 12): the
    # payments are too lopsided for one, the kinds of link are not.
    ("treemap", {"title": "Payments are most of the links", "query": {"name": "links_by_type"}}),
    ("histogram", {"title": "Most payments were small", "query": {"name": "payment_amounts"}}),
    ("scatter", {"title": "Meridian took one large payment", "query": {"name": "payments_count_and_total_by_party"}}),
]


@pytest.fixture()
def every_kind(demo):
    """One resolved chart of every statistical kind, cooked from the demo
    investigation with more payments (amounts to bin, parties to plot)."""
    from app.models.domain import Entity
    from app.services.relationships import create_relationship

    db, inv_id, _chart = demo
    ids = {e.caption: e.id for e in db.query(Entity).filter_by(investigation_id=inv_id).all()}
    for payer, payee, amount, date in (
        ("Meridian Holdings Ltd", "Harborline Civic Trust", "5000", "2023-12-02"), ("Meridian Holdings Ltd", "Harborline Civic Trust", "7000", "2024-01-15"),
        ("Meridian Holdings Ltd", "Alaric Voss", "45000", "2024-01-20"), ("Bellcastle Trading Group", "Alaric Voss", "250000", "2024-02-03"),
        ("Meridian Holdings Ltd", "Priya Kavanagh", "12000", "2024-02-11"),
    ):
        create_relationship(db, investigation_id=inv_id, schema="Payment", source_entity_id=ids[payer], target_entity_id=ids[payee],
                            properties={"amount": [amount], "currency": ["USD"], "date": [date]})
    charts = {}
    for kind, order in EVERY_KIND:
        chart, record = resolve_chart(db, inv_id, {"type": "chart", "spec": {"kind": kind, **order}})
        assert chart is not None, (kind, record)
        charts[kind] = chart
    return db, inv_id, charts


def _deck(db, inv_id, charts) -> Deck:
    row = Deck(investigation_id=inv_id, title="Every kind", slides_json=[
        {"kind": "content", "heading": f"A {kind} chart", "bullets": [], "evidence": [], "notes": "Section prose.", "visual": chart}
        for kind, chart in charts.items()
    ])
    db.add(row)
    db.commit()
    db.refresh(row)
    return row


def _alt_text(shape) -> str:
    return shape._element.xpath("./p:nvGraphicFramePr/p:cNvPr | ./p:nvPicPr/p:cNvPr")[0].get("descr")


def test_deck_draws_every_kind_natively_and_the_treemap_as_the_rendered_picture(every_kind, monkeypatch):
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    db, inv_id, charts = every_kind
    calls = []
    drawn = {}

    def fake_render(option, width, height, **kwargs):
        calls.append(option["series"][0]["type"])
        drawn.update(option=option, size=(width, height), scale=kwargs.get("scale"))
        return TINY_PNG

    monkeypatch.setattr("app.visuals.render_client.render_png", fake_render)
    row = _deck(db, inv_id, charts)
    prs = render_deck_pptx(row, describe_record_refs(db, investigation_id=inv_id, refs=deck_citation_refs(row.slides_json)))
    slides = list(prs.slides)
    by_kind = dict(zip(charts, slides[1:1 + len(charts)]))
    expected_types = {"line": XL_CHART_TYPE.LINE_MARKERS, "area": XL_CHART_TYPE.AREA, "stacked_bar_100": XL_CHART_TYPE.BAR_STACKED_100,
                      "donut": XL_CHART_TYPE.DOUGHNUT, "histogram": XL_CHART_TYPE.COLUMN_CLUSTERED, "scatter": XL_CHART_TYPE.XY_SCATTER}
    for kind, chart_type in expected_types.items():
        frames = [shape for shape in by_kind[kind].shapes if shape.has_chart]
        assert len(frames) == 1, kind
        assert frames[0].chart.chart_type == chart_type, kind
        assert _alt_text(frames[0]) == charts[kind]["text"]["alt_text"], kind
    # The treemap is the picture the render service drew, for the slide: at a
    # design size where its 12 px text reads at 10.5 pt on the slide (these
    # slides have no bullets, so the chart has the full width), filling the
    # chart's space, its title and source line in the picture.
    assert calls == ["treemap"]
    pictures = [shape for shape in by_kind["treemap"].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pictures) == 1 and _alt_text(pictures[0]) == charts["treemap"]["text"]["alt_text"]
    assert drawn["size"] == picture_size(11.8, 4.7) == (960, 382) and drawn["scale"] == 3.0
    assert (pictures[0].width, pictures[0].height) == (Inches(11.8), Inches(4.7))
    titles = {t["id"]: t["text"] for t in drawn["option"]["title"]}
    assert titles["title"] == charts["treemap"]["spec"]["title"] and titles["source"].startswith("Source: ")
    slide_text = "\n".join(shape.text_frame.text for shape in by_kind["treemap"].shapes if shape.has_text_frame)
    assert titles["title"] not in slide_text and "Source: " not in slide_text   # not written twice

    # The values drawn are the rows the data service cooked.
    line = next(s for s in by_kind["line"].shapes if s.has_chart).chart
    assert list(line.plots[0].series[0].values) == [r["value"] for r in charts["line"]["data"]["rows"]]
    whole = next(s for s in by_kind["stacked_bar_100"].shapes if s.has_chart).chart
    assert [s.name for s in whole.plots[0].series][0] == "Meridian Holdings Ltd"
    assert whole.plots[0].series[0].points[0].data_label.text_frame.text.startswith("Meridian Holdings Ltd")
    histogram = next(s for s in by_kind["histogram"].shapes if s.has_chart).chart
    assert histogram.plots[0].gap_width == 0   # touching bins
    assert sum(histogram.plots[0].series[0].values) == len(charts["histogram"]["data"]["rows"])
    scatter = next(s for s in by_kind["scatter"].shapes if s.has_chart).chart
    series = scatter.plots[0].series[0]
    xs = [float(v.text) for v in series._element.xpath("./c:xVal//c:pt/c:v")]
    rows = {r["label"]: r for r in charts["scatter"]["data"]["rows"]}
    assert sorted(zip(xs, series.values)) == sorted((r["x"], r["value"]) for r in rows.values())
    assert scatter.category_axis.axis_title.text_frame.text == "Number of payments received"

    # Under every native chart its source line; in the notes its description; in the footer its records.
    for kind, slide in by_kind.items():
        text = "\n".join(shape.text_frame.text for shape in slide.shapes if shape.has_text_frame)
        assert ("Source: " in text) is (kind != "treemap") and "Sources  [" in text, kind
        assert slide.notes_slide.notes_text_frame.text.startswith("Section prose.\n\nChart: "), kind


def test_a_treemap_falls_back_to_native_bars_when_the_render_service_is_down(every_kind, monkeypatch):
    db, inv_id, charts = every_kind

    def unreachable(*args, **kwargs):
        raise RenderUnavailable("render service at http://render:8070 could not be reached: ConnectError")

    monkeypatch.setattr("app.visuals.render_client.render_png", unreachable)
    row = _deck(db, inv_id, {"treemap": charts["treemap"]})
    slide = list(render_deck_pptx(row, []).slides)[1]
    frames = [shape for shape in slide.shapes if shape.has_chart]
    largest = max(charts["treemap"]["data"]["rows"], key=lambda r: r["value"])["label"]
    # Largest first, read top-down (PowerPoint lists a bar chart's categories bottom-up).
    assert len(frames) == 1 and list(frames[0].chart.plots[0].categories)[-1].startswith(largest)
    text = "\n".join(shape.text_frame.text for shape in slide.shapes if shape.has_text_frame)
    assert "Treemap image not available" in text and "the same parts are shown as bars" in text


def _stored_parts(kind: str, values: list[float]) -> dict:
    """A stored chart of four parties' payments, written out by hand so its
    shape is fixed: two payments behind each part."""
    names = ["Meridian Holdings Ltd", "Alaric Voss", "Harborline Civic Trust", "Orrin Ltd"]
    rows = [{"key": name, "label": name, "value": value, "record_count": 2,
             "record_refs": [{"record_type": "relationship", "record_id": f"{name}-{k}"} for k in range(2)]}
            for name, value in zip(names, values)]
    return {"type": "chart",
            "spec": {"kind": kind, "title": "Meridian received most of the money",
                     "query": {"name": "payments_by_party", "params": {"direction": "received"}}},
            "data": {"query": "payments_by_party", "params": {"direction": "received"}, "investigation_id": "inv", "series_kind": "category",
                     "unit_kind": "amount", "unit": "USD", "currency": "USD", "record_noun": "payments", "record_count": 8, "omitted": [], "rows": rows},
            "report": {"ok": True, "errors": [], "warnings": []}, "text": {}}


def test_a_treemap_that_cannot_name_its_parts_at_slide_size_is_shown_as_bars(demo, monkeypatch):
    """Drawn at 960 x 540 and shrunk into a slide, a treemap's text came out
    at 5.5 pt; drawn at the slide's scale, its small parts have no room for
    their names. The deck says so and shows the parts as bars, rather than
    a picture whose names no one in the room can read."""
    from app.services.decks import checked_chart
    from app.services.presentation_style import resolved_theme

    db, inv_id, _chart = demo
    lopsided = _stored_parts("treemap", [2_400_000, 355_000, 192_000, 107_000])
    theme = resolved_theme(None)
    assert checked_chart(lopsided, theme).report.ok   # it passes at the app's design size ...
    assert not checked_chart(lopsided, theme, width=960, height=382).report.ok   # ... but not at the slide's
    monkeypatch.setattr("app.visuals.render_client.render_png", lambda *args, **kwargs: pytest.fail("no picture should be drawn"))
    row = _deck(db, inv_id, {"treemap": lopsided})
    slide = list(render_deck_pptx(row, []).slides)[1]
    frames = [shape for shape in slide.shapes if shape.has_chart]
    assert len(frames) == 1 and list(frames[0].chart.plots[0].categories)[-1].startswith("Meridian Holdings Ltd")
    text = "\n".join(shape.text_frame.text for shape in slide.shapes if shape.has_text_frame)
    assert "Treemap not drawn at this size (rule 12: 'Harborline Civic Trust' is too small in a treemap to show its name" in text
    assert "The same parts are shown as bars." in text and "Source: " in text


def test_a_saved_deck_keeps_what_its_treemap_slide_shows(every_kind, monkeypatch):
    """The slide preview draws what the export draws: a saved deck keeps,
    with each treemap, the drawing at the slide's own scale -- or, when its
    parts cannot be named there, the parts as bars with the export's own
    words for why."""
    from app.services.decks import chart_region, picture_size, serialize_deck, update_deck

    db, inv_id, charts = every_kind
    drawn = {}
    monkeypatch.setattr("app.visuals.render_client.render_png",
                        lambda option, width, height, **kwargs: drawn.update(option=option, size=(width, height)) or TINY_PNG)
    row = _deck(db, inv_id, {"treemap": charts["treemap"], "donut": charts["donut"]})
    row = update_deck(db, row, {"slides_json": row.slides_json})
    slides = serialize_deck(row)["slides"]
    treemap, donut = slides[0]["visual"], slides[1]["visual"]
    assert "slide_drawing" not in donut                      # native charts are drawn by PowerPoint itself
    full_width = picture_size(*chart_region(False)[2:])
    assert (treemap["slide_drawing"]["width"], treemap["slide_drawing"]["height"]) == full_width
    assert "note" not in treemap["slide_drawing"]
    render_deck_pptx(row, [])
    assert drawn["size"] == full_width and drawn["option"] == treemap["slide_drawing"]["option"]   # what the export drew

    lopsided = _stored_parts("treemap", [2_400_000, 355_000, 192_000, 107_000])
    row = _deck(db, inv_id, {"treemap": lopsided})
    row = update_deck(db, row, {"slides_json": row.slides_json})
    drawing = serialize_deck(row)["slides"][0]["visual"]["slide_drawing"]
    assert drawing["note"].startswith("Treemap not drawn at this size (rule 12: 'Harborline Civic Trust' is too small")
    assert drawing["option"]["series"][0]["type"] == "bar"
    # Beside bullets the region is narrower, and so is the drawing.
    narrow = update_deck(db, row, {"slides_json": [{**row.slides_json[0], "bullets": ["One point"]}]})
    assert serialize_deck(narrow)["slides"][0]["visual"]["slide_drawing"]["width"] == picture_size(*chart_region(True)[2:])[0]


def test_picture_size_reads_at_ten_and_a_half_points_on_the_slide():
    assert picture_size(6.1, 4.7) == (502, 387)    # beside the bullets: 12 px is 10.5 pt across 6.1 in
    assert picture_size(11.8, 4.7) == (960, 382)   # the full width: never wider than the app's design


def _overlap(a: float, a_w: float, b: float, b_w: float) -> bool:
    return abs(a - b) < (a_w + b_w) / 2 - 1e-6


@pytest.mark.parametrize("width", [6.1 * 72, 11.8 * 72])
@pytest.mark.parametrize("values", [[2_400_000, 355_000, 192_000, 107_000], [40, 30, 20, 10], [30, 25, 20, 15, 10], [1, 1]])
def test_a_whole_bar_names_every_part_over_or_under_it_without_collisions(values, width):
    from app.services.decks import NATIVE_LABEL_GAP, whole_bar_layout

    labels = [f"Part number {i} with a long name\n{v} · 2 payments" for i, v in enumerate(values)]
    height, thickness = 3.72 * 72, 34.0
    layout = whole_bar_layout(labels, values, width, height, thickness)
    assert layout is not None
    size, bar_top, placed = layout
    total, before = sum(values), 0.0
    for label, value in zip(placed, values):
        start, end = before / total * width, (before + value) / total * width
        before += value
        assert 0 <= label.x - label.width / 2 and label.x + label.width / 2 <= width + 1e-6   # inside the frame
        assert label.x + label.width / 2 >= start and label.x - label.width / 2 <= end        # over or under its own part
        assert label.y + label.height / 2 <= bar_top or label.y - label.height / 2 >= bar_top + thickness   # off the bar
        assert label.anchor_x == pytest.approx((start + end) / 2) and label.anchor_y == pytest.approx(bar_top + thickness / 2)
    for i, a in enumerate(placed):
        for b in placed[i + 1:]:
            same_line = abs(a.y - b.y) < 1e-6
            assert not (same_line and _overlap(a.x, a.width + NATIVE_LABEL_GAP, b.x, b.width)), (a, b)
            assert not (_overlap(a.x, a.width, b.x, b.width) and _overlap(a.y, a.height, b.y, b.height)), (a, b)
    assert size in (11, 10, 9)


def test_a_whole_bar_too_crowded_for_its_names_has_no_layout():
    from app.services.decks import whole_bar_layout

    labels = [f"A party with a long name {i}\n12.5% · $1.2M · 12 payments" for i in range(8)]
    assert whole_bar_layout(labels, [1.0] * 8, 3 * 72, 3.72 * 72, 34.0) is None


@pytest.mark.parametrize("width", [6.1 * 72, 11.8 * 72])
@pytest.mark.parametrize("values", [[2_400_000, 355_000, 192_000, 107_000], [2_925_000, 129_000], [25, 25, 25, 25], [97, 1, 1, 1]])
def test_a_donut_labels_every_slice_outside_its_ring_on_its_own_side(values, width):
    import math

    from app.services.decks import DONUT_HOLE, DONUT_MIN_DIAMETER, donut_layout

    labels = [f"Party {i} Holdings Ltd\n{v} · 4 payments" for i, v in enumerate(values)]
    height = 3.72 * 72
    layout = donut_layout(labels, values, width, height)
    assert layout is not None
    size, diameter, centre, placed = layout
    assert diameter >= DONUT_MIN_DIAMETER and size in (11, 10, 9)
    radius, cy = diameter / 2, height / 2
    total, before = sum(values), 0.0
    for label, value in zip(placed, values):
        angle = 2 * math.pi * (before + value / 2) / total
        before += value
        assert 0 <= label.x - label.width / 2 and label.x + label.width / 2 <= width + 1e-6
        assert 0 <= label.y - label.height / 2 and label.y + label.height / 2 <= height + 1e-6
        # Clear of the ring, on the side of its slice.
        assert (label.x - label.width / 2 > centre + radius) if math.sin(angle) >= 0 else (label.x + label.width / 2 < centre - radius)
        # Moved from where PowerPoint writes it: the middle of its slice's band.
        middle = radius * (1 + DONUT_HOLE / 100) / 2
        assert (label.anchor_x, label.anchor_y) == pytest.approx((centre + middle * math.sin(angle), cy - middle * math.cos(angle)))
    for i, a in enumerate(placed):
        for b in placed[i + 1:]:
            assert not (_overlap(a.x, a.width, b.x, b.width) and _overlap(a.y, a.height, b.y, b.height)), (a, b)


def test_a_donut_with_no_room_for_its_labels_has_no_layout():
    from app.services.decks import donut_layout

    assert donut_layout(["A long party name\n50% · 2 payments", "Another long party name\n50% · 2 payments"], [1, 1], 3 * 72, 3 * 72) is None


def test_native_part_labels_are_moved_unwrapped_with_leader_lines(demo):
    """The 100% bar's and the donut's labels: moved by a manual layout from
    the middle of their marks to the places worked out for them, never
    wrapped, with a leader line back; the plot area pinned where the layout
    put it (the XML PowerPoint itself writes for a dragged label)."""
    from pptx.oxml.ns import qn

    from app.services.decks import whole_bar_layout

    db, inv_id, _chart = demo
    values = [2_400_000, 355_000, 192_000, 107_000]
    row = Deck(investigation_id=inv_id, title="Parts", slides_json=[
        {"kind": "content", "heading": f"A {kind}", "bullets": ["Meridian received most of the money"], "evidence": [], "notes": "",
         "visual": _stored_parts(kind, values)}
        for kind in ("stacked_bar_100", "donut")
    ])
    db.add(row)
    db.commit()
    slides = list(render_deck_pptx(row, []).slides)[1:3]
    whole, donut = (next(shape for shape in slide.shapes if shape.has_chart) for slide in slides)
    for frame in (whole, donut):
        chart = frame.chart
        layout = chart._chartSpace.find(qn("c:chart")).find(qn("c:plotArea")).find(qn("c:layout")).find(qn("c:manualLayout"))
        assert layout.find(qn("c:layoutTarget")).get("val") == "inner"
        dlbls = chart._chartSpace.xpath(".//c:dLbl")
        assert len(dlbls) == len(values)
        for dlbl in dlbls:
            assert dlbl[0].tag == qn("c:idx") and dlbl[1].tag == qn("c:layout") and dlbl[2].tag == qn("c:tx")   # CT_DLbl's order
            assert dlbl.find(qn("c:layout")).find(qn("c:manualLayout")).find(qn("c:x")) is not None
            assert dlbl.find(qn("c:tx")).find(qn("c:rich")).find(qn("a:bodyPr")).get("wrap") == "none"
        assert chart._chartSpace.xpath(".//c:dLbls/c:extLst/c:ext[@uri='{CE6537A1-D6FC-4f65-9D91-7224C49458BB}']")
    # Each 100% bar label moved by exactly what the layout worked out (fractions of the frame).
    width, height = whole.width / 914400 * 72, whole.height / 914400 * 72
    labels = [dlbl.find(qn("c:tx")).find(qn("c:rich")) for dlbl in whole.chart._chartSpace.xpath(".//c:dLbl")]
    texts = ["\n".join("".join(t.text for t in p.iter(qn("a:t"))) for p in rich.findall(qn("a:p"))) for rich in labels]
    _size, _top, placed = whole_bar_layout(texts, values, width, height, 34.0)
    moved = [float(x) for x in whole.chart._chartSpace.xpath(".//c:dLbl/c:layout/c:manualLayout/c:x/@val")]
    assert moved == pytest.approx([(p.x - p.anchor_x) / width for p in placed], abs=1e-4)


def test_a_line_writes_values_on_its_points_and_counts_under_its_periods(every_kind):
    db, inv_id, charts = every_kind
    styled = {**copy.deepcopy(charts["line"]), "style": {"preset": "line_values"}}
    row = _deck(db, inv_id, {"line": styled})
    chart = next(shape for shape in list(render_deck_pptx(row, []).slides)[1].shapes if shape.has_chart).chart
    rows = charts["line"]["data"]["rows"]
    categories = list(chart.plots[0].categories)
    assert all("\n" in label for label in categories)   # the period, then its count on a line of its own
    assert categories[0].split("\n")[0] == rows[0]["label"]
    points = chart.plots[0].series[0].points
    written = [points[i].data_label.text_frame.text for i in range(len(rows))]
    assert all(text.startswith("$") and "payment" not in text for text in written)   # the value alone


def test_a_chart_goes_into_the_deck_in_its_own_style_and_is_checked_again_when_saved(every_kind):
    from pptx.enum.chart import XL_CHART_TYPE

    from app.services.decks import create_deck, update_deck

    db, inv_id, charts = every_kind

    def content(visual):
        return {"kind": "content", "heading": "A styled chart", "bullets": [], "evidence": [], "notes": "", "visual": visual}

    area = {**copy.deepcopy(charts["line"]), "style": {"preset": "area"}}
    donut = {**copy.deepcopy(charts["stacked_bar_100"]), "style": {"preset": "donut"}}
    deck = create_deck(db, investigation_id=inv_id, title="Styled", slides=[content(area), content(donut)])
    kept = [slide["visual"] for slide in deck.slides_json]
    assert kept[0]["style"] == {"preset": "area", "options": {}} and "areaStyle" in kept[0]["option"]["series"][0]
    assert kept[1]["option"]["series"][0]["type"] == "pie" and kept[1]["report"]["ok"] is True
    # The rows are the ones the draft cooked; the order slip is the one the model wrote.
    assert kept[0]["data"] == charts["line"]["data"] and kept[0]["spec"] == charts["line"]["spec"]
    native = [shape.chart.chart_type for slide in render_deck_pptx(deck, []).slides for shape in slide.shapes if shape.has_chart]
    assert XL_CHART_TYPE.AREA in native and XL_CHART_TYPE.DOUGHNUT in native

    # A style that breaks a rule is kept as failing, and the export says so instead of drawing it:
    # the payments are too lopsided for a treemap to name every part (rule 12).
    broken = create_deck(db, investigation_id=inv_id, title="Broken", slides=[
        content({**copy.deepcopy(charts["stacked_bar_100"]), "style": {"preset": "treemap"}})])
    report = broken.slides_json[0]["visual"]["report"]
    assert report["ok"] is False and 12 in {e["rule"] for e in report["errors"]}
    slide = list(render_deck_pptx(broken, []).slides)[1]   # after the cover
    assert not any(shape.has_chart for shape in slide.shapes)
    assert "Chart not drawn: it fails the design checks (rule 12:" in "\n".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)

    # A new look for the deck: every chart is drawn again in it.
    newsprint = update_deck(db, deck, {"style_json": {"theme": "newsprint", "fonts": "theme"}})
    assert {slide["visual"]["option"]["backgroundColor"] for slide in newsprint.slides_json} == {"#FFFFFF"}


def test_a_stored_chart_that_now_fails_a_rule_is_not_drawn(every_kind):
    db, inv_id, charts = every_kind
    broken = copy.deepcopy(charts["donut"])
    broken["spec"]["limit"] = None
    # A fifth part (stored before a rule tightened, say): rule 5 refuses a donut past four slices.
    broken["data"]["rows"].append({"label": "Orrin Ltd", "value": 1000.0, "record_count": 1, "key": "Orrin Ltd",
                                   "record_refs": [{"record_type": "relationship", "record_id": "rel-orrin"}]})
    row = _deck(db, inv_id, {"donut": broken})
    slide = list(render_deck_pptx(row, []).slides)[1]
    assert not any(shape.has_chart for shape in slide.shapes)
    text = "\n".join(shape.text_frame.text for shape in slide.shapes if shape.has_text_frame)
    assert "Chart not drawn: it fails the design checks (rule 5:" in text


def test_writeup_has_the_png_and_a_data_table_for_every_kind(every_kind, monkeypatch):
    db, inv_id, charts = every_kind
    monkeypatch.setattr("app.visuals.render_client.render_png", lambda option, width, height, **kwargs: TINY_PNG)
    content = {"type": "doc", "content": [{"type": "evidenceVisual", "attrs": {"visual": chart}} for chart in charts.values()]}
    row = create_writeup(db, investigation_id=inv_id, title="Every kind", content_json=content)
    docx = render_writeup_docx(row, describe_record_refs(db, investigation_id=inv_id, refs=row.cited_record_refs))
    assert len(docx.inline_shapes) == len(charts)
    tables = {len(t.columns): t for t in docx.tables}
    assert set(tables) == {3, 4}   # three columns for one number, four for a share or a second number
    wide_rows = [[c.text for c in r.cells] for t in docx.tables if len(t.columns) == 4 for r in t.rows]
    assert any(cells[0].startswith("Alaric Voss") and cells[1] == "2" for cells in wide_rows)   # the scatter plot's x: two payments
    assert any(cells[2].endswith("%") for cells in wide_rows)   # a part's share of the whole
    assert all(cells[0] and cells[-1].isdigit() for cells in wide_rows)   # every row names its records' count
    alt_texts = [shape._inline.docPr.get("descr") for shape in docx.inline_shapes]
    paragraphs = [p.text for p in docx.paragraphs]
    for chart in charts.values():
        assert chart["text"]["alt_text"] in alt_texts
        assert chart["spec"]["title"] not in paragraphs and chart["text"]["source_line"] not in paragraphs


def test_writeup_tables_stand_in_for_every_kind_when_the_render_service_is_down(every_kind, monkeypatch):
    db, inv_id, charts = every_kind

    def unreachable(*args, **kwargs):
        raise RenderUnavailable("render service could not be reached")

    monkeypatch.setattr("app.visuals.render_client.render_png", unreachable)
    content = {"type": "doc", "content": [{"type": "evidenceVisual", "attrs": {"visual": chart}} for chart in charts.values()]}
    row = create_writeup(db, investigation_id=inv_id, title="Every kind", content_json=content)
    docx = render_writeup_docx(row, [])
    assert not docx.inline_shapes
    paragraphs = [p.text for p in docx.paragraphs]
    assert sum("Chart image not available" in p for p in paragraphs) == len(charts)
    # With no image, the title and source line are written out once each.
    # (Two fixture charts share a title, so count per chart, not per title.)
    titles = Counter(chart["spec"]["title"] for chart in charts.values())
    sources = Counter(chart["text"]["source_line"] for chart in charts.values())
    for title, times in titles.items():
        assert sum(p.startswith(title) for p in paragraphs) == times
    for source, times in sources.items():
        assert paragraphs.count(source) == times
    cells = [c.text for t in docx.tables for r in t.rows for c in r.cells]
    # A histogram's bars count payments: its table writes the count, never money.
    first_bin = charts["histogram"]["text"]["data_table"]["rows"][0]
    assert f"{first_bin['formatted']} payments" in cells and f"${first_bin['formatted']}" not in cells
    # The parts of a whole and the scatter plot keep every number in their data tables.
    assert any(text.endswith("%") for text in cells)


def test_the_old_three_visuals_still_export_as_before(demo):
    db, inv_id, _chart = demo
    ev = [{"record_type": "claim", "record_id": "clm-1"}]
    row = Deck(investigation_id=inv_id, title="Old visuals", slides_json=[
        {"kind": "content", "heading": "Figures", "bullets": ["a"], "evidence": [], "notes": "", "visual": {"type": "figures", "items": [{"value": "$2.4M", "label": "wired", "evidence": ev}]}},
        {"kind": "content", "heading": "Timeline", "bullets": [], "evidence": [], "notes": "", "visual": {"type": "timeline", "items": [
            {"date": "2023-11-03", "label": "Signed", "evidence": ev}, {"date": "2023-11-14", "label": "Wired", "evidence": ev}]}},
        {"kind": "content", "heading": "Bars", "bullets": [], "evidence": [], "notes": "", "visual": {"type": "bar", "title": "Fees", "unit": "USD", "items": [
            {"label": "Bellcastle", "value": 2400000, "evidence": ev}, {"label": "Arden", "value": 285000, "evidence": ev}]}},
    ])
    db.add(row)
    db.commit()
    db.refresh(row)
    prs = render_deck_pptx(row, [])
    slides = list(prs.slides)
    assert len(slides) == 5
    assert sum(1 for s in slides for sh in s.shapes if sh.has_chart) == 1
    assert writeups_module.EMPTY_DOC["type"] == "doc"
