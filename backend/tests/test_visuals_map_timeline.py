"""The map with its timeline (the design doc's "map linked to a timeline":
in the app a strip under the map; on a page two pictures side by side).
Each place carries the dated events at it (data_service.places_on_map);
the "Map with its timeline" preset draws them as a timeline beside the map,
each event naming its place, checked like any timeline, and in decks and
briefs as one picture of two, or the places and their events listed."""

from __future__ import annotations

import json

import httpx
import pytest
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.ai.chart_check import resolve_chart
from app.core.config import Settings
from app.db.session import SessionLocal
from app.models.domain import Deck, Investigation
from app.schemas.api import EntityCreate, TimelineEventCreate
from app.services.decks import render_deck_pptx
from app.services.entities import create_entity
from app.services.presentation_style import THEMES
from app.services.relationships import create_relationship
from app.services.timeline import create_timeline_event
from app.services.writeups import create_writeup, render_writeup_docx
from app.visuals.checker import check_visual
from app.visuals.data_service import VisualData, run_query
from app.visuals.echarts import BESIDE_LABEL_WIDTH, BESIDE_TITLE, TIMELINE_PLOT_MIN, ChartDoesNotFit, DrawOptions, build_chart, timeline_plot_room
from app.visuals.render_client import RenderUnavailable, render_png
from app.visuals.spec import QueryRef, parse_spec
from app.visuals.style import styled_chart

DARK = THEMES["dark_academia"]
TINY_PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4890000000d4944415478da63f8cfc0f01f0005020203"
    "f1a2b9780000000049454e44ae426082"
)
LINK = "Meridian registered office — Other link → Villa Kyrenia — startDate"   # the timeline's title for the link's start
SPEC = {"kind": "map", "question": "place", "title": "The money ended up in a villa in Cyprus", "query": {"name": "places_on_map"}}


def _world(*, events: bool = True) -> dict[str, str]:
    with SessionLocal() as db:
        inv = Investigation(name="Map and time")
        db.add(inv)
        db.commit()
        db.refresh(inv)

        def entity(schema: str, caption: str, **props: list[str]) -> str:
            return create_entity(db, EntityCreate(investigation_id=inv.id, schema=schema, caption=caption, properties=props)).id

        villa = entity("RealEstate", "Villa Kyrenia", latitude=["35.3364"], longitude=["33.3199"])
        office = entity("Address", "Meridian registered office", latitude=["51.5072"], longitude=["-0.1276"])
        if events:
            # An event about the villa, one about the office, and a link between the two that started in a month.
            create_timeline_event(db, TimelineEventCreate(investigation_id=inv.id, title="Villa searched", date_start="2023-11-14",
                                                          precision="day", entity_id=villa))
            create_timeline_event(db, TimelineEventCreate(investigation_id=inv.id, title="Office opened", date_start="2019-06",
                                                          precision="month", entity_id=office))
            create_relationship(db, investigation_id=inv.id, schema="UnknownLink", source_entity_id=office, target_entity_id=villa,
                                properties={"startDate": ["2021-02"]})
        return {"inv": inv.id, "villa": villa, "office": office}


def _data(ids: dict[str, str]) -> VisualData:
    with SessionLocal() as db:
        return run_query(db, ids["inv"], QueryRef(name="places_on_map"))


def _built(data: VisualData, **extra):
    return build_chart(parse_spec({**SPEC, **extra}), data, DARK, draw=DrawOptions(with_timeline=True))


# --- the data --------------------------------------------------------------------------

def test_each_place_carries_its_dated_events_each_cited_by_its_own_records():
    ids = _world()
    data = _data(ids)
    places = {row.label: row for row in data.rows}
    villa, office = places["Villa Kyrenia"], places["Meridian registered office"]
    assert [(e.label, e.start, e.status) for e in villa.events or []] == [
        (LINK, "2021-02", "asserted"),
        ("Villa searched", "2023-11-14", "asserted")]
    assert [e.label for e in office.events or []] == ["Office opened", LINK]
    searched = (villa.events or [])[1]
    assert [ref.record_type for ref in searched.record_refs] == ["timeline_event"]
    # The place itself stays one record, cited as itself.
    assert villa.record_count == 1 and [ref.record_id for ref in villa.record_refs] == [ids["villa"]]
    # Stored and read back with its events.
    again = VisualData.from_dict(json.loads(json.dumps(data.as_dict())))
    assert [e.label for e in again.rows[1].events or []] == [e.label for e in villa.events or []]


# --- the drawing -----------------------------------------------------------------------

def test_the_preset_draws_the_events_beside_the_map_each_naming_its_place():
    ids = _world()
    data = _data(ids)
    spec, draw = styled_chart(parse_spec(SPEC), data, {"preset": "map_timeline"})
    assert draw.with_timeline
    built = build_chart(spec, data, DARK, draw=draw)
    beside = built.option["beside"]
    assert beside["title"][0]["text"] == BESIDE_TITLE
    # The map takes 40% of the width and the timeline 60%: their titles wrap at their own widths.
    assert built.option["title"][0]["textStyle"]["width"] == 384 - 48 and beside["title"][0]["textStyle"]["width"] == 576 - 48
    # Each row is named in two lines, the event and its place; the link between the two places is drawn at each.
    assert [row.label for row in built.beside_rows] == [
        "Office opened\nMeridian registered office", f"{LINK}\nMeridian registered office", f"{LINK}\nVilla Kyrenia",
        "Villa searched\nVilla Kyrenia"]
    assert beside["yAxis"]["axisLabel"]["width"] == BESIDE_LABEL_WIDTH
    # Beside each mark, its date and status, as on any timeline.
    written = {item["name"]: (item.get("label") or {}).get("formatter") for series in beside["series"] for item in series["data"]
               if (item.get("label") or {}).get("show") is not False}
    assert written["Villa searched\nVilla Kyrenia"] == "14 Nov 2023 · asserted"
    assert written["Office opened\nMeridian registered office"] == "Jun 2019 (month only) · asserted"
    assert beside["title"][1]["text"] == "Source: 4 dated events at these places, from this investigation's timeline"
    # The words and the data table carry every event with its place (rule 13).
    assert "Beside it, a timeline of 4 dated events at these places, from Jun 2019 (month only) (Office opened, Meridian registered office)" in built.alt_text
    table = built.data_table["events"]
    assert table["columns"] == ["Event", "Place", "When", "Status", "Records"]
    assert [(r["label"], r["place"]) for r in table["rows"]][3] == ("Villa searched", "Villa Kyrenia")


def test_a_picked_place_picks_out_its_events_beside_the_map():
    data = _data(_world())
    built = _built(data, highlight=["Villa Kyrenia"])
    # The rest go grey: one colour and a highlight, which the check allows (rule 9).
    report = check_visual(parse_spec({**SPEC, "highlight": ["Villa Kyrenia"]}), data, DARK, built)
    assert report.ok, report.errors
    bold = {item["name"] for series in built.option["beside"]["series"] for item in series["data"]
            if (item.get("label") or {}).get("fontWeight") == "bold"}
    assert bold == {"Villa searched\nVilla Kyrenia", f"{LINK}\nVilla Kyrenia"}


def test_a_map_with_no_dated_event_at_its_places_has_no_timeline_to_put_beside_it():
    data = _data(_world(events=False))
    with pytest.raises(ChartDoesNotFit, match="no event on the timeline is dated at these places"):
        _built(data)
    assert "beside" not in build_chart(parse_spec(SPEC), data, DARK).option


@pytest.mark.parametrize("theme_id", sorted(THEMES))
def test_both_pictures_pass_the_check_in_every_theme(theme_id):
    data = _data(_world())
    spec = parse_spec({**SPEC, "theme": theme_id})
    built = build_chart(spec, data, THEMES[theme_id], draw=DrawOptions(with_timeline=True))
    report = check_visual(spec, data, THEMES[theme_id], built)
    assert report.ok, report.errors


def test_the_checker_reads_the_timeline_beside_the_map_back_too():
    data = _data(_world())
    spec = parse_spec(SPEC)
    built = _built(data)
    broken = json.loads(json.dumps(built.option))
    for series in broken["beside"]["series"]:
        for item in series["data"]:
            if item["name"] == "Villa searched\nVilla Kyrenia":
                item["value"][0] += 86_400_000              # a day late
                item["label"]["formatter"] = "14 Nov 2023"  # its status not written
    report = check_visual(spec, data, DARK, built.__class__(**{**built.__dict__, "option": broken}))
    messages = [e.message for e in report.errors]
    assert "the timeline beside the map: the dates drawn do not match the events' records for: Villa searched at Villa Kyrenia" in messages
    assert "the timeline beside the map: the date and status of 'Villa searched at Villa Kyrenia' are not written beside it" in messages


def test_a_timeline_squeezed_until_its_axis_has_no_room_is_refused():
    """Rule 14 across: the names on the left and the dates on the right must
    leave the time axis room for its first and last dates. Narrower,
    ECharts drew the axis backwards (seen with the render service)."""
    data = _data(_world())
    spec = parse_spec(SPEC)
    narrow = build_chart(spec, data, DARK, width=520, draw=DrawOptions(with_timeline=True))
    report = check_visual(spec, data, DARK, narrow)
    room = [e.message for e in report.errors if e.rule == 14]
    assert room and room[0].startswith("the timeline beside the map: the events' names and the dates written beside them leave the time axis")
    assert f"under the {TIMELINE_PLOT_MIN}px it needs" in room[0]
    assert timeline_plot_room(_built(data).option["beside"], 576) >= TIMELINE_PLOT_MIN


# --- the picture --------------------------------------------------------------------------

def test_the_two_pictures_are_drawn_their_own_ways_and_placed_side_by_side():
    built = _built(_data(_world()))
    asked: list[tuple[str, dict]] = []

    def answer(request: httpx.Request) -> httpx.Response:
        body = json.loads(request.content)
        asked.append((request.url.path, body))
        if request.url.path == "/map":
            return httpx.Response(200, headers={"content-type": "image/png"}, content=TINY_PNG)
        return httpx.Response(200, headers={"content-type": "image/svg+xml"}, text='<svg xmlns="http://www.w3.org/2000/svg" width="10" height="10"/>')

    settings = Settings(render_service_url="http://render.test", _env_file=None)
    assert render_png(built.option, 960, 540, settings=settings, transport=httpx.MockTransport(answer)).startswith(b"\x89PNG")
    paths = [path for path, _body in asked]
    assert paths == ["/map", "/echarts", "/echarts", "/echarts"]   # the map area, the map framed, the timeline, the page
    assert asked[0][1]["width"] == 384 and (asked[1][1]["width"], asked[2][1]["width"]) == (384, 576)
    assert "beside" not in asked[1][1]["option"] and asked[2][1]["option"]["title"][0]["text"] == BESIDE_TITLE
    page = asked[3][1]
    assert (page["width"], page["height"]) == (960, 540)
    assert [(g["left"], g["style"]["width"]) for g in page["option"]["graphic"]] == [(0, 384), (384, 576)]


# --- exports -----------------------------------------------------------------------------

@pytest.fixture()
def paired_chart():
    ids = _world()
    with SessionLocal() as db:
        chart, record = resolve_chart(db, ids["inv"], {"type": "chart", "spec": SPEC})
        assert chart is not None, record
        chart = {**chart, "style": {"preset": "map_timeline", "options": {}}}
        yield db, ids["inv"], chart


def test_a_deck_holds_one_picture_of_both_or_the_places_and_their_events_listed(paired_chart, monkeypatch):
    db, inv_id, chart = paired_chart
    drawn: list[bool] = []

    def fake_render(option, width, height, **kwargs):
        drawn.append("beside" in option)
        return TINY_PNG

    monkeypatch.setattr("app.visuals.render_client.render_png", fake_render)
    deck = Deck(investigation_id=inv_id, title="Where and when", slides_json=[
        {"kind": "content", "heading": "Where the money went, and when", "bullets": [], "evidence": [], "notes": "", "visual": chart}])
    db.add(deck)
    db.commit()
    slide = list(render_deck_pptx(deck, []).slides)[1]
    assert drawn == [True] and len([s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]) == 1

    def unreachable(*args, **kwargs):
        raise RenderUnavailable("render service at http://render:8070 could not be reached: ConnectError")

    monkeypatch.setattr("app.visuals.render_client.render_png", unreachable)
    slide = list(render_deck_pptx(deck, []).slides)[1]
    tables = [[[cell.text for cell in row.cells] for row in shape.table.rows] for shape in slide.shapes if shape.has_table]
    assert [table[0] for table in tables] == [["Place", "Where"], ["Event", "Place", "When"]]
    assert ["Villa searched", "Villa Kyrenia", "14 Nov 2023"] in tables[1]
    note = " ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
    assert "the places and the events at them are listed instead" in note


def test_a_brief_lists_the_events_with_their_places_after_the_places(paired_chart, monkeypatch):
    db, inv_id, chart = paired_chart
    monkeypatch.setattr("app.visuals.render_client.render_png", lambda *args, **kwargs: TINY_PNG)
    row = create_writeup(db, investigation_id=inv_id, title="Where and when", content_json={"type": "doc", "content": [
        {"type": "evidenceVisual", "attrs": {"visual": chart}}]})
    tables = [[[c.text for c in r.cells] for r in t.rows] for t in render_writeup_docx(row, []).tables]
    events = next(t for t in tables if any(r[0].startswith("Villa searched") for r in t))
    searched = next(r for r in events if r[0].startswith("Villa searched"))
    assert searched[1:] == ["Villa Kyrenia", "14 Nov 2023", "asserted"]
