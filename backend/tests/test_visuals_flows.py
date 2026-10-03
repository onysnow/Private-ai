"""The flow diagram (visuals pass 4, money flows; app/visuals/echarts.py
_sankey, data_service.payment_flows): who paid whom, each band as wide as
its money, every flow's amount and payments on it or in its tooltip, and
checked like every other chart."""

from __future__ import annotations

import copy

import pytest
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.ai.chart_check import resolve_chart
from app.db.session import SessionLocal
from app.models.domain import Deck, Entity
from app.services.decks import render_deck_pptx
from app.services.demo_seed import seed_demo_investigation
from app.services.presentation_style import THEMES
from app.services.relationships import create_relationship
from app.visuals.checker import MAX_FLOW_PARTIES, check_visual
from app.visuals.data_service import DataRow, RecordRef, VisualData
from app.visuals.echarts import (
    SINGLE_RECORD_MARK, ChartDoesNotFit, build_chart, flow_stages, flows_circle, sankey_layout, sankey_node_ids,
)
from app.visuals.render_client import RenderUnavailable
from app.visuals.spec import parse_spec
from app.visuals.style import normalize_chart_style, styled_chart

DARK = THEMES["dark_academia"]
TINY_PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4890000000d4944415478da63f8cfc0f01f0005020203"
    "f1a2b9780000000049454e44ae426082"
)


def _flow(source: str, target: str, value: float, payments: int) -> DataRow:
    label = f"{source} → {target}"
    return DataRow(label, value, [RecordRef("relationship", f"{source}-{target}-{i}") for i in range(payments)], payments, label,
                   source=source, target=target)


def _data(*rows: DataRow) -> VisualData:
    return VisualData("inv", "payment_flows", {}, "USD", "amount", "flows", "payments", list(rows), "USD", sum(r.record_count for r in rows))


CHAIN = _data(_flow("Meridian", "Bellcastle", 2_400_000, 2), _flow("Bellcastle", "Orrin", 900_000, 1), _flow("Meridian", "Orrin", 150_000, 3))
CIRCLE = _data(*CHAIN.rows, _flow("Orrin", "Meridian", 350_000, 2))


def _spec(**extra):
    return parse_spec({"kind": "sankey", "title": "Most of Meridian's money went to Bellcastle", "query": {"name": "payment_flows"}, **extra})


def _sankey(option: dict) -> dict:
    return next(s for s in option["series"] if s["type"] == "sankey")


def test_a_chain_of_payments_is_drawn_left_to_right_and_passes_its_check():
    built = build_chart(_spec(), CHAIN, DARK)
    series = _sankey(built.option)
    depths = {node["name"]: node["depth"] for node in series["data"]}
    # A party's column is the longest chain of payments into it; one that pays no one is last.
    assert depths == {"Meridian": 0, "Bellcastle": 1, "Orrin": 2}
    assert {(l["name"], l["source"], l["target"], l["value"]) for l in series["links"]} == {
        ("Meridian → Bellcastle", "Meridian", "Bellcastle", 2_400_000), ("Bellcastle → Orrin", "Bellcastle", "Orrin", 900_000),
        ("Meridian → Orrin", "Meridian", "Orrin", 150_000)}
    links = {l["name"]: l for l in series["links"]}
    # Every flow says its amount and payments in its tooltip; on its band only
    # when the band goes one column and is thick enough to hold the words.
    assert links["Meridian → Orrin"]["tooltip"]["formatter"] == "Meridian → Orrin: $150K · 3 payments"
    assert links["Bellcastle → Orrin"]["tooltip"]["formatter"].endswith(f"1 payment {SINGLE_RECORD_MARK}")
    assert links["Meridian → Bellcastle"]["edgeLabel"]["show"] is True
    assert links["Meridian → Orrin"]["edgeLabel"]["show"] is False     # passes behind Bellcastle's bar
    names = {node["name"]: node["label"] for node in series["data"]}
    assert names["Orrin"]["position"] == "left" and names["Meridian"]["position"] == "right"
    report = check_visual(_spec(), CHAIN, DARK, built)
    assert report.ok, report.errors
    assert {0, 2, 3, 9, 10, 11, 12, 13, 14, 15, 16, 20} <= set(report.checked_rules)
    assert {1, 5, 8} <= set(report.not_applicable)
    assert built.data_table["columns"] == ["Flow", "From", "To", "Value (USD)", "Records"]
    assert built.alt_text.startswith("Flow diagram: Most of Meridian's money went to Bellcastle. 3 flows between 3 parties")


def test_money_that_comes_back_round_is_drawn_payers_left_recipients_right():
    assert not flows_circle(CHAIN.rows) and flows_circle(CIRCLE.rows)
    built = build_chart(_spec(), CIRCLE, DARK)
    series = _sankey(built.option)
    depths = {node["name"]: node["depth"] for node in series["data"]}
    assert depths == {"pays: Meridian": 0, "pays: Bellcastle": 0, "pays: Orrin": 0,
                      "receives: Bellcastle": 1, "receives: Orrin": 1, "receives: Meridian": 1}
    assert {node["label"]["formatter"] for node in series["data"]} == {"Meridian", "Bellcastle", "Orrin"}
    assert "payers are drawn on the left and recipients on the right" in built.option["title"][0]["subtext"]
    assert check_visual(_spec(), CIRCLE, DARK, built).ok


def test_the_columns_and_band_thickness_follow_echarts_own_layout():
    ids, split = sankey_node_ids(CHAIN.rows)
    layout = sankey_layout(CHAIN.rows, ids, 412)
    assert not split and layout.depth == {"Meridian": 0, "Bellcastle": 1, "Orrin": 2}
    # The tightest column decides: Meridian's 2.55M fills the first column alone.
    assert layout.ky == pytest.approx(412 / 2_550_000)


def test_orders_a_flow_diagram_cannot_take_are_refused_in_plain_words():
    with pytest.raises(ChartDoesNotFit, match="limit folds ranked rows"):
        build_chart(_spec(limit=2), CHAIN, DARK)
    with pytest.raises(ChartDoesNotFit, match="sort does not apply to a flow diagram"):
        build_chart(_spec(sort="value_desc"), CHAIN, DARK)
    donut = parse_spec({"kind": "donut", "title": "Most of Meridian's money went to Bellcastle", "query": {"name": "payment_flows"}})
    with pytest.raises(ChartDoesNotFit, match="draw it as a flow diagram"):
        build_chart(donut, CHAIN, DARK)
    by_party = VisualData("inv", "payments_by_party", {}, "USD", "amount", "category", "payments", [_flow("A", "B", 1, 1)], "USD", 1)
    with pytest.raises(ChartDoesNotFit, match="a sankey needs money moving between parties"):
        build_chart(_spec(), by_party, DARK)
    # Ranked as bars, the flows are folded past a limit like any ranking.
    bars = build_chart(parse_spec({"kind": "bar", "title": "Most of Meridian's money went to Bellcastle", "query": {"name": "payment_flows"},
                                   "limit": 1}), CHAIN, DARK)
    assert [row.label for row in bars.rows] == ["Meridian → Bellcastle", "Other"]


def _damaged(mutate) -> list[int]:
    built = build_chart(_spec(), CHAIN, DARK)
    built.option = copy.deepcopy(built.option)
    mutate(_sankey(built.option))
    return [item.rule for item in check_visual(_spec(), CHAIN, DARK, built).errors]


def test_the_checker_reads_the_bands_and_the_parties():
    def wrong_value(series):
        series["links"][0]["value"] = 1
    assert 0 in _damaged(wrong_value)

    def no_count(series):
        series["links"][1]["tooltip"]["formatter"] = "Bellcastle → Orrin: $900K"
    assert {16, 20} <= set(_damaged(no_count))

    def nameless(series):
        series["data"][0]["label"]["show"] = False
    assert 12 in _damaged(nameless)

    def pale(series):
        for node in series["data"]:
            node["itemStyle"]["color"] = "#" + DARK.page
    assert 11 in _damaged(pale)

    many = _data(*[_flow(f"Payer {i}", f"Recipient {i}", 1000 + i, 2) for i in range(MAX_FLOW_PARTIES)])
    assert 10 in [w.rule for w in check_visual(_spec(), many, DARK).warnings]


def test_a_party_or_a_flow_can_be_picked_out():
    # The style keeps the parties as well as the flows the data has.
    clean = normalize_chart_style(_spec(), CHAIN, {"options": {"highlight": ["Orrin", "Meridian → Bellcastle", "Nobody"]}})
    assert clean == {"preset": "", "options": {"highlight": ["Meridian → Bellcastle", "Orrin"]}}
    styled, draw = styled_chart(_spec(), CHAIN, clean)
    built = build_chart(styled, CHAIN, DARK, draw=draw)
    series = _sankey(built.option)
    links = {l["name"]: l for l in series["links"]}
    accent = "#" + DARK.accent
    assert {name for name, l in links.items() if l["lineStyle"]["color"] == accent} == {"Meridian → Bellcastle", "Bellcastle → Orrin", "Meridian → Orrin"}
    # A picked flow writes its amount in bold, whatever the room (its colour is not the only cue).
    assert links["Meridian → Orrin"]["edgeLabel"]["show"] is True and links["Meridian → Orrin"]["edgeLabel"]["fontWeight"] == "bold"
    orrin = next(n for n in series["data"] if n["name"] == "Orrin")
    assert orrin["label"]["fontWeight"] == "bold"
    assert check_visual(styled, CHAIN, DARK, built).ok


@pytest.fixture()
def flows_chart():
    with SessionLocal() as db:
        inv, _ = seed_demo_investigation(db)
        ids = {e.caption: e.id for e in db.query(Entity).filter_by(investigation_id=inv.id)}
        create_relationship(db, investigation_id=inv.id, schema="Payment", source_entity_id=ids["Meridian Holdings Ltd"],
                            target_entity_id=ids["Alaric Voss"], properties={"amount": ["45000"], "currency": ["USD"], "date": ["2024-01-20"]})
        chart, record = resolve_chart(db, inv.id, {"type": "chart", "spec": {
            "kind": "sankey", "title": "Bellcastle's money went on through Meridian", "query": {"name": "payment_flows"}}})
        assert chart is not None, record
        yield db, inv.id, chart


def _deck(db, inv_id, chart) -> Deck:
    row = Deck(investigation_id=inv_id, title="Flows", slides_json=[
        {"kind": "content", "heading": "Where the money went", "bullets": [], "evidence": [], "notes": "", "visual": chart}])
    db.add(row)
    db.commit()
    db.refresh(row)
    return row


def test_a_deck_holds_the_flow_diagram_as_the_rendered_picture_or_ranked_bars(flows_chart, monkeypatch):
    db, inv_id, chart = flows_chart
    assert chart["report"]["ok"] and chart["data"]["rows"][0]["source"] == "Bellcastle Trading Group"
    drawn = []

    def fake_render(option, width, height, **kwargs):
        drawn.append(option["series"][0]["type"])
        return TINY_PNG

    monkeypatch.setattr("app.visuals.render_client.render_png", fake_render)
    slide = list(render_deck_pptx(_deck(db, inv_id, chart), []).slides)[1]
    assert drawn == ["sankey"]
    assert len([s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]) == 1

    def unreachable(*args, **kwargs):
        raise RenderUnavailable("render service at http://render:8070 could not be reached: ConnectError")

    monkeypatch.setattr("app.visuals.render_client.render_png", unreachable)
    slide = list(render_deck_pptx(_deck(db, inv_id, chart), []).slides)[1]
    frames = [shape for shape in slide.shapes if shape.has_chart]
    assert len(frames) == 1 and list(frames[0].chart.plots[0].categories)[-1].startswith("Bellcastle Trading Group → Meridian Holdings Ltd")


# --- the flow in stages (for a slide or print) --------------------------------------------

LAUNDROMAT = _data(_flow("Shell A", "Moldindcon", 900_000, 2), _flow("Shell B", "Moldindcon", 600_000, 1),
                   _flow("Moldindcon", "Trasta", 1_400_000, 3), _flow("Trasta", "Luxury goods", 700_000, 2),
                   _flow("Trasta", "Real estate", 500_000, 1), _flow("Shell A", "Real estate", 50_000, 1))


def test_the_parties_at_each_place_in_the_chain_are_grouped_into_one_stage_and_no_money_is_lost():
    rows, stages = flow_stages(LAUNDROMAT.rows)
    assert stages == {"Stage 1: Shell A, Shell B": ["Shell A", "Shell B"], "Stage 2: Moldindcon": ["Moldindcon"],
                      "Stage 3: Trasta": ["Trasta"], "Stage 4: Luxury goods, Real estate": ["Luxury goods", "Real estate"]}
    assert [(r.source, r.target, r.value, r.record_count) for r in rows] == [
        ("Stage 1: Shell A, Shell B", "Stage 2: Moldindcon", 1_500_000, 3), ("Stage 2: Moldindcon", "Stage 3: Trasta", 1_400_000, 3),
        ("Stage 3: Trasta", "Stage 4: Luxury goods, Real estate", 1_200_000, 3),
        ("Stage 1: Shell A, Shell B", "Stage 4: Luxury goods, Real estate", 50_000, 1)]
    # Every payment is in exactly one staged flow, and every amount is in the totals.
    assert sorted(ref.record_id for r in rows for ref in r.record_refs) == sorted(ref.record_id for r in LAUNDROMAT.rows for ref in r.record_refs)
    assert sum(r.value for r in rows) == sum(r.value for r in LAUNDROMAT.rows)
    many = flow_stages(_data(*[_flow(f"Payer {i}", "Hub", 10, 1) for i in range(5)]).rows)[1]
    assert list(many) == ["Stage 1: 5 parties", "Stage 2: Hub"]


def test_the_staged_preset_draws_the_stages_names_them_above_their_bars_and_passes_its_check():
    styled, draw = styled_chart(_spec(), LAUNDROMAT, {"preset": "staged", "options": {"highlight": ["Moldindcon"]}})
    assert draw.stages
    built = build_chart(styled, LAUNDROMAT, DARK, draw=draw)
    series = _sankey(built.option)
    nodes = {n["name"]: n for n in series["data"]}
    assert nodes["Stage 1: Shell A, Shell B"]["label"]["formatter"] == "Stage 1\nShell A, Shell B"
    assert nodes["Stage 1: Shell A, Shell B"]["label"]["position"] == "top" and nodes["Stage 1: Shell A, Shell B"]["label"]["align"] == "left"
    assert nodes["Stage 4: Luxury goods, Real estate"]["label"]["align"] == "right"
    # A party picked out picks out its stage.
    assert nodes["Stage 2: Moldindcon"]["label"]["fontWeight"] == "bold"
    assert "a stage groups the parties at the same place in the chain of payments" in built.option["title"][0]["subtext"]
    assert "Stage 4: Luxury goods, Real estate." in built.description
    assert check_visual(styled, LAUNDROMAT, DARK, built).ok


def test_money_that_comes_back_round_cannot_be_put_in_stages():
    styled, draw = styled_chart(_spec(), CIRCLE, {"preset": "staged"})
    with pytest.raises(ChartDoesNotFit, match="no order to group into stages"):
        build_chart(styled, CIRCLE, DARK, draw=draw)
