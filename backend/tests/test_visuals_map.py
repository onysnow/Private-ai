"""The map (visuals pass 5; app/visuals/echarts.py _map,
data_service.places_on_map, render_client.render_map_png): every record
with coordinates as a named dot where its record puts it, drawn on the
self-hosted base map by MapLibre (the app) and MapLibre Native (the render
service), framed by the chart's title and source line; checked like every
other chart, and in decks and briefs."""

from __future__ import annotations

import base64
import json

import httpx
import pytest
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.ai.chart_check import resolve_chart
from app.core.config import Settings
from app.db.session import SessionLocal
from app.models.domain import Deck, Investigation
from app.schemas.api import EntityCreate
from app.services.decks import render_deck_pptx
from app.services.entities import create_entity
from app.services.presentation_style import THEMES
from app.services.writeups import create_writeup, extract_plain_text, render_writeup_docx
from app.visuals.checker import check_visual
from app.visuals.data_service import DataRow, RecordRef, VisualData, run_query
from app.visuals.echarts import MAP_BOTTOM, MAP_TOP, ChartDoesNotFit, build_chart
from app.visuals.render_client import RenderUnavailable, map_area, render_map_png
from app.visuals.spec import QueryRef, parse_spec

DARK = THEMES["dark_academia"]
TINY_PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4890000000d4944415478da63f8cfc0f01f0005020203"
    "f1a2b9780000000049454e44ae426082"
)


def _place(label: str, latitude: float, longitude: float) -> DataRow:
    return DataRow(label, 1.0, [RecordRef("entity", label)], 1, label.lower(), latitude=latitude, longitude=longitude)


PLACES = VisualData("inv", "places_on_map", {}, "places", "count", "places", "places",
                    [_place("Meridian registered office", 51.5072, -0.1276), _place("Villa Kyrenia", 35.3364, 33.3199)], record_count=2)


def _spec(**extra):
    return parse_spec({"kind": "map", "question": "place", "title": "The money ended up in a villa in Cyprus",
                       "query": {"name": "places_on_map"}, **extra})


# --- the drawing ------------------------------------------------------------------------

def test_each_place_is_a_named_dot_where_its_record_puts_it():
    built = build_chart(_spec(), PLACES, DARK)
    drawn = built.option["map"]
    assert (drawn["top"], drawn["bottom"]) == (MAP_TOP, MAP_BOTTOM)
    assert [(p["name"], p["latitude"], p["longitude"], p["label"]["text"]) for p in drawn["places"]] == [
        ("Meridian registered office", 51.5072, -0.1276, "Meridian registered office"), ("Villa Kyrenia", 35.3364, 33.3199, "Villa Kyrenia")]
    assert {p["colour"] for p in drawn["places"]} == {"#8A2F1E"}
    # The map's data is credited on the source line (ODbL), and the subtitle says nothing is looked up.
    assert built.source_line.endswith("Map data © OpenStreetMap contributors")
    assert "nothing is looked up" in built.option["title"][0]["subtext"]
    assert built.data_table["columns"] == ["Place", "Where", "Records"]
    assert built.data_table["rows"][1]["formatted"] == "35.3364° N, 33.3199° E"
    assert "Villa Kyrenia: 35.3364° N, 33.3199° E." in built.description


@pytest.mark.parametrize("theme_id", sorted(THEMES))
def test_a_map_passes_its_check_in_every_theme(theme_id):
    report = check_visual(_spec(theme=theme_id), PLACES, THEMES[theme_id])
    assert report.ok, report.errors


def test_a_picked_place_keeps_its_colour_with_a_larger_dot_and_a_bold_name():
    built = build_chart(_spec(highlight=["Villa Kyrenia"]), PLACES, DARK)
    places = {p["name"]: p for p in built.option["map"]["places"]}
    assert places["Villa Kyrenia"]["radius"] > places["Meridian registered office"]["radius"]
    assert places["Villa Kyrenia"]["label"]["fontWeight"] == "bold" and places["Meridian registered office"]["label"]["fontWeight"] == "normal"
    assert places["Villa Kyrenia"]["colour"] == places["Meridian registered office"]["colour"]   # a grey dot would vanish into the base map
    assert check_visual(_spec(highlight=["Villa Kyrenia"]), PLACES, DARK, built).ok


def test_orders_a_map_cannot_take_are_refused_in_plain_words():
    with pytest.raises(ChartDoesNotFit, match="sort, limit and a log scale do not apply"):
        build_chart(_spec(limit=1), PLACES, DARK)
    with pytest.raises(ChartDoesNotFit, match="draw them on a map"):
        build_chart(parse_spec({"kind": "bar", "title": "Two places", "query": {"name": "places_on_map"}}), PLACES, DARK)
    events = VisualData("inv", "events_in_order", {}, "events", "count", "events", "events", [], record_count=0)
    with pytest.raises(ChartDoesNotFit, match="a map needs records with coordinates"):
        build_chart(parse_spec({"kind": "map", "title": "Where it happened", "query": {"name": "events_in_order"}}), events, DARK)


def test_the_checker_reads_every_place_back_off_the_drawing():
    built = build_chart(_spec(), PLACES, DARK)
    moved = json.loads(json.dumps(built.option))
    moved["map"]["places"][1]["latitude"] = 36.0
    report = check_visual(_spec(), PLACES, DARK, built.__class__(**{**built.__dict__, "option": moved}))
    assert 0 in report.error_rules()
    unnamed = json.loads(json.dumps(built.option))
    unnamed["map"]["places"][0]["label"]["text"] = ""
    assert 12 in check_visual(_spec(), PLACES, DARK, built.__class__(**{**built.__dict__, "option": unnamed})).error_rules()
    faint = json.loads(json.dumps(built.option))
    faint["map"]["places"][0]["colour"] = "#B0B0B0"
    assert 11 in check_visual(_spec(), PLACES, DARK, built.__class__(**{**built.__dict__, "option": faint})).error_rules()
    plain = build_chart(_spec(highlight=["Villa Kyrenia"]), PLACES, DARK)
    unbold = json.loads(json.dumps(plain.option))
    for place in unbold["map"]["places"]:
        place["label"]["fontWeight"] = "normal"
    report = check_visual(_spec(highlight=["Villa Kyrenia"]), PLACES, DARK, plain.__class__(**{**plain.__dict__, "option": unbold}))
    assert any(e.rule == 12 and "bold" in e.message for e in report.errors)


# --- the query ---------------------------------------------------------------------------

def _investigation() -> str:
    with SessionLocal() as db:
        inv = Investigation(name="Where the money went")
        db.add(inv)
        db.commit()
        db.refresh(inv)

        def entity(schema: str, caption: str, **props: list[str]) -> str:
            return create_entity(db, EntityCreate(investigation_id=inv.id, schema=schema, caption=caption, properties=props)).id

        entity("RealEstate", "Villa Kyrenia", latitude=["35.3364"], longitude=["33.3199"])
        entity("Address", "Meridian registered office", latitude=["51.5072"], longitude=["-0.1276"])
        entity("Address", "Branch office", latitude=["56.9496"], longitude=["24.1052"])
        entity("Address", "Branch office", latitude=["47.0105"], longitude=["28.8638"])
        entity("Address", "Typo address", latitude=["abc"], longitude=["200"])
        entity("Company", "Meridian Holdings Ltd")
        return inv.id


def test_places_on_map_reads_each_record_with_its_coordinates_and_says_which_it_left_out():
    inv_id = _investigation()
    with SessionLocal() as db:
        data = run_query(db, inv_id, QueryRef(name="places_on_map"))
    assert data.series_kind == "places" and data.record_count == 4
    labels = [r.label for r in data.rows]
    # Two places of one name are told apart by their record id.
    assert labels[2:] == ["Meridian registered office", "Villa Kyrenia"]
    assert all(label.startswith("Branch office [") for label in labels[:2]) and labels[0] != labels[1]
    assert all(r.record_count == 1 and r.record_refs[0].record_type == "entity" for r in data.rows)
    assert [(o["reason"], o["detail"].split(":")[0]) for o in data.omitted] == [("coordinates not usable", "Typo address")]
    assert VisualData.from_dict(data.as_dict()).rows[3].latitude == 35.3364


# --- exports -----------------------------------------------------------------------------

@pytest.fixture()
def map_chart():
    inv_id = _investigation()
    with SessionLocal() as db:
        chart, record = resolve_chart(db, inv_id, {"type": "chart", "spec": {
            "kind": "map", "question": "place", "title": "The money ended up in a villa in Cyprus", "query": {"name": "places_on_map"}}})
        assert chart is not None, record
        yield db, inv_id, chart


def _deck(db, inv_id, chart) -> Deck:
    row = Deck(investigation_id=inv_id, title="Where", slides_json=[
        {"kind": "content", "heading": "Where the money went", "bullets": [], "evidence": [], "notes": "", "visual": chart}])
    db.add(row)
    db.commit()
    db.refresh(row)
    return row


def test_a_deck_holds_the_map_as_the_rendered_picture_or_its_places_listed(map_chart, monkeypatch):
    db, inv_id, chart = map_chart
    drawn = []

    def fake_render(option, width, height, **kwargs):
        drawn.append(len(option["map"]["places"]))
        return TINY_PNG

    monkeypatch.setattr("app.visuals.render_client.render_png", fake_render)
    slide = list(render_deck_pptx(_deck(db, inv_id, chart), []).slides)[1]
    assert drawn == [4]
    assert len([s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]) == 1

    def unreachable(*args, **kwargs):
        raise RenderUnavailable("render service at http://render:8070 could not be reached: ConnectError")

    monkeypatch.setattr("app.visuals.render_client.render_png", unreachable)
    slide = list(render_deck_pptx(_deck(db, inv_id, chart), []).slides)[1]
    tables = [shape for shape in slide.shapes if shape.has_table]
    assert len(tables) == 1 and not [s for s in slide.shapes if s.has_chart]
    cells = [[cell.text for cell in row.cells] for row in tables[0].table.rows]
    assert cells[0] == ["Place", "Where"]
    assert ["Villa Kyrenia", "35.3364° N, 33.3199° E"] in cells
    note = " ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
    assert "image not available" in note and "the places are listed instead" in note


def test_a_brief_lists_every_place_with_where_it_is_beside_the_picture_or_without_it(map_chart, monkeypatch):
    db, inv_id, chart = map_chart
    content = {"type": "doc", "content": [{"type": "evidenceVisual", "attrs": {"visual": chart}}]}
    row = create_writeup(db, investigation_id=inv_id, title="Where", content_json=content)
    assert "Villa Kyrenia — 35.3364° N, 33.3199° E" in extract_plain_text(row.content_json)

    def unreachable(*args, **kwargs):
        raise RenderUnavailable("render service at http://render:8070 could not be reached: ConnectError")

    monkeypatch.setattr("app.visuals.render_client.render_png", unreachable)
    rows = [[c.text for c in r.cells] for t in render_writeup_docx(row, []).tables for r in t.rows]
    # Each place with its citation number, never the single-record mark: a place is one record.
    assert ["Villa Kyrenia [4]", "35.3364° N, 33.3199° E", "1"] in rows
    assert not any("†" in r[0] for r in rows)

    # A chart that no longer passes its check is not drawn; its places are still listed.
    broken = json.loads(json.dumps(chart))
    broken["spec"]["limit"] = 2
    row = create_writeup(db, investigation_id=inv_id, title="Where", content_json={"type": "doc", "content": [
        {"type": "evidenceVisual", "attrs": {"visual": broken}}]})
    docx = render_writeup_docx(row, [])
    rows = [[c.text for c in r.cells] for t in docx.tables for r in t.rows]
    assert ["Villa Kyrenia [4]", "35.3364° N, 33.3199° E"] in rows
    assert any("sort, limit and a log scale do not apply" in p.text for p in docx.paragraphs)


# --- the render service ----------------------------------------------------------------------

def test_the_render_service_is_asked_for_the_map_area_and_the_picture_framed_around_it():
    built = build_chart(_spec(highlight=["Villa Kyrenia"]), PLACES, DARK)
    asked: list[tuple[str, dict]] = []

    def answer(request: httpx.Request) -> httpx.Response:
        body = json.loads(request.content)
        asked.append((request.url.path, body))
        if request.url.path == "/map":
            return httpx.Response(200, headers={"content-type": "image/png"}, content=TINY_PNG)
        return httpx.Response(200, headers={"content-type": "image/svg+xml"}, text='<svg xmlns="http://www.w3.org/2000/svg" width="960" height="540"/>')

    settings = Settings(render_service_url="http://render.test", _env_file=None)
    assert map_area(built.option, 960, 540, settings=settings, transport=httpx.MockTransport(answer)) == TINY_PNG
    path, body = asked[0]
    assert path == "/map" and (body["width"], body["height"]) == (960, 540 - MAP_TOP - MAP_BOTTOM)
    assert body["places"][1] == {"longitude": 33.3199, "latitude": 35.3364, "label": "Villa Kyrenia", "radius": 10, "picked": True}
    assert body["places"][0]["picked"] is False

    asked.clear()
    assert render_map_png(built.option, 960, 540, settings=settings, transport=httpx.MockTransport(answer)).startswith(b"\x89PNG")
    framed = asked[1][1]["option"]
    assert "map" not in framed and framed["title"] == built.option["title"]
    image = framed["graphic"][0]
    assert (image["top"], image["style"]["height"]) == (MAP_TOP, 540 - MAP_TOP - MAP_BOTTOM)
    assert image["style"]["image"] == "data:image/png;base64," + base64.b64encode(TINY_PNG).decode("ascii")


def test_a_render_service_without_a_base_map_is_reported_not_drawn_around():
    built = build_chart(_spec(), PLACES, DARK)
    settings = Settings(render_service_url="http://render.test", _env_file=None)
    no_map = httpx.MockTransport(lambda request: httpx.Response(503, text="no base map at /basemap/world.pmtiles"))
    with pytest.raises(RenderUnavailable, match="503: no base map"):
        map_area(built.option, 960, 540, settings=settings, transport=no_map)
