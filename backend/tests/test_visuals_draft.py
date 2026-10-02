"""A presentation draft whose model orders a chart (app/ai/chart_check.py
inside app/ai/reasoning.py): the order is cooked from the records, drawn
and checked; a failing order goes back to the model at most twice, then
is removed with its reasons recorded. Follows test_presentation_draft.py's
fake-model pattern: no network, no API key."""

import json
from io import BytesIO

import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from pptx import Presentation

import app.ai.reasoning as reasoning_module
from app.ai.governor import governor
from app.core.config import settings
from app.db.session import SessionLocal
from app.main import app
from app.models.domain import AIAnalysisCandidate, Claim, Evidence
from app.services.demo_seed import seed_demo_investigation
from app.visuals.render_client import RenderUnavailable

client = TestClient(app)


class _FakeLLMClient:
    """Answers with each response in turn, then repeats the last one."""

    def __init__(self, *responses: dict):
        self._responses = [json.dumps(r) for r in responses]
        self.calls = 0
        self.user_prompts: list[str] = []

    def generate(self, system_prompt: str, user_prompt: str, *, max_tokens: int = 4096) -> str:
        self.user_prompts.append(user_prompt)
        index = min(self.calls, len(self._responses) - 1)
        self.calls += 1
        return self._responses[index]


@pytest.fixture(autouse=True)
def _ai_on():
    original = (settings.enable_ai_features, settings.ai_provider, settings.ai_runs_per_minute, settings.ai_max_concurrent_runs)
    settings.enable_ai_features = True
    settings.ai_provider = "anthropic"
    settings.ai_runs_per_minute = 100
    governor.reset()
    yield
    settings.enable_ai_features, settings.ai_provider, settings.ai_runs_per_minute, settings.ai_max_concurrent_runs = original
    governor.reset()


@pytest.fixture()
def demo() -> dict:
    with SessionLocal() as db:
        inv, _ = seed_demo_investigation(db)
        claim = db.query(Claim).filter(Claim.investigation_id == inv.id, Claim.text.like("Meridian Holdings Ltd is wholly%")).one()
        evidence = db.query(Evidence).filter(Evidence.quote.like("Bank records show%")).one()
        return {"investigation_id": inv.id, "claim_id": claim.id, "evidence_id": evidence.id}


def _draft(demo: dict, visual: dict | None) -> dict:
    claim_ref = {"record_type": "claim", "record_id": demo["claim_id"], "quote": "wholly owned by Bellcastle"}
    evidence_ref = {"record_type": "evidence", "record_id": demo["evidence_id"], "quote": "transfer of $2,400,000 from Bellcastle"}
    return {
        "title": "Where the Bellcastle money went",
        "goal_restated": "Show who received the money.",
        "executive_takeaway": "Meridian received the whole $2.4M.",
        "sections": [{
            "heading": "Who got paid",
            "body": "Bank records show one transfer, to Meridian.",
            "key_points": [{"text": "One transfer, to Meridian", "evidence": [evidence_ref]}],
            "evidence": [claim_ref, evidence_ref],
            "visual": visual,
        }],
        "open_questions": ["Whether other transfers exist."],
    }


def _chart(**spec_overrides) -> dict:
    spec = {"kind": "bar", "title": "Meridian received all of the money", "query": {"name": "payments_by_party", "params": {"direction": "received"}}}
    spec.update(spec_overrides)
    return {"type": "chart", "spec": spec}


def _run(monkeypatch, demo: dict, fake: _FakeLLMClient):
    monkeypatch.setattr(reasoning_module, "get_llm_client", lambda *a, **k: fake)
    return client.post(
        f"/api/investigations/{demo['investigation_id']}/assistant/presentation-draft",
        json={"goal": "Show who received the money.", "record_refs": [
            {"record_type": "claim", "record_id": demo["claim_id"]},
            {"record_type": "evidence", "record_id": demo["evidence_id"]},
        ]},
    )


def test_a_good_chart_order_is_cooked_and_stored_with_the_draft(monkeypatch, demo):
    fake = _FakeLLMClient(_draft(demo, _chart()))
    response = _run(monkeypatch, demo, fake)
    assert response.status_code == 200, response.text
    assert fake.calls == 1
    assert '"visual_catalog"' in fake.user_prompts[0] and '"payments_by_party"' in fake.user_prompts[0]
    body = response.json()
    assert body["review_status"] == "proposed"
    visual = body["payload"]["sections"][0]["visual"]
    assert visual["type"] == "chart"
    assert visual["spec"]["kind"] == "bar" and visual["spec"]["theme"] == "dark_academia"
    assert visual["data"]["rows"] == [{"label": "Meridian Holdings Ltd", "value": 2400000.0, "record_count": 1, "key": "Meridian Holdings Ltd",
                                       "record_refs": visual["data"]["rows"][0]["record_refs"]}]
    assert visual["data"]["rows"][0]["record_refs"][0]["record_type"] == "relationship"
    assert visual["report"]["ok"] is True
    assert visual["text"]["alt_text"].startswith("Bar chart: Meridian received all of the money")
    assert "data_table" in visual["text"]

    with SessionLocal() as db:
        row = db.get(AIAnalysisCandidate, body["candidate_id"])
        assert row is not None
        assert row.trace["chart_check"]["removed"] == []
        assert row.trace["chart_check"]["resolved"][0]["where"] == "Who got paid: chart"
        assert "chart_retries" not in row.trace

    detail = client.get(f"/api/ai-analysis-candidates/{body['candidate_id']}").json()
    by_type = {d["record_type"]: d for d in detail["citation_details"]}
    assert "relationship" in by_type and not by_type["relationship"]["missing"]
    assert by_type["relationship"]["label"] == "Bellcastle Trading Group — Payment → Meridian Holdings Ltd"
    assert detail["deck_preview"][2]["visual"]["type"] == "chart"


def test_a_failing_chart_goes_back_to_the_model_and_the_fixed_order_is_kept(monkeypatch, demo):
    fake = _FakeLLMClient(_draft(demo, _chart(title="Payments")), _draft(demo, _chart()))
    response = _run(monkeypatch, demo, fake)
    assert response.status_code == 200, response.text
    assert fake.calls == 2
    assert "CHART CHECK" in fake.user_prompts[1] and "rule 13" in fake.user_prompts[1]
    assert "Never add numbers" in fake.user_prompts[1]
    body = response.json()
    assert body["payload"]["sections"][0]["visual"]["report"]["ok"] is True
    with SessionLocal() as db:
        row = db.get(AIAnalysisCandidate, body["candidate_id"])
        assert row is not None
        retries = row.trace["chart_retries"]
        assert len(retries) == 1 and retries[0]["attempt"] == 1
        assert retries[0]["removed"][0]["rules"] == [13]
        assert row.trace["chart_check"]["removed"] == []


def test_after_two_retries_a_failing_chart_is_removed_and_the_reasons_recorded(monkeypatch, demo):
    fake = _FakeLLMClient(_draft(demo, _chart(title="Payments")))
    response = _run(monkeypatch, demo, fake)
    assert response.status_code == 200, response.text
    assert fake.calls == 3  # the first answer and two re-asks
    body = response.json()
    assert body["review_status"] == "proposed"  # the rest of the draft stands
    assert body["payload"]["sections"][0]["visual"] is None
    with SessionLocal() as db:
        row = db.get(AIAnalysisCandidate, body["candidate_id"])
        assert row is not None
        removed = row.trace["chart_check"]["removed"]
        assert len(removed) == 1
        assert removed[0]["where"] == "Who got paid: chart" and removed[0]["title"] == "Payments"
        assert removed[0]["rules"] == [13]
        assert "state the finding" in removed[0]["errors"][0]["message"]
        assert len(row.trace["chart_retries"]) == 2


def test_a_chart_over_a_query_with_no_records_is_removed_under_rule_15(monkeypatch, demo):
    fake = _FakeLLMClient(_draft(demo, _chart(title="Most documents have been read", query={"name": "documents_by_status"})))
    response = _run(monkeypatch, demo, fake)
    assert response.status_code == 200, response.text
    assert response.json()["payload"]["sections"][0]["visual"] is None
    with SessionLocal() as db:
        row = db.get(AIAnalysisCandidate, response.json()["candidate_id"])
        assert row is not None
        assert row.trace["chart_check"]["removed"][0]["rules"] == [15]


def test_a_kind_not_drawn_yet_is_removed_with_a_plain_reason(monkeypatch, demo):
    fake = _FakeLLMClient(_draft(demo, _chart(kind="sankey", title="The money flowed from Bellcastle to Meridian")))
    response = _run(monkeypatch, demo, fake)
    assert response.status_code == 200, response.text
    with SessionLocal() as db:
        row = db.get(AIAnalysisCandidate, response.json()["candidate_id"])
        assert row is not None
        assert "not drawn yet" in row.trace["chart_check"]["removed"][0]["reason"]


def test_a_kind_that_cannot_draw_its_query_goes_back_with_the_reason_and_a_fitting_kind_is_kept(monkeypatch, demo):
    histogram_over_time = _chart(kind="histogram", title="Payments came in one burst", query={"name": "payments_over_time"})
    fixed = _chart(kind="stacked_bar_100", title="Meridian received all of the money")
    fake = _FakeLLMClient(_draft(demo, histogram_over_time), _draft(demo, fixed))
    response = _run(monkeypatch, demo, fake)
    assert response.status_code == 200, response.text
    assert fake.calls == 2
    assert "a histogram needs one value per record, such as payment_amounts" in fake.user_prompts[1]
    visual = response.json()["payload"]["sections"][0]["visual"]
    assert visual["spec"]["kind"] == "stacked_bar_100" and visual["report"]["ok"] is True
    assert visual["report"]["not_applicable"]["8"].startswith("a stacked_bar_100 joins no points")
    assert visual["text"]["data_table"]["rows"][0]["share"] == 100.0


@pytest.mark.parametrize("bad_visual, fragment", [
    (_chart(query={"name": "payments_by_payee"}), "unknown query 'payments_by_payee'"),
    ({"type": "chart", "spec": {**_chart()["spec"], "values": [2400000]}}, "Extra inputs are not permitted"),
    ({"type": "chart"}, "spec must be an object"),
    (_chart(query={"name": "payments_by_party", "params": {"direction": "up"}}), "must be one of received, paid"),
])
def test_a_malformed_order_slip_is_a_schema_rejection(monkeypatch, demo, bad_visual, fragment):
    fake = _FakeLLMClient(_draft(demo, bad_visual))
    response = _run(monkeypatch, demo, fake)
    assert response.status_code == 422, response.text
    assert fragment in response.text
    assert fake.calls == 1


def test_the_three_old_visuals_still_work_beside_a_chart(monkeypatch, demo):
    ev = [{"record_type": "evidence", "record_id": demo["evidence_id"], "quote": "transfer of $2,400,000"}]
    draft = _draft(demo, _chart())
    draft["sections"].append({"heading": "Headline", "body": "The sum.", "key_points": [{"text": "$2.4M", "evidence": ev}], "evidence": ev,
                              "visual": {"type": "figures", "items": [{"value": "$2.4M", "label": "transferred", "evidence": ev}]}})
    fake = _FakeLLMClient(draft)
    response = _run(monkeypatch, demo, fake)
    assert response.status_code == 200, response.text
    kinds = [s["visual"]["type"] for s in response.json()["payload"]["sections"]]
    assert kinds == ["chart", "figures"]


def test_promotion_exports_the_chart_in_the_deck_and_the_brief(monkeypatch, demo):
    fake = _FakeLLMClient(_draft(demo, _chart()))
    response = _run(monkeypatch, demo, fake)
    candidate_id = response.json()["candidate_id"]
    assert client.post(f"/api/ai-analysis-candidates/{candidate_id}/review", json={"decision": "accept"}).status_code == 200

    deck = client.post(f"/api/ai-analysis-candidates/{candidate_id}/promote-to-deck", json={}).json()
    assert deck["slides"][2]["visual"]["type"] == "chart"
    exported = client.get(f"/api/decks/{deck['id']}/export.pptx")
    assert exported.status_code == 200
    prs = Presentation(BytesIO(exported.content))
    charts = [sh for s in prs.slides for sh in s.shapes if sh.has_chart]
    assert len(charts) == 1
    assert list(charts[0].chart.plots[0].series[0].values) == [2400000.0]
    text = "\n".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    assert "Bellcastle Trading Group — Payment → Meridian Holdings Ltd" in text  # on the sources slide

    def unreachable(*args, **kwargs):
        raise RenderUnavailable("render service at http://render:8070 could not be reached")

    monkeypatch.setattr("app.visuals.render_client.render_png", unreachable)
    writeup = client.post(f"/api/ai-analysis-candidates/{candidate_id}/promote-to-writeup", json={}).json()
    assert any(n["type"] == "evidenceVisual" and n["attrs"]["visual"]["type"] == "chart" for n in writeup["content_json"]["content"])
    assert {"record_type": "relationship", "record_id": writeup["cited_record_refs"][-1]["record_id"]} in writeup["cited_record_refs"]
    exported_docx = client.get(f"/api/writeups/{writeup['id']}/export.docx")
    assert exported_docx.status_code == 200
    docx = DocxDocument(BytesIO(exported_docx.content))
    paragraphs = [p.text for p in docx.paragraphs]
    assert any("Chart image not available" in p for p in paragraphs)
    cells = "\n".join(c.text for t in docx.tables for r in t.rows for c in r.cells)
    assert "Meridian Holdings Ltd · 1 payment †" in cells
