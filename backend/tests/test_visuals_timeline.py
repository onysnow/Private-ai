"""The timeline (visuals pass 4; app/visuals/echarts.py _timeline,
data_service.events_in_order): every dated event on its own row in date
order, a dot for a day, a line from start to end, dashed where a date is
known only to the month or year; its date and status written beside it;
checked like every other chart, and in decks and briefs."""

from __future__ import annotations

import copy
from datetime import date

import pytest
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.ai.chart_check import resolve_chart
from app.db.session import SessionLocal
from app.models.domain import Deck
from app.schemas.api import TimelineEventCreate
from app.services.decks import render_deck_pptx, slide_picture
from app.services.demo_seed import seed_demo_investigation
from app.services.presentation_style import THEMES
from app.services.timeline import create_timeline_event
from app.services.writeups import create_writeup, render_writeup_docx
from app.visuals.checker import check_visual
from app.visuals.data_service import DataRow, RecordRef, VisualData, date_bounds, event_span, run_query
from app.visuals.echarts import ChartDoesNotFit, build_chart, timeline_marks, when_text
from app.visuals.render_client import RenderUnavailable
from app.visuals.spec import QueryRef, parse_spec

DARK = THEMES["dark_academia"]
DAY = 86_400_000
TINY_PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4890000000d4944415478da63f8cfc0f01f0005020203"
    "f1a2b9780000000049454e44ae426082"
)


def _event(label: str, start: str, end: str | None = None, status: str = "verified") -> DataRow:
    return DataRow(label, 1.0, [RecordRef("timeline_event", label)], 1, f"{start}|{label}", start=start, end=end, status=status)


def _data(*rows: DataRow) -> VisualData:
    return VisualData("inv", "events_in_order", {}, "events", "count", "events", "events", list(rows), record_count=len(rows))


EVENTS = _data(_event("Meridian incorporated", "2019-03-04"), _event("Bellcastle holds the stake", "2019-06-01", "2021-02-10", "asserted"),
               _event("Orrin audit", "2020", status="disputed"))


def _spec(**extra):
    return parse_spec({"kind": "timeline", "question": "change_over_time", "title": "Bellcastle held the stake for two years",
                       "query": {"name": "events_in_order"}, **extra})


def _ms(day: date) -> int:
    return int((day - date(1970, 1, 1)).days) * DAY


# --- dates as recorded ------------------------------------------------------------------

def test_a_date_means_the_days_it_was_recorded_to_and_no_more():
    assert date_bounds("2024-02") == (date(2024, 2, 1), date(2024, 2, 29), "month")      # a leap year's February
    assert date_bounds("2021") == (date(2021, 1, 1), date(2021, 12, 31), "year")
    assert date_bounds("2021-03-04") == (date(2021, 3, 4), date(2021, 3, 4), "day")
    assert date_bounds("2021-02-30") is None and date_bounds("March 2021") is None and date_bounds(None) is None
    assert event_span("2019-06", "2021-02-10") == (date(2019, 6, 1), date(2021, 2, 10), False)
    assert event_span("2019-06-01", "2021-02-10") == (date(2019, 6, 1), date(2021, 2, 10), True)
    assert event_span("2019", "someday") is None


def test_an_event_is_drawn_at_its_day_or_across_the_days_it_covers():
    day, ranged, year = EVENTS.rows[0], EVENTS.rows[1], EVENTS.rows[2]
    assert timeline_marks(day) == (_ms(date(2019, 3, 4)),)
    # A line runs from the start of the first day to the end of the last.
    assert timeline_marks(ranged) == (_ms(date(2019, 6, 1)), _ms(date(2021, 2, 10)) + DAY)
    assert timeline_marks(year) == (_ms(date(2020, 1, 1)), _ms(date(2021, 1, 1)))
    assert when_text(day) == "4 Mar 2019"
    assert when_text(ranged) == "1 Jun 2019 – 10 Feb 2021"
    assert when_text(year) == "2020 (year only)"
    assert when_text(_event("x", "2019-06", "2021-02-10")) == "Jun 2019 – 10 Feb 2021 (not known to the day)"


# --- the drawing ------------------------------------------------------------------------

def test_events_run_down_in_date_order_each_named_with_its_date_and_status():
    built = build_chart(_spec(), EVENTS, DARK)
    option = built.option
    assert option["useUTC"] is True and option["xAxis"]["type"] == "time"
    # A category axis runs upward, so the rows are given latest first: the earliest event is at the top. No axis is turned over.
    assert [item["value"] for item in option["yAxis"]["data"]] == ["Orrin audit", "Bellcastle holds the stake", "Meridian incorporated"]
    assert not option["yAxis"].get("inverse")
    lines = {s["name"]: s for s in option["series"] if s["type"] == "line"}
    dots = next(s for s in option["series"] if s["type"] == "scatter")
    assert [d["name"] for d in dots["data"]] == ["Meridian incorporated"]
    assert dots["data"][0]["label"]["formatter"] == "4 Mar 2019 · verified"
    # Known to the day: a solid line; known only to the year: dashed (rule 12: not colour alone).
    assert lines["Bellcastle holds the stake"]["lineStyle"]["type"] == "solid"
    assert lines["Orrin audit"]["lineStyle"]["type"] == "dashed"
    assert lines["Orrin audit"]["data"][-1]["label"]["formatter"] == "2020 (year only) · disputed"
    assert "dashed where a date is known only to the month or year" in option["title"][0]["subtext"]
    assert built.alt_text == ("Timeline: Bellcastle held the stake for two years. 3 events from 4 Mar 2019 (Meridian incorporated) "
                              "to 2020 (year only) (Orrin audit).")
    assert built.data_table["columns"] == ["Event", "When", "Status", "Records"]
    assert [(r["formatted"], r["status"]) for r in built.data_table["rows"]] == [
        ("4 Mar 2019", "verified"), ("1 Jun 2019 – 10 Feb 2021", "asserted"), ("2020 (year only)", "disputed")]
    # One event is one record: nothing is flagged as resting on one record.
    assert "†" not in built.source_line


@pytest.mark.parametrize("theme_id", sorted(THEMES))
def test_a_timeline_passes_its_check_in_every_theme(theme_id):
    report = check_visual(_spec(theme=theme_id), EVENTS, THEMES[theme_id])
    assert report.ok, report.as_dict()
    assert {0, 7, 12, 14} <= set(report.checked_rules)
    assert {16, 20, 8} <= set(report.not_applicable)


def test_highlighted_events_stay_in_colour_and_bold_and_the_rest_turn_grey():
    built = build_chart(_spec(highlight=["Orrin audit"]), EVENTS, DARK)
    lines = {s["name"]: s for s in built.option["series"] if s["type"] == "line"}
    assert lines["Orrin audit"]["data"][-1]["label"]["fontWeight"] == "bold"
    assert lines["Orrin audit"]["lineStyle"]["color"] != lines["Bellcastle holds the stake"]["lineStyle"]["color"]
    assert check_visual(_spec(highlight=["Orrin audit"]), EVENTS, DARK, built).ok


def test_orders_a_timeline_cannot_take_are_refused_in_plain_words():
    with pytest.raises(ChartDoesNotFit, match="runs in date order"):
        build_chart(_spec(sort="label"), EVENTS, DARK)
    with pytest.raises(ChartDoesNotFit, match="log scale does not apply"):
        build_chart(_spec(scale="log"), EVENTS, DARK)
    with pytest.raises(ChartDoesNotFit, match="draw them as a timeline"):
        build_chart(parse_spec({"kind": "bar", "title": "Most events came in 2023", "query": {"name": "events_in_order"}}), EVENTS, DARK)
    with pytest.raises(ValueError, match="limit cannot cut"):
        _spec(limit=3)
    with pytest.raises(ValueError, match="is not valid"):
        parse_spec({"kind": "timeline", "title": "Bellcastle held the stake", "query": {"name": "events_in_order", "params": {"from": "March"}}})


# --- the checker -------------------------------------------------------------------------

def test_the_checker_reads_every_date_back_off_the_drawing():
    built = build_chart(_spec(), EVENTS, DARK)
    moved = copy.deepcopy(built)
    line = next(s for s in moved.option["series"] if s["type"] == "line" and s["name"] == "Bellcastle holds the stake")
    line["data"][1]["value"][0] += 30 * DAY
    report = check_visual(_spec(), EVENTS, DARK, moved)
    assert any(e.rule == 0 and "Bellcastle holds the stake" in e.message for e in report.errors)

    unnamed = copy.deepcopy(built)
    unnamed.option["yAxis"]["data"] = unnamed.option["yAxis"]["data"][1:]
    assert any(e.rule == 12 and "Orrin audit" in e.message for e in check_visual(_spec(), EVENTS, DARK, unnamed).errors)

    shuffled = _data(EVENTS.rows[2], EVENTS.rows[0], EVENTS.rows[1])
    out_of_order = copy.deepcopy(build_chart(_spec(), EVENTS, DARK))
    out_of_order.rows = shuffled.rows
    assert any(e.rule == 7 for e in check_visual(_spec(), EVENTS, DARK, out_of_order).errors)


def test_too_many_events_for_the_picture_is_refused_with_how_to_narrow_it():
    crowded = _data(*[_event(f"Board meeting {i:02d}", f"2020-01-{i:02d}") for i in range(1, 31)])
    report = check_visual(_spec(), crowded, DARK)
    assert any(e.rule == 14 and "narrow the timeline with from, to or status" in e.message for e in report.errors)


# --- the query, on a real investigation ------------------------------------------------------

def test_events_in_order_reads_the_timeline_with_its_dates_as_recorded():
    with SessionLocal() as db:
        inv, _ = seed_demo_investigation(db)
        data = run_query(db, inv.id, QueryRef(name="events_in_order"))
        assert data.series_kind == "events" and data.omitted == []
        assert [(r.start, r.end, r.status) for r in data.rows][:2] == [("2019-06-11", None, "verified"), ("2021-02", None, "verified")]
        assert [r.start for r in data.rows] == sorted(r.start or "" for r in data.rows)
        assert all(r.value == 1.0 and r.record_count == 1 and r.record_refs for r in data.rows)
        assert VisualData.from_dict(data.as_dict()).rows[1].start == "2021-02"

        create_timeline_event(db, TimelineEventCreate(investigation_id=inv.id, title="Riverside Gazette publishes \"Offshore ties revealed\"",
                                                      date_start="2025-01-09", precision="day"))
        create_timeline_event(db, TimelineEventCreate(investigation_id=inv.id, title="Lease signed", date_start="2022-05-01",
                                                      date_end="2022-03-01", precision="day"))
        again = run_query(db, inv.id, QueryRef(name="events_in_order"))
        # Two events of one name are told apart by their date; one that ends before it starts is left out, with the reason.
        assert {r.label for r in again.rows if r.label.startswith("Riverside")} == {
            'Riverside Gazette publishes "Offshore ties revealed" (2024-03-02)', 'Riverside Gazette publishes "Offshore ties revealed" (2025-01-09)'}
        assert [(o["reason"], o["detail"]) for o in again.omitted] == [("ends before it starts", "Lease signed: 2022-05-01 to 2022-03-01")]

        window = run_query(db, inv.id, QueryRef(name="events_in_order", params={"from": "2021", "to": "2023-12"}))
        assert [r.start for r in window.rows] == ["2021-02", "2023-11-14", "2023-11-14"]
        disputed = run_query(db, inv.id, QueryRef(name="events_in_order", params={"status": "disputed"}))
        assert [r.label for r in disputed.rows] == ["Disputed $2.4M payment from Bellcastle to Meridian"]


@pytest.fixture()
def timeline_chart():
    with SessionLocal() as db:
        inv, _ = seed_demo_investigation(db)
        chart, record = resolve_chart(db, inv.id, {"type": "chart", "spec": {
            "kind": "timeline", "question": "change_over_time", "title": "The payment came two years after Bellcastle took control",
            "query": {"name": "events_in_order"}}})
        assert chart is not None, record
        yield db, inv.id, chart


def _deck(db, inv_id, chart) -> Deck:
    row = Deck(investigation_id=inv_id, title="Chronology", slides_json=[
        {"kind": "content", "heading": "What happened when", "bullets": [], "evidence": [], "notes": "", "visual": chart}])
    db.add(row)
    db.commit()
    db.refresh(row)
    return row


def test_a_deck_holds_the_timeline_as_the_rendered_picture_or_its_events_listed(timeline_chart, monkeypatch):
    db, inv_id, chart = timeline_chart
    assert chart["report"]["ok"] and chart["data"]["series_kind"] == "events"
    drawn = []

    def fake_render(option, width, height, **kwargs):
        drawn.append(option["xAxis"]["type"])
        return TINY_PNG

    monkeypatch.setattr("app.visuals.render_client.render_png", fake_render)
    slide = list(render_deck_pptx(_deck(db, inv_id, chart), []).slides)[1]
    assert drawn == ["time"]
    assert len([s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]) == 1

    def unreachable(*args, **kwargs):
        raise RenderUnavailable("render service at http://render:8070 could not be reached: ConnectError")

    monkeypatch.setattr("app.visuals.render_client.render_png", unreachable)
    slide = list(render_deck_pptx(_deck(db, inv_id, chart), []).slides)[1]
    tables = [shape for shape in slide.shapes if shape.has_table]
    assert len(tables) == 1 and not [s for s in slide.shapes if s.has_chart]
    cells = [[cell.text for cell in row.cells] for row in tables[0].table.rows]
    assert cells[0] == ["Event", "When", "Status"]
    assert cells[1] == ["Alaric Voss appointed director of Meridian Holdings Ltd", "11 Jun 2019", "verified"]
    assert cells[2][1] == "Feb 2021 (month only)"
    note = " ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
    assert "image not available" in note and "the events are listed instead" in note


def test_a_slide_too_small_for_the_events_lists_them_instead_of_drawing_bars(timeline_chart):
    _db, _inv_id, chart = timeline_chart
    sized, note = slide_picture(chart, DARK, 4.0, 1.4)
    assert note is not None and "rule 14" in note and note.endswith("The events are listed instead.")


def test_a_brief_lists_every_event_with_its_date_and_status_beside_the_picture_or_without_it(timeline_chart, monkeypatch):
    db, inv_id, chart = timeline_chart
    content = {"type": "doc", "content": [{"type": "evidenceVisual", "attrs": {"visual": chart}}]}
    row = create_writeup(db, investigation_id=inv_id, title="Chronology", content_json=content)

    def unreachable(*args, **kwargs):
        raise RenderUnavailable("render service at http://render:8070 could not be reached: ConnectError")

    from app.services.writeups import extract_plain_text
    text = extract_plain_text(row.content_json)
    assert "2019-06-11 — Alaric Voss appointed director of Meridian Holdings Ltd" in text and "2021-02 — Bellcastle Trading Group" in text

    monkeypatch.setattr("app.visuals.render_client.render_png", unreachable)
    docx = render_writeup_docx(row, [])
    rows = [[c.text for c in r.cells] for t in docx.tables for r in t.rows]
    # Each event with its citation number, never the single-record mark: an event is one record.
    assert ["Alaric Voss appointed director of Meridian Holdings Ltd [1]", "11 Jun 2019", "verified", "1"] in rows
    assert not any("†" in r[0] for r in rows)
    assert any(r[1] == "Feb 2021 (month only)" for r in rows)
