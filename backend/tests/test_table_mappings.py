"""Spreadsheets to entities (ADR-0006, Wave C; app/services/table_mappings.py):
a ledger's columns are mapped with a FollowTheMoney mapping, what it makes
arrives as proposals, and a link is accepted only after both its ends, with
the row it was read from as its evidence."""

from __future__ import annotations

from datetime import datetime

from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.db.session import SessionLocal
from app.main import app
from app.models.domain import Document, Evidence, Investigation, RelationshipEdge, Source, Statement

client = TestClient(app)

LEDGER = (
    "Date;Payer;Payee;Amount;Currency;Payer number\n"
    "2024-03-03;Meridian Holdings Ltd;Bellcastle Trading Group;2,400,000;GBP;GB-123\n"
    "14/03/2024;Meridian Holdings Ltd;Orrin Shipping SA;150000;GBP;GB-123\n"
    "2024-04-01;Meridian Holdings Ltd;Orrin Shipping SA;lots;GBP;GB-123\n"
    ";;;;;\n"
    "2024-05-09;;Orrin Shipping SA;900;EUR;\n"
).encode()

PAYMENTS = {
    "payer": {"schema": "Company", "keys": ["Payer"], "properties": {
        "name": {"column": "Payer"}, "registrationNumber": {"column": "Payer number"}}},
    "payee": {"schema": "LegalEntity", "keys": ["Payee"], "properties": {"name": {"column": "Payee"}}},
    "payment": {"schema": "Payment", "keys": ["Date", "Payer", "Payee", "Amount"], "properties": {
        "payer": {"entity": "payer"}, "beneficiary": {"entity": "payee"},
        "amount": {"column": "Amount"}, "currency": {"column": "Currency"}, "date": {"column": "Date"}}},
}


def _investigation() -> str:
    with SessionLocal() as db:
        inv = Investigation(name="Ledger")
        db.add(inv)
        db.commit()
        return inv.id


def _ledger(inv: str) -> dict:
    reply = client.post("/api/documents/upload", data={"investigation_id": inv, "title": "ledger.csv"},
                        files={"file": ("ledger.csv", LEDGER, "text/csv")})
    assert reply.status_code == 200, reply.text
    return reply.json()


def _proposals(document_id: str, kind: str) -> list[dict]:
    rows = client.get(f"/api/documents/{document_id}/candidates", params={"candidate_type": kind}).json()
    return [row for row in rows if (row["payload"].get("rationale") or {}).get("rule") == "column_mapping"]


def test_the_table_reads_with_its_own_row_numbers_and_the_kinds_come_from_followthemoney():
    doc = _ledger(_investigation())
    table = client.get(f"/api/documents/{doc['id']}/table").json()
    assert table["columns"] == ["Date", "Payer", "Payee", "Amount", "Currency", "Payer number"]   # ; sniffed
    assert [row["number"] for row in table["rows"]] == [2, 3, 4, 6]                              # the blank row 5 is left out
    assert table["rows"][0]["cells"]["Amount"] == "2,400,000" and table["row_count"] == 4 and table["sheets"] == []

    schemas = client.get("/api/table-mapping/schemas").json()
    assert [s["name"] for s in schemas["entities"]] == ["Person", "Company", "Organization", "PublicBody", "LegalEntity"]
    payment = next(s for s in schemas["links"] if s["name"] == "Payment")
    assert (payment["source_prop"], payment["target_prop"]) == ("payer", "beneficiary")
    assert {"name": "amount", "label": "Amount", "type": "number"} in payment["properties"]
    assert not any(p["type"] == "entity" for p in payment["properties"])
    assert schemas["entities"][1]["properties"][0]["name"] == "name"


def test_a_preview_shows_each_row_and_what_it_leaves_out():
    doc = _ledger(_investigation())
    reply = client.post(f"/api/documents/{doc['id']}/table-mappings/preview", json={"entities": PAYMENTS})
    assert reply.status_code == 200, reply.text
    rows = reply.json()["rows"]
    first = {e["name"]: e for e in rows[0]["entities"]}
    assert first["payer"]["caption"] == "Meridian Holdings Ltd" and first["payer"]["properties"]["registrationNumber"] == ["GB-123"]
    assert first["payment"]["caption"] == "Meridian Holdings Ltd → Bellcastle Trading Group"
    assert first["payment"]["properties"]["date"] == ["2024-03-03"]
    assert rows[1]["problems"] == ["“14/03/2024” is not a valid date; it was left out"]
    # A date format reads it.
    dated = {**PAYMENTS, "payment": {**PAYMENTS["payment"], "properties": {**PAYMENTS["payment"]["properties"],
                                                                         "date": {"column": "Date", "format": "%d/%m/%Y"}}}}
    second = client.post(f"/api/documents/{doc['id']}/table-mappings/preview", json={"entities": dated}).json()["rows"][1]
    assert next(e for e in second["entities"] if e["name"] == "payment")["properties"]["date"] == ["2024-03-14"]


def test_running_proposes_each_entity_once_and_each_link_with_its_rows():
    inv = _investigation()
    doc = _ledger(inv)
    run = client.post(f"/api/documents/{doc['id']}/table-mappings", json={"entities": PAYMENTS})
    assert run.status_code == 200, run.text
    body = run.json()
    assert (body["entities_proposed"], body["links_proposed"], body["already_proposed"]) == (3, 2, 0)
    reasons = {(s["row"], s["reason"]) for s in body["skipped"]}
    assert (3, "“14/03/2024” is not a valid date; it was left out") in reasons
    assert (4, "“payment”: Payment amount must be a number: 'lots'") in reasons
    assert (6, "“payer” is left out: its key column is empty in this row") in reasons
    assert (6, "“payment” is missing its payer in this row") in reasons
    assert len(reasons) == 4
    assert body["mapping"]["rows_read"] == 4 and body["mapping"]["query"]["payment"]["schema"] == "Payment"

    entities = {row["payload"]["caption"]: row for row in _proposals(doc["id"], "entity")}
    meridian = entities["Meridian Holdings Ltd"]
    assert meridian["payload"]["rows"] == [2, 3, 4] and meridian["payload"]["locator"] == "rows 2, 3, 4"
    assert meridian["payload"]["suggested_schema"] == "Company" and meridian["confidence"] == 1.0
    links = {row["payload"]["caption"]: row for row in _proposals(doc["id"], "relationship")}
    assert set(links) == {"Meridian Holdings Ltd → Bellcastle Trading Group", "Meridian Holdings Ltd → Orrin Shipping SA"}
    wire = links["Meridian Holdings Ltd → Bellcastle Trading Group"]
    assert wire["payload"]["properties"]["amount"] == ["2400000"] and wire["payload"]["locator"] == "row 2"

    # The same mapping again proposes nothing new.
    again = client.post(f"/api/documents/{doc['id']}/table-mappings", json={"entities": PAYMENTS}).json()
    assert (again["entities_proposed"], again["links_proposed"], again["already_proposed"]) == (0, 0, 5)
    assert [m["id"] for m in client.get(f"/api/documents/{doc['id']}/table-mappings").json()] == [again["mapping"]["id"], body["mapping"]["id"]]

    # Explained in the mapping's terms.
    why = client.get(f"/api/extraction-candidates/{wire['id']}/explanation").json()
    assert why["method"]["id"] == "column_mapping" and why["summary"] == "Payment: Meridian Holdings Ltd → Bellcastle Trading Group"
    assert any(part["reasoning"] == ["Read from the column “Amount”, at row 2."] for part in why["parts"])

    # A link waits for its ends; then the row is its evidence.
    review = f"/api/extraction-candidates/{wire['id']}/review"
    refused = client.post(review, json={"decision": "accept"})
    assert refused.status_code == 400 and refused.json()["detail"] == "Accept “Meridian Holdings Ltd” first"
    for caption in ("Meridian Holdings Ltd", "Bellcastle Trading Group"):
        accepted = client.post(f"/api/extraction-candidates/{entities[caption]['id']}/review", json={"decision": "accept"})
        assert accepted.status_code == 200, accepted.text
    edge = client.post(review, json={"decision": "accept"})
    assert edge.status_code == 200, edge.text
    with SessionLocal() as db:
        relationship = db.get(RelationshipEdge, edge.json()["record"]["id"])
        assert relationship is not None and relationship.schema == "Payment"
        evidence = db.query(Evidence).filter(Evidence.locator == "row 2").one()
        assert evidence.quote == "Date: 2024-03-03; Payer: Meridian Holdings Ltd; Payee: Bellcastle Trading Group; Amount: 2,400,000; Currency: GBP; Payer number: GB-123"
        assert db.get(Source, evidence.source_id).id == doc["source_id"]
        name = db.query(Statement).filter(Statement.entity_id == relationship.source_entity_id, Statement.prop == "name").one()
        assert name.origin == f"document:{doc['id']}#rows 2, 3, 4"

    # An end the reporter rejected stops its links.
    client.post(f"/api/extraction-candidates/{entities['Orrin Shipping SA']['id']}/review", json={"decision": "reject"})
    orrin = client.post(f"/api/extraction-candidates/{links['Meridian Holdings Ltd → Orrin Shipping SA']['id']}/review", json={"decision": "accept"})
    assert orrin.status_code == 400 and orrin.json()["detail"] == "“Orrin Shipping SA” was rejected, so this link cannot be accepted"


def test_a_mapping_the_workbench_would_not_keep_is_refused_in_plain_words():
    doc = _ledger(_investigation())
    url = f"/api/documents/{doc['id']}/table-mappings/preview"

    def refused(entities: dict) -> str:
        reply = client.post(url, json={"entities": entities})
        assert reply.status_code == 400, reply.text
        return reply.json()["detail"]

    assert refused({"x": {"schema": "Vessel", "keys": ["Payer"]}}) == "“x”: Vessel is not a kind the Workbench keeps"
    assert refused({"x": {"schema": "Company", "keys": ["Payor"]}}) == "“x”: no column named “Payor”"
    assert refused({"x": {"schema": "Company", "keys": ["Payer"], "properties": {"parent": {"entity": "x"}}}}) \
        == "“x”: parent points at another entity; map that as a link"
    assert refused({"p": PAYMENTS["payer"], "pay": {"schema": "Payment", "keys": ["Date"], "properties": {"payer": {"entity": "p"}}}}) \
        == "“pay”: choose the entity that is the beneficiary"
    assert refused({"x": {"schema": "Company", "keys": ["Payer"], "properties": {"colour": {"column": "Payer"}}}}) == "Invalid property: colour"
    assert client.post(url, json={"entities": {}}).status_code == 422
    assert client.post(url, json={"entities": {"x": {"schema": "Company", "keys": ["Payer"], "properties": {"name": {"template": "{{Payer}}"}}}}}).status_code == 422


def test_a_workbook_is_read_sheet_by_sheet(tmp_path):
    book = Workbook()
    first = book.active
    first.title = "Notes"
    first.append(["Nothing to map"])
    ledger = book.create_sheet("Payments")
    ledger.append([None])
    ledger.append(["When", "Who", "Amount", None, "Who"])
    ledger.append([datetime(2024, 3, 3), "Meridian Holdings Ltd", 2400000.0, "x", "again"])
    ledger.append([datetime(2024, 3, 3, 14, 30), "Bellcastle", 12.5, None, None])
    path = tmp_path / "ledger.xlsx"
    book.save(path)
    inv = _investigation()
    with SessionLocal() as db:
        source = Source(investigation_id=inv, title="ledger.xlsx", source_type="document")
        db.add(source)
        db.flush()
        doc = Document(investigation_id=inv, source_id=source.id, filename="ledger.xlsx", sha256="0" * 64,
                       mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                       storage_path=str(path), extraction_status="complete")
        db.add(doc)
        db.commit()
        doc_id = doc.id

    table = client.get(f"/api/documents/{doc_id}/table", params={"sheet": "Payments"}).json()
    assert table["sheets"] == ["Notes", "Payments"] and table["sheet"] == "Payments"
    assert table["columns"] == ["When", "Who", "Amount", "Column 4", "Who (2)"]
    assert table["rows"] == [
        {"number": 3, "cells": {"When": "2024-03-03", "Who": "Meridian Holdings Ltd", "Amount": "2400000", "Column 4": "x", "Who (2)": "again"}},
        {"number": 4, "cells": {"When": "2024-03-03 14:30:00", "Who": "Bellcastle", "Amount": "12.5", "Column 4": "", "Who (2)": ""}},
    ]
    assert client.get(f"/api/documents/{doc_id}/table").json()["sheet"] == "Notes"
    missing = client.get(f"/api/documents/{doc_id}/table", params={"sheet": "Ledger"})
    assert missing.status_code == 400 and missing.json()["detail"] == "The workbook has no sheet named “Ledger”"

    run = client.post(f"/api/documents/{doc_id}/table-mappings", json={"sheet": "Payments", "entities": {
        "who": {"schema": "Company", "keys": ["Who"], "properties": {"name": {"column": "Who"}}}}}).json()
    assert run["entities_proposed"] == 2
    located = {row["payload"]["caption"]: row["payload"]["locator"] for row in _proposals(doc_id, "entity")}
    assert located == {"Meridian Holdings Ltd": "Payments, row 3", "Bellcastle": "Payments, row 4"}


def test_only_tables_can_be_mapped():
    inv = _investigation()
    memo = client.post("/api/documents/upload", data={"investigation_id": inv, "title": "memo.txt"},
                       files={"file": ("memo.txt", b"Meridian Holdings Ltd paid.", "text/plain")}).json()
    reply = client.get(f"/api/documents/{memo['id']}/table")
    assert reply.status_code == 400 and reply.json()["detail"] == "Only CSV, TSV and XLSX files can be mapped to entities"
    assert client.get("/api/documents/nope/table").status_code == 404


def test_a_mapped_table_is_not_read_again_by_the_text_name_finder(tmp_path, monkeypatch):
    """A table whose columns are already mapped (as the demo's payments table is,
    before it is read) is read for its text, but its cells are not run through
    the pattern name-finder: its names come from the mapping."""
    from app.core.config import settings
    from app.services.documents import extract_document, store_document
    from app.services.table_mappings import run_mapping

    monkeypatch.setattr(settings, "enable_local_entity_suggestions", True)
    inv = _investigation()
    with SessionLocal() as db:
        doc = store_document(db, investigation_id=inv, title="ledger.csv", filename="ledger.csv", mime_type="text/csv",
                             data=LEDGER, storage_dir=tmp_path)
        run_mapping(db, doc, PAYMENTS, sheet=None, created_by=None)
        read = extract_document(db, doc.id)
        assert read is not None and read.extraction_status == "complete"
        doc_id = doc.id
    entities = client.get(f"/api/documents/{doc_id}/candidates", params={"candidate_type": "entity"}).json()
    assert entities and all((row["payload"].get("rationale") or {}).get("rule") == "column_mapping" for row in entities)


def test_a_table_not_yet_mapped_is_not_read_as_running_text_either(tmp_path, monkeypatch):
    """A CSV read before anyone maps it (the usual order: the upload is read at once) proposes no names, claims or
    passages from its cells joined as text; mapping it later withdraws any such proposal kept from before."""
    from app.core.config import settings
    from app.models.domain import ExtractionCandidate
    from app.services.documents import extract_document, store_document
    from app.services.table_mappings import run_mapping

    monkeypatch.setattr(settings, "enable_local_entity_suggestions", True)
    inv = _investigation()
    with SessionLocal() as db:
        doc = store_document(db, investigation_id=inv, title="ledger.csv", filename="ledger.csv", mime_type="text/csv",
                             data=LEDGER, storage_dir=tmp_path)
        read = extract_document(db, doc.id)
        assert read is not None and read.extraction_status == "complete"
        doc_id = doc.id
        assert not db.query(ExtractionCandidate).filter(ExtractionCandidate.document_id == doc_id).count()
        # One left from an older reading is withdrawn when the table is mapped.
        db.add(ExtractionCandidate(investigation_id=inv, document_id=doc_id, candidate_type="entity",
                                   payload={"caption": "Payer Payee Amount", "rationale": {"rule": "capitalized_name_sequence"}}, confidence=0.5))
        db.commit()
        run_mapping(db, db.get(Document, doc_id), PAYMENTS, sheet=None, created_by=None)
        db.commit()
        old = db.query(ExtractionCandidate).filter(ExtractionCandidate.document_id == doc_id,
                                                   ExtractionCandidate.candidate_type == "entity").all()
        assert {c.payload["caption"]: c.review_status for c in old}["Payer Payee Amount"] == "rejected"
        assert all(c.review_status == "proposed" for c in old if c.payload["caption"] != "Payer Payee Amount")
