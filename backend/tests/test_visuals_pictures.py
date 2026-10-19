"""Before and after (the design doc's "before-and-after images"): pictures
upload as documents (kept as pictures, their text read by OCR), an evidence
record can be a picture with the date it was taken, and the `before_after`
kind compares two of them, checked like every chart and placed in decks and
briefs as two pictures with their dates."""

from __future__ import annotations

import io
import json

import pytest
from fastapi.testclient import TestClient
from PIL import Image
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.ai.chart_check import resolve_chart, visual_catalog_for_prompt
from app.db.session import SessionLocal
from app.main import app
from app.models.domain import Deck
from app.services.exports import _lineage_integrity
from app.services.presentation_style import THEMES
from app.services.writeups import create_writeup, render_writeup_docx
from app.visuals.checker import check_visual
from app.visuals.data_service import QueryError, VisualData, run_query
from app.visuals.echarts import ChartDoesNotFit, build_chart
from app.visuals.spec import QueryRef, parse_spec
from app.visuals.style import styled_chart

client = TestClient(app)
DARK = THEMES["dark_academia"]


def _png(colour: str, size: tuple[int, int] = (40, 30), fmt: str = "PNG") -> bytes:
    out = io.BytesIO()
    Image.new("RGB", size, colour).save(out, format=fmt)
    return out.getvalue()


def _upload(inv: str, name: str, body: bytes, mime: str = "image/png") -> dict:
    reply = client.post("/api/documents/upload", data={"investigation_id": inv, "title": name}, files={"file": (name, body, mime)})
    assert reply.status_code == 200, reply.text
    return reply.json()


def _world() -> dict[str, str]:
    inv = client.post("/api/investigations", json={"name": "Before and after"}).json()["id"]
    before = _upload(inv, "before.png", _png("green"))
    after = _upload(inv, "after.jpg", _png("grey", fmt="JPEG"), "image/jpeg")
    source = before["source_id"]

    def picture(document: dict, quote: str, taken: str | None) -> str:
        body = {"source_id": source, "quote": quote, "locator": "image", "document_id": document["id"], **({"captured_on": taken} if taken else {})}
        reply = client.post("/api/evidence", json=body)
        assert reply.status_code == 200, reply.text
        return reply.json()["id"]

    first = picture(before, "The market hall, intact", "2023-10-03")
    second = picture(after, "The market hall, destroyed", "2023-11-14")
    undated = picture(after, "A copy with no date", None)
    return {"inv": inv, "before": first, "after": second, "undated": undated, "before_doc": before["id"], "after_doc": after["id"],
            "source": source}


# --- pictures as documents and as evidence ----------------------------------------------

def test_a_picture_uploads_as_a_document_and_is_sent_back_as_the_picture_it_is():
    inv = client.post("/api/investigations", json={"name": "Pictures"}).json()["id"]
    document = _upload(inv, "scan.jpg", _png("blue", fmt="JPEG"), "image/jpeg")
    assert document["extraction_status"] == "complete"
    detail = client.get(f"/api/documents/{document['id']}").json()
    assert any("kept as a picture" in issue for issue in detail.get("extraction_issues") or [])
    reply = client.get(f"/api/documents/{document['id']}/picture")
    assert reply.status_code == 200 and reply.headers["content-type"] == "image/jpeg"
    assert reply.headers["x-content-type-options"] == "nosniff" and reply.headers["content-security-policy"] == "default-src 'none'"
    assert Image.open(io.BytesIO(reply.content)).format == "JPEG"
    # A document that is not a picture is refused as one; a missing one is not found.
    text = client.post("/api/documents/upload", data={"investigation_id": inv, "title": "Memo"},
                       files={"file": ("memo.txt", b"A memo.\n", "text/plain")}).json()
    assert client.get(f"/api/documents/{text['id']}/picture").status_code == 415
    assert client.get("/api/documents/no-such-document/picture").status_code == 404
    # A WebP is not taken (OpenAleph does not read it); bytes that are not a picture are not read as one, whatever the name says.
    assert client.post("/api/documents/upload", data={"investigation_id": inv, "title": "WebP"},
                       files={"file": ("scan.webp", _png("blue", fmt="WEBP"), "image/webp")}).status_code == 415
    broken = client.post("/api/documents/upload", data={"investigation_id": inv, "title": "Broken"},
                         files={"file": ("broken.png", b"not a picture", "image/png")})
    saved = broken.json()["detail"]["document"]   # saved, and the failed read said so (422), as for any unreadable file
    assert broken.status_code == 422 and saved["extraction_status"] == "failed"
    assert saved["extraction_error"] == "the file does not open as a picture (UnidentifiedImageError)"


def test_an_evidence_picture_must_be_a_picture_of_its_own_investigation():
    ids = _world()
    other = client.post("/api/investigations", json={"name": "Elsewhere"}).json()["id"]
    elsewhere = _upload(other, "other.png", _png("red"))
    memo = client.post("/api/documents/upload", data={"investigation_id": ids["inv"], "title": "Memo"},
                       files={"file": ("memo.txt", b"A memo.\n", "text/plain")}).json()

    def evidence(**body):
        return client.post("/api/evidence", json={"source_id": ids["source"], "quote": "A picture", **body})

    assert evidence(document_id=elsewhere["id"]).status_code == 400
    assert "not a picture" in evidence(document_id=memo["id"]).json()["detail"]
    assert "belongs to a picture" in evidence(captured_on="2023-10-03").json()["detail"]
    assert evidence(document_id=ids["before_doc"], captured_on="3 Oct 2023").status_code == 422
    assert evidence(document_id=ids["before_doc"], captured_on="2023-02-30").status_code == 400
    made = evidence(document_id=ids["before_doc"], captured_on="2023-10").json()
    assert (made["document_id"], made["captured_on"]) == (ids["before_doc"], "2023-10")


# --- the query and the drawing ------------------------------------------------------------

def test_pictures_compared_holds_the_dated_pictures_and_picks_two():
    ids = _world()
    with SessionLocal() as db:
        every = run_query(db, ids["inv"], QueryRef(name="pictures_compared"))
        pair = run_query(db, ids["inv"], QueryRef(name="pictures_compared", params={"before": ids["before"], "after": ids["after"]}))
        with pytest.raises(QueryError, match="both before and after"):
            run_query(db, ids["inv"], QueryRef(name="pictures_compared", params={"before": ids["before"]}))
        with pytest.raises(QueryError, match="no dated picture"):
            run_query(db, ids["inv"], QueryRef(name="pictures_compared", params={"before": ids["before"], "after": ids["undated"]}))
        menu = visual_catalog_for_prompt(db, ids["inv"])
    assert [(r.label, r.start, r.picture) for r in every.rows] == [
        ("The market hall, intact", "2023-10-03", ids["before_doc"]), ("The market hall, destroyed", "2023-11-14", ids["after_doc"])]
    assert [(o["reason"], o["record_id"]) for o in every.omitted] == [("date not known", ids["undated"])]
    assert [r.record_refs[0].record_id for r in pair.rows] == [ids["before"], ids["after"]]
    assert VisualData.from_dict(json.loads(json.dumps(pair.as_dict()))).rows[1].picture == ids["after_doc"]
    entry = next(q for q in menu["queries"] if q["name"] == "pictures_compared")
    assert entry["has_data"] and entry["kinds"] == ["before_after"] and "place" in entry["questions"]
    assert entry["picks"] == [{"evidence_id": ids["before"], "caption": "The market hall, intact", "taken": "2023-10-03"},
                              {"evidence_id": ids["after"], "caption": "The market hall, destroyed", "taken": "2023-11-14"}]


def _order(ids: dict[str, str], **extra) -> dict:
    return {"kind": "before_after", "question": "place", "title": "The market hall was destroyed within six weeks",
            "query": {"name": "pictures_compared", "params": {"before": ids["before"], "after": ids["after"]}}, **extra}


def _pair(ids: dict[str, str]) -> VisualData:
    with SessionLocal() as db:
        return run_query(db, ids["inv"], QueryRef(name="pictures_compared", params={"before": ids["before"], "after": ids["after"]}))


def test_the_pair_is_drawn_with_both_dates_in_its_layout_and_passes_its_check_in_every_theme():
    ids = _world()
    data = _pair(ids)
    for theme_id, theme in THEMES.items():
        report = check_visual(parse_spec({**_order(ids), "theme": theme_id}), data, theme)
        assert report.ok, (theme_id, report.errors)
        assert {1, 2, 9, 16, 20} <= set(report.not_applicable) and {0, 7, 12, 13, 15} <= set(report.checked_rules)
    built = build_chart(parse_spec(_order(ids)), data, DARK)
    pictures = built.option["pictures"]
    assert pictures["layout"] == "slider" and pictures["divider"] == 50
    assert [(p["role"], p["document_id"], p["caption"]) for p in pictures["items"]] == [
        ("before", ids["before_doc"], "Before · taken 3 Oct 2023"), ("after", ids["after_doc"], "After · taken 14 Nov 2023")]
    assert built.option["title"][0]["subtext"] == "Before: taken 3 Oct 2023. After: taken 14 Nov 2023"
    assert built.alt_text == ("Before and after: The market hall was destroyed within six weeks. Before: The market hall, intact, "
                              "taken 3 Oct 2023. After: The market hall, destroyed, taken 14 Nov 2023.")
    assert built.data_table["columns"] == ["Picture", "Taken", "Records"]
    # The design's presets and options: the layout, the divider, the captions.
    spec, draw = styled_chart(parse_spec(_order(ids)), data, {"preset": "stacked", "options": {"divider": "33", "date_captions": False}})
    restyled = build_chart(spec, data, DARK, draw=draw)
    assert restyled.option["pictures"]["layout"] == "stacked" and restyled.option["pictures"]["divider"] == 33
    assert [p["caption"] for p in restyled.option["pictures"]["items"]] == [None, None]
    assert check_visual(spec, data, DARK, restyled).ok   # the subtitle still states both dates


def test_the_checker_refuses_a_pair_out_of_order_or_not_its_records():
    ids = _world()
    with SessionLocal() as db:
        swapped = run_query(db, ids["inv"], QueryRef(name="pictures_compared", params={"before": ids["after"], "after": ids["before"]}))
    report = check_visual(parse_spec(_order(ids)), swapped, DARK)
    assert [e.rule for e in report.errors] == [7]
    assert "is not before" in report.errors[0].message
    data = _pair(ids)
    built = build_chart(parse_spec(_order(ids)), data, DARK)
    broken = json.loads(json.dumps(built.option))
    broken["pictures"]["items"][1]["document_id"] = ids["before_doc"]          # the wrong picture
    broken["pictures"]["items"][0]["caption"] = "Before · taken 2022"          # the wrong date
    broken["title"][0]["subtext"] = "Two pictures"                              # the dates not written
    report = check_visual(parse_spec(_order(ids)), data, DARK, built.__class__(**{**built.__dict__, "option": broken}))
    assert {0, 12} <= set(report.error_rules())
    with SessionLocal() as db:
        every = run_query(db, ids["inv"], QueryRef(name="pictures_compared"))
    every.rows.append(every.rows[0])
    with pytest.raises(ChartDoesNotFit, match="the query found 3"):
        build_chart(parse_spec(_order(ids)), every, DARK)
    with pytest.raises(ChartDoesNotFit, match="sort, limit and a log scale do not apply"):
        build_chart(parse_spec(_order(ids, sort="value_desc")), data, DARK)


# --- exports -----------------------------------------------------------------------------

@pytest.fixture()
def pair_chart():
    ids = _world()
    with SessionLocal() as db:
        chart, record = resolve_chart(db, ids["inv"], {"type": "chart", "spec": _order(ids)})
        assert chart is not None, record
        yield db, ids, chart


def test_a_deck_places_both_pictures_with_their_dates_and_only_its_own_investigations(pair_chart):
    db, ids, chart = pair_chart
    deck = Deck(investigation_id=ids["inv"], title="Market hall", slides_json=[
        {"kind": "content", "heading": "The market hall was destroyed", "bullets": [], "evidence": [], "notes": "", "visual": chart}])
    db.add(deck)
    db.commit()
    from app.services.decks import render_deck_pptx

    slide = list(render_deck_pptx(deck, []).slides)[1]
    pictures = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pictures) == 2 and pictures[0].left < pictures[1].left          # side by side: the slider is the app's
    text = " ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
    assert "Before · taken 3 Oct 2023" in text and "After · taken 14 Nov 2023" in text and "Before: taken 3 Oct 2023" in text

    # A stored slide naming another investigation's picture shows the reason, not the picture.
    other = client.post("/api/investigations", json={"name": "Elsewhere"}).json()["id"]
    elsewhere = _upload(other, "other.png", _png("red"))
    foreign = json.loads(json.dumps(chart))
    foreign["data"]["rows"][1]["picture"] = elsewhere["id"]
    deck.slides_json = [{**deck.slides_json[0], "visual": foreign}]
    db.commit()
    slide = list(render_deck_pptx(deck, []).slides)[1]
    text = " ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
    assert len([s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]) == 1
    assert "The after picture is not shown: the picture is not a document of this investigation." in text


def test_a_brief_places_both_pictures_with_their_dates_then_lists_them(pair_chart):
    db, ids, chart = pair_chart
    row = create_writeup(db, investigation_id=ids["inv"], title="Market hall", content_json={"type": "doc", "content": [
        {"type": "evidenceVisual", "attrs": {"visual": {**chart, "style": {"preset": "stacked"}}}}]})
    docx = render_writeup_docx(row, [])
    pair = docx.tables[0]
    assert (len(pair.rows), len(pair.columns)) == (2, 1)                       # stacked
    assert "Before · taken 3 Oct 2023" in pair.cell(0, 0).text and "After · taken 14 Nov 2023" in pair.cell(1, 0).text
    assert len(docx.inline_shapes) == 2
    listed = [[c.text for c in r.cells] for r in docx.tables[1].rows]
    assert any(r[0].startswith("The market hall, intact") and r[1] == "3 Oct 2023" for r in listed)


def test_a_backup_names_a_picture_its_evidence_points_at():
    records = {"documents": [], "sources": [], "evidence": [{"id": "e1", "source_id": "s1", "document_id": "gone"}]}
    assert "evidence e1 references missing picture document gone" in _lineage_integrity(records)["errors"]
