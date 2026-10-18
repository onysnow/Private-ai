"""extraction_issues: non-fatal parse problems surfaced alongside a
'complete' document, distinct from extraction_error (hard failure).

Feature request (2026-09-30): bulk-upload + per-document review needs the
app to tell the reporter what it could NOT parse, not just silently drop
it. These tests build real fixture files (a mixed text/blank PDF via
PyMuPDF, a DOCX with only a table, an empty text file) and confirm each
produces the expected extraction_issues entry through the real upload
route, not just the service function in isolation.
"""
import io
from fastapi.testclient import TestClient
from app.main import app

client = TestClient(app)


def new_inv(name='Issues'):
    r = client.post('/api/investigations', json={'name': name})
    assert r.status_code == 200
    return r.json()['id']


def _mixed_text_and_blank_pdf() -> bytes:
    import fitz
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Playwright QA text page for extraction issue testing.")
    doc.new_page()  # a genuinely blank page: no text, nothing to OCR into text either
    buf = doc.tobytes()
    doc.close()
    return buf


def _docx_with_only_a_table() -> bytes:
    import docx
    d = docx.Document()
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Role"
    table.cell(1, 0).text = "Alaric Voss"
    table.cell(1, 1).text = "Director"
    stream = io.BytesIO()
    d.save(stream)
    return stream.getvalue()


def test_pdf_with_a_blank_page_reports_issue_but_stays_complete():
    inv = new_inv()
    r = client.post(
        '/api/documents/upload',
        data={'investigation_id': inv, 'title': 'Mixed PDF'},
        files={'file': ('mixed.pdf', _mixed_text_and_blank_pdf(), 'application/pdf')},
    )
    assert r.status_code == 200, r.text
    doc = r.json()
    assert doc['extraction_status'] == 'complete'
    assert doc['chunk_count'] == 1  # only the text page produced a chunk
    issues = doc['extraction_issues']
    assert len(issues) == 1
    assert '1 of 2 page(s)' in issues[0]


def test_docx_with_only_a_table_reports_issue():
    inv = new_inv()
    r = client.post(
        '/api/documents/upload',
        data={'investigation_id': inv, 'title': 'Table-only DOCX'},
        files={'file': ('table.docx', _docx_with_only_a_table(), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')},
    )
    assert r.status_code == 200, r.text
    doc = r.json()
    assert doc['extraction_status'] == 'complete'
    assert doc['chunk_count'] == 0
    issues = doc['extraction_issues']
    assert len(issues) == 1
    assert 'tables' in issues[0]


def test_empty_text_file_reports_issue():
    inv = new_inv()
    r = client.post(
        '/api/documents/upload',
        data={'investigation_id': inv, 'title': 'Empty note'},
        files={'file': ('empty.txt', b'   \n\n  \n', 'text/plain')},
    )
    assert r.status_code == 200, r.text
    doc = r.json()
    assert doc['extraction_status'] == 'complete'
    assert doc['chunk_count'] == 0
    assert doc['extraction_issues'] == ['This file has no readable text content.']


def test_normal_document_has_no_issues():
    inv = new_inv()
    r = client.post(
        '/api/documents/upload',
        data={'investigation_id': inv, 'title': 'Normal memo'},
        files={'file': ('memo.txt', b'A perfectly ordinary paragraph of text.', 'text/plain')},
    )
    assert r.status_code == 200, r.text
    doc = r.json()
    assert doc['extraction_issues'] == []
