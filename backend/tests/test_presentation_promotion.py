"""Tests for promoting an accepted presentation_draft AIAnalysisCandidate
into a Writeup (brief) and/or Deck -- app/services/presentation.py, wired
up as /api/ai-analysis-candidates/{id}/promote-to-writeup and
.../promote-to-deck in app/api/routes_system.py.
"""

import json

import pytest
from fastapi.testclient import TestClient

from app.db.session import SessionLocal
from app.main import app
from app.models.domain import AIAnalysisCandidate
from app.core.config import settings
from app.core.time import utcnow_naive
import app.ai.reasoning as reasoning_module
from app.ai.governor import governor

client = TestClient(app)


class _FakeLLMClient:
    def __init__(self, response_text: str):
        self._response_text = response_text

    def generate(self, system_prompt: str, user_prompt: str, *, max_tokens: int = 4096) -> str:
        return self._response_text


@pytest.fixture(autouse=True)
def _restore_ai_settings():
    original = (settings.enable_ai_features, settings.ai_provider, settings.ai_runs_per_minute, settings.ai_max_concurrent_runs)
    governor.reset()
    yield
    settings.enable_ai_features, settings.ai_provider, settings.ai_runs_per_minute, settings.ai_max_concurrent_runs = original
    governor.reset()


def _seeded_investigation() -> dict:
    inv = client.post("/api/investigations", json={"name": "Promotion tests"}).json()
    claim = client.post("/api/claims", json={"investigation_id": inv["id"], "text": "Meridian paid Bellcastle a consulting fee in November 2023."}).json()
    source = client.post("/api/sources", json={"investigation_id": inv["id"], "title": "Wire transfer log", "source_type": "filing", "url": "https://example.test/wire-log"}).json()
    evidence = client.post("/api/evidence", json={"source_id": source["id"], "locator": "row 41",
                                                  "quote": "$2.4M wired to Bellcastle Holdings on 2023-11-14. Agreement signed 2023-11-03. Arden was paid 285,000 USD."}).json()
    return {"investigation": inv, "claim": claim, "evidence": evidence}


def _draft_response(claim_id: str, evidence_id: str) -> str:
    return json.dumps({
        "title": "The Bellcastle consulting fee",
        "goal_restated": "Summarize what the selected evidence shows about the Bellcastle payment.",
        "executive_takeaway": "Meridian paid Bellcastle $2.4M in November 2023, consistent with the wire log.",
        "sections": [{
            "heading": "The payment",
            "body": "Meridian paid Bellcastle Holdings a $2.4M consulting fee in November 2023.",
            "key_points": [
                {"text": "$2.4M paid to Bellcastle", "evidence": [{"record_type": "evidence", "record_id": evidence_id, "quote": "$2.4M wired to Bellcastle Holdings"}]},
                {"text": "Dated November 2023", "evidence": [{"record_type": "claim", "record_id": claim_id, "quote": "consulting fee in  November 2023"}]},
            ],
            "evidence": [
                {"record_type": "claim", "record_id": claim_id, "quote": "Meridian paid Bellcastle a consulting fee"},
                {"record_type": "evidence", "record_id": evidence_id, "quote": "$2.4M wired to Bellcastle"},
            ],
        }],
        "open_questions": ["Whether the fee was disclosed to the board."],
    })


def _accepted_candidate(monkeypatch) -> dict:
    """Seeds an investigation, runs presentation_draft against a fake LLM,
    and accepts the resulting candidate. Returns the seeded dict plus the
    candidate id."""
    settings.enable_ai_features = True
    settings.ai_provider = "anthropic"
    seeded = _seeded_investigation()
    claim_id, evidence_id = seeded["claim"]["id"], seeded["evidence"]["id"]
    monkeypatch.setattr(reasoning_module, "get_llm_client", lambda *a, **k: _FakeLLMClient(_draft_response(claim_id, evidence_id)))

    r = client.post(
        f"/api/investigations/{seeded['investigation']['id']}/assistant/presentation-draft",
        json={
            "goal": "Summarize the Bellcastle payment.",
            "record_refs": [
                {"record_type": "claim", "record_id": claim_id},
                {"record_type": "evidence", "record_id": evidence_id},
            ],
        },
    )
    assert r.status_code == 200, r.text
    candidate_id = r.json()["candidate_id"]

    reviewed = client.post(f"/api/ai-analysis-candidates/{candidate_id}/review", json={"decision": "accept", "note": "looks right"})
    assert reviewed.status_code == 200, reviewed.text
    assert reviewed.json()["review_status"] == "accepted"

    seeded["candidate_id"] = candidate_id
    return seeded


def test_promote_to_writeup_builds_a_content_json_with_matching_citations(monkeypatch):
    seeded = _accepted_candidate(monkeypatch)
    r = client.post(f"/api/ai-analysis-candidates/{seeded['candidate_id']}/promote-to-writeup", json={})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["title"] == "The Bellcastle consulting fee"
    assert body["status"] == "draft"
    assert body["source_ai_candidate_id"] == seeded["candidate_id"]
    assert sorted(body["cited_record_refs"], key=str) == sorted(
        [
            {"record_type": "claim", "record_id": seeded["claim"]["id"]},
            {"record_type": "evidence", "record_id": seeded["evidence"]["id"]},
        ],
        key=str,
    )
    nodes = body["content_json"]["content"]
    # The title is the write-up's own field, not repeated as a heading.
    assert not any(n["type"] == "heading" and n["content"][0]["text"] == body["title"] for n in nodes)
    takeaway = next(n for n in nodes if n["type"] == "callout" and n["attrs"]["tone"] == "takeaway")
    assert "Meridian paid Bellcastle $2.4M" in takeaway["content"][0]["content"][0]["text"]
    body_paragraph = next(n for n in nodes if n["type"] == "paragraph" and "consulting fee in November 2023" in (n.get("content") or [{}])[0].get("text", ""))
    citations = [c for c in body_paragraph["content"] if c["type"] == "citation"]
    # Citation labels are readable (evidence -> its locator), never a raw UUID.
    assert {c["attrs"]["label"] for c in citations} >= {"row 41"}
    assert any(n["type"] == "callout" and n["attrs"]["tone"] == "questions" for n in nodes)
    assert body["word_count"] > 0


def test_promote_to_writeup_title_override(monkeypatch):
    seeded = _accepted_candidate(monkeypatch)
    r = client.post(f"/api/ai-analysis-candidates/{seeded['candidate_id']}/promote-to-writeup", json={"title": "Custom title"})
    assert r.status_code == 200, r.text
    assert r.json()["title"] == "Custom title"


def test_promote_to_deck_builds_one_slide_per_section_plus_takeaway_and_open_questions(monkeypatch):
    seeded = _accepted_candidate(monkeypatch)
    r = client.post(f"/api/ai-analysis-candidates/{seeded['candidate_id']}/promote-to-deck", json={})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["title"] == "The Bellcastle consulting fee"
    assert body["source_ai_candidate_id"] == seeded["candidate_id"]
    kinds = [slide["kind"] for slide in body["slides"]]
    # The last slide before the open questions says how we know and what the evidence cannot show.
    assert kinds == ["title", "statement", "content", "questions", "questions"]
    assert body["slides"][3]["heading"] == "How we know, and our limits"
    assert body["slides"][3]["bullets"][0].startswith("This rests on ") and body["slides"][3]["bullets"][0].endswith("† marks a statement that rests on one record.")
    assert body["slides"][0]["heading"] == "The Bellcastle consulting fee"
    payment_slide = body["slides"][2]
    assert payment_slide["heading"] == "The payment"
    # Each key point rests on one record here, so each carries the single-record mark.
    assert payment_slide["bullets"] == ["$2.4M paid to Bellcastle †", "Dated November 2023 †"]
    # The section's full prose rides along as speaker notes.
    assert payment_slide["notes"].startswith("Meridian paid Bellcastle Holdings")
    assert {"record_type": "claim", "record_id": seeded["claim"]["id"]} in payment_slide["evidence"]
    assert {"record_type": "evidence", "record_id": seeded["evidence"]["id"]} in payment_slide["evidence"]


def test_promoting_the_same_candidate_twice_creates_two_separate_writeups(monkeypatch):
    """Not idempotent by design -- Ony's 'both, picked per use' requirement
    means a reporter may legitimately want more than one artifact from the
    same accepted draft."""
    seeded = _accepted_candidate(monkeypatch)
    first = client.post(f"/api/ai-analysis-candidates/{seeded['candidate_id']}/promote-to-writeup", json={})
    second = client.post(f"/api/ai-analysis-candidates/{seeded['candidate_id']}/promote-to-writeup", json={})
    assert first.status_code == 200 and second.status_code == 200
    assert first.json()["id"] != second.json()["id"]


def test_promotion_requires_accepted_review_status(monkeypatch):
    settings.enable_ai_features = True
    settings.ai_provider = "anthropic"
    seeded = _seeded_investigation()
    claim_id, evidence_id = seeded["claim"]["id"], seeded["evidence"]["id"]
    monkeypatch.setattr(reasoning_module, "get_llm_client", lambda *a, **k: _FakeLLMClient(_draft_response(claim_id, evidence_id)))
    r = client.post(
        f"/api/investigations/{seeded['investigation']['id']}/assistant/presentation-draft",
        json={
            "goal": "Summarize the Bellcastle payment.",
            "record_refs": [
                {"record_type": "claim", "record_id": claim_id},
                {"record_type": "evidence", "record_id": evidence_id},
            ],
        },
    )
    assert r.status_code == 200, r.text
    candidate_id = r.json()["candidate_id"]  # still "proposed" -- never reviewed

    promote = client.post(f"/api/ai-analysis-candidates/{candidate_id}/promote-to-writeup", json={})
    assert promote.status_code == 400, promote.text
    assert "accepted" in promote.text


def test_promotion_rejects_wrong_module():
    """Only presentation_draft candidates map to a writeup/deck shape --
    case_synthesis/hypothesis_test payloads don't have sections/key_points.
    Built directly against the DB (bypassing the reasoning endpoint) so this
    doesn't depend on the vendored TAS submodule being checked out, which
    neither this check nor its candidate's existence requires."""
    inv = client.post("/api/investigations", json={"name": "Wrong module promotion test"}).json()
    with SessionLocal() as db:
        row = AIAnalysisCandidate(
            investigation_id=inv["id"], module="case_synthesis",
            payload={"summary": "not a presentation_draft shape"},
            review_status="accepted", reviewed_at=utcnow_naive(),
        )
        db.add(row)
        db.commit()
        candidate_id = row.id

    promote = client.post(f"/api/ai-analysis-candidates/{candidate_id}/promote-to-writeup", json={})
    assert promote.status_code == 400, promote.text
    assert "presentation_draft" in promote.text

    promote_deck = client.post(f"/api/ai-analysis-candidates/{candidate_id}/promote-to-deck", json={})
    assert promote_deck.status_code == 400, promote_deck.text
    assert "presentation_draft" in promote_deck.text


# --- v2: readable citations, visuals, and the designed exports ----------------

def _run_draft(monkeypatch, seeded: dict, payload: dict):
    settings.enable_ai_features = True
    settings.ai_provider = "anthropic"
    monkeypatch.setattr(reasoning_module, "get_llm_client", lambda *a, **k: _FakeLLMClient(json.dumps(payload)))
    return client.post(
        f"/api/investigations/{seeded['investigation']['id']}/assistant/presentation-draft",
        json={"goal": "Summarize the Bellcastle payment.", "record_refs": [
            {"record_type": "claim", "record_id": seeded["claim"]["id"]},
            {"record_type": "evidence", "record_id": seeded["evidence"]["id"]},
        ]},
    )


def _payload_with_visuals(claim_id: str, evidence_id: str, **section_overrides) -> dict:
    payload = json.loads(_draft_response(claim_id, evidence_id))
    ev = [{"record_type": "evidence", "record_id": evidence_id, "quote": "$2.4M wired to Bellcastle"}]

    def cites(quote: str) -> list[dict]:
        # Each visual item cites the words that state its own value or date.
        return [{"record_type": "evidence", "record_id": evidence_id, "quote": quote}]

    payload["sections"][0]["visual"] = {"type": "figures", "items": [{"value": "$2.4M", "label": "wired 14 Nov 2023", "evidence": ev}]}
    payload["sections"].append({
        "heading": "Sequence", "body": "The wire followed the agreement.", "key_points": [{"text": "Agreement, then wire", "evidence": ev}], "evidence": ev,
        "visual": {"type": "timeline", "items": [
            {"date": "2023-11-03", "label": "Agreement signed", "evidence": cites("Agreement signed 2023-11-03")},
            {"date": "2023-11-14", "label": "$2.4M wired", "evidence": cites("Bellcastle Holdings on 2023-11-14")},
        ]},
    })
    payload["sections"].append({
        "heading": "Compared", "body": "The fee dwarfs the others.", "key_points": [{"text": "Largest fee", "evidence": ev}], "evidence": ev,
        "visual": {"type": "bar", "title": "Consulting fees", "unit": "USD", "items": [
            {"label": "Bellcastle", "value": 2400000, "evidence": ev},
            {"label": "Arden", "value": 285000, "evidence": cites("Arden was paid 285,000 USD")},
        ]},
    })
    payload["sections"][-1].update(section_overrides)
    return payload


def test_draft_with_cited_visuals_is_accepted(monkeypatch):
    seeded = _seeded_investigation()
    r = _run_draft(monkeypatch, seeded, _payload_with_visuals(seeded["claim"]["id"], seeded["evidence"]["id"]))
    assert r.status_code == 200, r.text
    assert r.json()["review_status"] == "proposed"


@pytest.mark.parametrize("bad_visual, message", [
    ({"type": "bar", "title": "x", "unit": "USD", "items": [{"label": "a", "value": "lots", "evidence": "EV"}, {"label": "b", "value": 2, "evidence": "EV"}]}, "value must be a number"),
    ({"type": "bar", "title": "x", "unit": "USD", "items": [{"label": "a", "value": 1, "evidence": []}, {"label": "b", "value": 2, "evidence": "EV"}]}, "must cite at least one record"),
    ({"type": "timeline", "items": [{"date": "2023", "label": "only one point", "evidence": "EV"}]}, "2-8 items"),
    ({"type": "pie", "items": []}, "type must be one of"),
])
def test_malformed_or_uncited_visual_is_rejected(monkeypatch, bad_visual, message):
    seeded = _seeded_investigation()
    ev = [{"record_type": "evidence", "record_id": seeded["evidence"]["id"]}]
    def fill(v):
        return json.loads(json.dumps(v).replace('"EV"', json.dumps(ev)))
    r = _run_draft(monkeypatch, seeded, _payload_with_visuals(seeded["claim"]["id"], seeded["evidence"]["id"], visual=fill(bad_visual)))
    assert r.status_code == 422, r.text
    assert message in r.text


def test_visual_citing_a_record_outside_the_selection_is_rejected(monkeypatch):
    seeded = _seeded_investigation()
    payload = _payload_with_visuals(seeded["claim"]["id"], seeded["evidence"]["id"])
    payload["sections"][0]["visual"]["items"][0]["evidence"] = [{"record_type": "evidence", "record_id": "NOT-SELECTED"}]
    r = _run_draft(monkeypatch, seeded, payload)
    assert r.status_code == 422, r.text
    assert "NOT-SELECTED" in r.text


def test_candidate_detail_has_readable_citations_and_a_deck_preview(monkeypatch):
    seeded = _seeded_investigation()
    r = _run_draft(monkeypatch, seeded, _payload_with_visuals(seeded["claim"]["id"], seeded["evidence"]["id"]))
    detail = client.get(f"/api/ai-analysis-candidates/{r.json()['candidate_id']}").json()
    by_type = {d["record_type"]: d for d in detail["citation_details"]}
    assert by_type["evidence"]["label"] == "row 41"
    assert by_type["evidence"]["source_title"] == "Wire transfer log"
    assert by_type["evidence"]["quote"].startswith("$2.4M wired to Bellcastle")
    assert by_type["claim"]["label"].startswith("Meridian paid Bellcastle a consulting fee")
    kinds = [s["kind"] for s in detail["deck_preview"]]
    assert kinds == ["title", "statement", "content", "content", "content", "questions", "questions"]
    assert [s["visual"]["type"] for s in detail["deck_preview"] if s.get("visual")] == ["figures", "timeline", "bar"]


def test_describe_record_refs_handles_missing_records_and_task_spellings():
    from app.services.assistant_context import describe_record_refs
    inv = client.post("/api/investigations", json={"name": "describe refs"}).json()
    with SessionLocal() as db:
        out = describe_record_refs(db, investigation_id=inv["id"], refs=[
            {"record_type": "claim", "record_id": "gone"},
            {"record_type": "reporting_task", "record_id": "gone-too"},
            {"record_type": "claim", "record_id": "gone"},  # duplicate
        ])
    assert len(out) == 2
    assert all(d["missing"] for d in out)
    assert out[1]["trace_record_type"] == "task" and out[1]["type_label"] == "Task"


def test_promoted_deck_exports_native_chart_timeline_and_speaker_notes(monkeypatch):
    from io import BytesIO
    from pptx import Presentation
    seeded = _seeded_investigation()
    r = _run_draft(monkeypatch, seeded, _payload_with_visuals(seeded["claim"]["id"], seeded["evidence"]["id"]))
    candidate_id = r.json()["candidate_id"]
    client.post(f"/api/ai-analysis-candidates/{candidate_id}/review", json={"decision": "accept"})
    deck = client.post(f"/api/ai-analysis-candidates/{candidate_id}/promote-to-deck", json={}).json()
    exported = client.get(f"/api/decks/{deck['id']}/export.pptx")
    assert exported.status_code == 200
    prs = Presentation(BytesIO(exported.content))
    slides = list(prs.slides)
    # title, takeaway, 3 content, how we know, questions, sources
    assert len(slides) == 8
    assert sum(1 for s in slides for sh in s.shapes if sh.has_chart) == 1
    text = "\n".join(sh.text_frame.text for s in slides for sh in s.shapes if sh.has_text_frame)
    assert "3 Nov 2023" in text and "14 Nov 2023" in text  # timeline dates, human-formatted
    assert "$2.4M" in text  # figure card
    assert "row 41" in text and "Wire transfer log" in text  # readable sources
    assert seeded["evidence"]["id"] not in text
    notes = [s.notes_slide.notes_text_frame.text for s in slides if s.has_notes_slide]
    assert any(n.startswith("Meridian paid Bellcastle Holdings") for n in notes)


def test_promoted_writeup_counts_visual_citations_and_exports_tables(monkeypatch):
    from io import BytesIO
    from docx import Document as DocxDocument
    seeded = _seeded_investigation()
    r = _run_draft(monkeypatch, seeded, _payload_with_visuals(seeded["claim"]["id"], seeded["evidence"]["id"]))
    candidate_id = r.json()["candidate_id"]
    client.post(f"/api/ai-analysis-candidates/{candidate_id}/review", json={"decision": "accept"})
    writeup = client.post(f"/api/ai-analysis-candidates/{candidate_id}/promote-to-writeup", json={}).json()
    assert any(n["type"] == "evidenceVisual" for n in writeup["content_json"]["content"])
    exported = client.get(f"/api/writeups/{writeup['id']}/export.docx")
    assert exported.status_code == 200
    docx = DocxDocument(BytesIO(exported.content))
    paragraphs = [p.text for p in docx.paragraphs]
    # Masthead title appears once; no duplicate heading of the same text.
    assert paragraphs.count(writeup["title"]) == 1
    cells = "\n".join(c.text for t in docx.tables for row in t.rows for c in row.cells)
    assert "THE TAKEAWAY" in cells and "OPEN QUESTIONS" in cells
    assert "3 Nov 2023" in cells and "$2.4M" in cells
    assert "Sources" in paragraphs


def test_a_slidedoc_puts_each_sections_prose_on_its_slide_and_the_export_writes_it_plain(monkeypatch):
    """Duarte's two kinds of deck: a deck to present keeps the prose in the
    speaker notes; a slidedoc, read alone, puts it on the slide in place of
    the bullets, as plain paragraphs."""
    from io import BytesIO

    from pptx import Presentation

    seeded = _accepted_candidate(monkeypatch)
    presenting = client.post(f"/api/ai-analysis-candidates/{seeded['candidate_id']}/promote-to-deck", json={}).json()
    assert all(slide["prose"] == "" for slide in presenting["slides"])
    deck = client.post(f"/api/ai-analysis-candidates/{seeded['candidate_id']}/promote-to-deck", json={"deck_kind": "slidedoc"}).json()
    slide = deck["slides"][2]
    assert slide["prose"].startswith("Meridian paid Bellcastle Holdings") and slide["notes"] == slide["prose"]
    assert client.post(f"/api/ai-analysis-candidates/{seeded['candidate_id']}/promote-to-deck", json={"deck_kind": "poster"}).status_code == 422

    prs = Presentation(BytesIO(client.get(f"/api/decks/{deck['id']}/export.pptx").content))
    shapes = [shape for shape in list(prs.slides)[2].shapes if shape.has_text_frame and shape.text_frame.text.startswith("Meridian paid Bellcastle Holdings")]
    assert shapes, "the prose is on the slide"
    paragraph = shapes[0].text_frame.paragraphs[0]._p
    assert paragraph.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}buChar") is None   # plain prose, no bullet
